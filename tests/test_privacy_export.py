"""脱敏导出必须阻断原数据和各类缓存，同时能在本机还原匿名代码的数据接口。

Copyright (c) 2026 AbeLin · MIT License
"""
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
import pandas as pd
import pytest

from desktop.bridge import selection_config
from desktop.privacy_export import export_privacy, inspect_privacy, load_ai_context
from desktop.privacy_formula import FormulaSanitizer

PRIVATE = 'PRIVATE_CUSTOMER_913579246821'
AMOUNT = 87654321987


@pytest.fixture
def source(tmp_path):
    path = tmp_path / f'{PRIVATE}.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '机密销售'
    ws.merge_cells('A1:F2')
    ws['A1'] = PRIVATE
    ws['A1'].font = Font(name=PRIVATE, bold=True, color='FFFFFF', size=20)
    ws['A1'].fill = PatternFill('solid', fgColor='155D45')
    for col, value in enumerate([PRIVATE, '数量', '金额', '计算', '类别', '隐藏字段'], 1):
        ws.cell(3, col, value)
    for r in range(4, 304):
        ws.append([PRIVATE, 99, AMOUNT, f'=B{r}*C{r}*0.13579+9876543', PRIVATE, PRIVATE])
    ws['A4'].comment = Comment(PRIVATE, PRIVATE)
    ws['A4'].hyperlink = f'https://example.invalid/{PRIVATE}'
    ws['C4'].number_format = f'0.00"{PRIVATE}"'
    ws.column_dimensions['F'].hidden = True
    ws.row_dimensions[9].hidden = True
    ws.column_dimensions['C'].width = 24
    ws.conditional_formatting.add('C4:C303', CellIsRule(operator='greaterThan', formula=[str(AMOUNT)], font=Font(color='FF0000')))
    validation = DataValidation(type='list', formula1=f'"{PRIVATE},B"')
    ws.add_data_validation(validation)
    validation.add('E4:E303')
    chart = BarChart()
    chart.title = PRIVATE
    chart.add_data(Reference(ws, min_col=3, min_row=3, max_row=8), titles_from_data=True)
    ws.add_chart(chart, 'H3')
    secret = wb.create_sheet('隐藏机密')
    secret['A1'] = PRIVATE
    secret.sheet_state = 'veryHidden'
    wb.defined_names.add(DefinedName(PRIVATE, attr_text="'机密销售'!$C$4"))
    wb.properties.creator = PRIVATE
    wb.properties.description = PRIVATE
    wb.save(path)
    wb.close()
    with ZipFile(path, 'a') as z:
        z.writestr('xl/pivotCache/pivotCacheRecords99.xml', f'<cache>{PRIVATE}{AMOUNT}</cache>')
        z.writestr('customXml/item1.xml', f'<secret>{PRIVATE}</secret>')
    return path


def package(source, directory, mode='synthetic', count=20, overrides=None):
    """经过真实选择校验和准备 token 导出到隔离目录。"""
    directory.mkdir(parents=True, exist_ok=True)
    config = selection_config({'inputPath': str(source), 'selection': [{'path': str(source), 'sheets': ['机密销售']}], 'headerRow': 3})
    plan = inspect_privacy(config)
    request = {'mode': mode, 'sampleRows': count, 'token': plan['token'],
               'fieldKinds': {table['id']: [f['kind'] for f in table['fields']] for table in plan['tables']},
               'outputPath': str(directory / 'share.zip'), 'exportId': str(uuid4()), 'mappingDir': str(directory / 'private-app-data')}
    if overrides:
        request.update(overrides)
    before = source.read_bytes()
    result = export_privacy(request, config)
    assert source.read_bytes() == before
    return request, result, plan


def contents(path):
    """解开外层包和内部工作簿检查，不能只扫描压缩后的字节。"""
    result = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            data = archive.read(name)
            if name.endswith('.xlsx'):
                with ZipFile(BytesIO(data)) as workbook:
                    result.update({name + '/' + key: workbook.read(key) for key in workbook.namelist()})
            else:
                result[name] = data
    return result


def test_synthetic_export_contains_no_private_cells_metadata_or_caches(source, tmp_path):
    request, result, plan = package(source, tmp_path / 'out')
    assert len(plan['tables'][0]['fields']) == 5
    assert plan['tables'][0]['fields'][0]['name'] == PRIVATE  # 只在本机选择器出现
    expanded = contents(result['output'])
    for name, payload in expanded.items():
        text = payload.decode('utf-8')
        for private in (PRIVATE, str(AMOUNT), '机密销售', '隐藏机密', '0.13579', '9876543', str(source.parent)):
            assert private not in text, (name, private)
        assert 'pivotCache' not in name and 'drawing' not in name and 'comments' not in name
        assert 'TargetMode="External"' not in text
    with ZipFile(result['output']) as archive:
        assert set(archive.namelist()) == {'结构说明.json', '给AI的说明.txt', '虚构演示.xlsx'}
        book = openpyxl.load_workbook(BytesIO(archive.read('虚构演示.xlsx')))
        assert book.sheetnames == ['Sheet_001', '演示参数']
        ws = book['Sheet_001']
        assert ws.max_row == 23  # 源表有 300 条记录，导出始终只有指定的 20 行
        assert str(ws.merged_cells) == 'A1:F2' or str(ws.merged_cells) == ''  # 隐藏列不带入标题合并
        assert ws['A3'].value == '字段_001'
        assert ws['A4'].value == '示例文本_01'
        assert ws['C4'].value == 1022.5
        assert ws['D4'].value == "=B4*C4*'演示参数'!$B$2+'演示参数'!$B$3"
        assert ws['D23'].value == "=B23*C23*'演示参数'!$B$2+'演示参数'!$B$3"
        assert ws['F4'].value is None and ws.column_dimensions['F'].hidden
        assert not ws._charts and not ws._images and not ws.data_validations.dataValidation
        assert not ws.conditional_formatting and not book.defined_names
        assert ws['C4'].number_format == '#,##0.00'
        assert ws.column_dimensions['C'].width == 24
        book.close()
    private = json.loads((Path(request['mappingDir']) / f"{result['exportId']}.json").read_text(encoding='utf-8'))
    assert private['tables'][0]['fields'][0]['name'] == PRIVATE
    assert private['parameters']['参数_0001']['value'] == .13579
    assert private['parameters']['参数_0002']['value'] == 9876543


@pytest.mark.parametrize('mode', ['blank', 'schema'])
def test_blank_and_schema_modes_do_not_emit_values_or_real_counts(source, tmp_path, mode):
    _request, result, _plan = package(source, tmp_path / mode, mode=mode, count=5)
    with ZipFile(result['output']) as archive:
        schema = json.loads(archive.read('结构说明.json'))
        if mode == 'schema':
            assert len(archive.namelist()) == 2
            assert 'sampleRows' not in schema and not schema['parameters']
        else:
            book = openpyxl.load_workbook(BytesIO(archive.read('空白模板.xlsx')))
            assert book['Sheet_001'].max_row == 8
            assert book['Sheet_001']['A4'].value is None
            assert book['Sheet_001']['D4'].data_type == 'f'
            assert book['演示参数']['B2'].value is None
            book.close()
        assert 'rows' not in schema['tables'][0]
        assert set(schema['tables'][0]) == ({'id', 'fields'} if mode == 'schema' else {'id', 'fields', 'headerRow'})


def test_demo_values_do_not_depend_on_original_values_or_row_count(source, tmp_path):
    _, first, _ = package(source, tmp_path / 'first')
    wb = openpyxl.load_workbook(source)
    ws = wb['机密销售']
    for row in ws.iter_rows(min_row=4, max_row=603):
        for cell in row:
            if cell.column != 4:
                cell.value = 'DIFFERENT_PRIVATE_001' if cell.column in (1, 5, 6) else 123456.78
    wb.save(source)
    wb.close()
    _, second, _ = package(source, tmp_path / 'second')
    a, b = contents(first['output']), contents(second['output'])
    for name in a:
        if name.endswith('.xml') and not name.endswith('core.xml'):
            assert a[name] == b[name], name


def test_local_context_returns_real_data_with_anonymous_columns_and_validates_drift(source, tmp_path):
    request, result, _ = package(source, tmp_path / 'out')
    frame = pd.read_excel(source, sheet_name='机密销售', header=2).iloc[:, :5]
    tables = {'original': frame}
    inputs = [{'key': 'original', 'path': str(source), 'sheet': '机密销售', 'headerRow': 3}]
    aliases, params = load_ai_context(result['exportId'], request['mappingDir'], tables, inputs)
    assert aliases['Sheet_001'].iloc[0, 2] == AMOUNT
    assert aliases['Sheet_001'].columns.tolist() == [f'字段_{i:03d}' for i in range(1, 6)]
    assert params['参数_0002'] == 9876543
    assert frame.columns[0] == PRIVATE  # 对原数据框不改名
    with pytest.raises(ValueError, match='Sheet'):
        load_ai_context(result['exportId'], request['mappingDir'], tables, [])
    with pytest.raises(ValueError, match='列结构'):
        load_ai_context(result['exportId'], request['mappingDir'], {'original': frame.rename(columns={PRIVATE: 'changed'})}, inputs)
    with pytest.raises(ValueError, match='本机没有'):
        load_ai_context('../file', request['mappingDir'], tables, inputs)
    wb = openpyxl.load_workbook(source)
    wb['机密销售']['D4'] = '=B4*C4*0.5'
    wb.save(source)
    wb.close()
    with pytest.raises(ValueError, match='源公式'):
        load_ai_context(result['exportId'], request['mappingDir'], tables, inputs)


def test_narrow_columns_match_offline_runner_selection(source, tmp_path):
    from core.xlio import hidden_info
    wb = openpyxl.load_workbook(source)
    wb['机密销售'].column_dimensions['B'].width = 0.25
    wb.save(source)
    wb.close()
    request, result, plan = package(source, tmp_path / 'out')
    assert [field['column'] for field in plan['tables'][0]['fields']] == [1, 3, 4, 5]
    frame = pd.read_excel(source, sheet_name='机密销售', header=2)
    excluded = hidden_info(str(source))['机密销售'][1]
    frame = frame.iloc[:, [i for i in range(len(frame.columns)) if i+1 not in excluded]]
    inputs = [{'key': 'original', 'path': str(source), 'sheet': '机密销售', 'headerRow': 3}]
    aliases, _ = load_ai_context(result['exportId'], request['mappingDir'], {'original': frame}, inputs)
    assert aliases['Sheet_001'].shape == (300, 4)
    assert aliases['Sheet_001'].iloc[0, 1] == AMOUNT


@pytest.mark.parametrize('formula', ['=HYPERLINK("https://secret.invalid/", "private")', '=WEBSERVICE("https://secret.invalid/")',
    '=INDIRECT("private!A1")', '=PRIVATE_NAME', '=SUM([secret.xlsx]Sheet1!A1:A3)', '=私密函数(A1)', '=TableSecret[工资]'])
def test_formula_allowlist_rejects_private_names_external_or_dynamic_refs(formula):
    sheet = {'id': 'Sheet_001', 'path': 'local.xlsx', 'sheet': '私密', 'width': 3, 'visible': [1, 2, 3], 'header': 1}
    sanitizer = FormulaSanitizer([sheet], 20, 'synthetic')
    with pytest.raises(ValueError):
        sanitizer.sanitize(formula, 'C2', 'C2', sheet, '字段_003')
    assert not sanitizer.parameters


def test_formula_ranges_and_text_constants_are_redacted():
    first = {'id': 'Sheet_001', 'path': 'local.xlsx', 'sheet': '机密一', 'width': 3, 'visible': [1, 2, 3], 'header': 1}
    second = {**first, 'id': 'Sheet_002', 'sheet': '机密二'}
    sanitizer = FormulaSanitizer([first, second], 20, 'synthetic')
    result = sanitizer.sanitize('=SUMIF(\'机密二\'!A2:A98765,"PRIVATE",\'机密二\'!C2:C98765)', 'C2', 'C2', first, '字段_003')
    assert result == "=SUMIF('Sheet_002'!A2:A21,'演示参数'!$B$2,'Sheet_002'!C2:C21)"
    assert sanitizer.parameters['参数_0001']['value'] == 'PRIVATE'


def test_stale_plan_or_invalid_options_never_publish_package(source, tmp_path):
    with pytest.raises(ValueError, match='变化'):
        package(source, tmp_path / 'out', overrides={'token': 'stale'})
    assert not (tmp_path / 'out' / 'share.zip').exists()
    with pytest.raises(ValueError, match='演示数据'):
        package(source, tmp_path / 'out', count=300)
    assert not (tmp_path / 'out' / 'share.zip').exists()
