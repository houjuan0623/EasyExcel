"""Electron / Python 接口的真实工作簿回归测试。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import threading

import openpyxl
from openpyxl.styles import PatternFill
import pytest

from core.collect import merge_tables
from core.splitter import list_values, run_split
from desktop import bridge


@pytest.fixture
def workbook(tmp_path):
    """三个可见 Sheet（一个结构不同），一个隐藏 Sheet，另有一个空白 Sheet。"""
    path = tmp_path / "多 Sheet 测试 ' 空格.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, department, amount in [("十月", "研发部", 100), ("十一月", "设计部", 200), ("内部", "隐藏部", 999)]:
        ws = wb.create_sheet(name)
        ws.append(["月度汇总"])
        ws.append(["部门", "姓名", "金额", "日期"])
        ws.append([department, "测试人员", amount, datetime.datetime(2026, 10, 1)])
        ws["A2"].fill = PatternFill("solid", fgColor="21845E")
    wb["内部"].sheet_state = "hidden"
    ws = wb.create_sheet("库存")
    ws.append(["物料", "库存"])
    ws.append(["物料甲", 7])
    wb.create_sheet("空白页")
    wb.save(path)
    wb.close()
    return path


def request(path, out, sheets=None, **overrides):
    req = {"inputPath": str(path), "outputPath": str(out), "mode": "split",
           "selection": [{"path": str(path), "sheets": sheets if sheets is not None else ["十一月"]}],
           "headerRow": None, "column": "部门", "values": None, "preserveFormat": True,
           "keepFormulas": False, "makeZip": False, "mergeAcrossFiles": False}
    req.update(overrides)
    return req


def test_inspect_includes_hidden_and_empty_sheets_without_selecting_them(workbook):
    result = bridge.inspect_source({"path": str(workbook)})
    sheets = {s["name"]: s for s in result["workbooks"][0]["sheets"]}
    assert set(sheets) == {"十月", "十一月", "内部", "库存", "空白页"}
    assert sheets["十一月"]["headerRow"] == 2
    assert sheets["十一月"]["columns"] == ["部门", "姓名", "金额", "日期"]
    assert sheets["内部"]["selectable"] is False
    assert sheets["空白页"]["selectable"] is False


def test_preview_uses_requested_sheet_and_original_row_numbers(workbook):
    result = bridge.preview_sheet({"path": str(workbook), "sheet": "十一月"})
    assert [row["index"] for row in result["rows"]] == [1, 2, 3]
    assert result["rows"][0]["cells"][0]["text"] == "月度汇总"
    assert [cell["text"] for cell in result["rows"][2]["cells"]] == ["设计部", "测试人员", "200", "2026-10-01 0:00:00"]
    with pytest.raises(ValueError, match="隐藏"):
        bridge.preview_sheet({"path": str(workbook), "sheet": "内部"})


def test_groups_and_export_share_selection_and_preserve_source(workbook, tmp_path):
    before = hashlib.sha256(workbook.read_bytes()).digest()
    req = request(workbook, tmp_path / "out")
    assert bridge.groups(req) == {"values": ["设计部"], "missing": []}
    events = []
    result = bridge.run_job(req, threading.Event(), events.append)
    assert result["summary"]["rows"] == 1
    assert result["summary"]["groups"] == 1
    files = list(Path(result["output"]).rglob("*.xlsx"))
    assert len(files) == 1
    wb = openpyxl.load_workbook(files[0])
    assert wb.sheetnames == ["十一月"]
    assert wb.active["A3"].value == "设计部"
    assert wb.active["A2"].fill.fgColor.rgb == "0021845E"
    wb.close()
    assert hashlib.sha256(workbook.read_bytes()).digest() == before
    assert any(e["type"] == "progress" for e in events)


def test_multiselect_merges_only_selected_sheets(workbook, tmp_path):
    req = request(workbook, tmp_path / "merged", ["十月", "十一月"], mode="merge")
    result = bridge.run_job(req, threading.Event(), lambda _: None)
    wb = openpyxl.load_workbook(result["output"], data_only=True)
    rows = [row for ws in wb if ws.title != "汇总说明" for row in ws.values]
    assert any("研发部" in row for row in rows)
    assert any("设计部" in row for row in rows)
    assert not any("隐藏部" in row or "物料甲" in row for row in rows)
    wb.close()


@pytest.mark.parametrize("sheets", [[], ["内部"], ["不存在"], ["空白页"]])
def test_invalid_or_empty_selection_never_becomes_all_sheets(workbook, tmp_path, sheets):
    with pytest.raises(ValueError):
        bridge.run_job(request(workbook, tmp_path / "out", sheets), threading.Event(), lambda _: None)
    assert not (tmp_path / "out").exists()


def test_missing_field_is_rejected_before_output(workbook, tmp_path):
    req = request(workbook, tmp_path / "out", ["十月", "库存"])
    assert bridge.groups(req)["missing"] == [f"{workbook.name} / 库存"]
    with pytest.raises(ValueError, match="缺少拆分字段"):
        bridge.run_job(req, threading.Event(), lambda _: None)
    assert not (tmp_path / "out").exists()


@pytest.mark.parametrize("values", [[], ["不存在的分组"], "设计部"])
def test_invalid_group_selection_is_rejected(workbook, tmp_path, values):
    with pytest.raises(ValueError, match="分组选择"):
        bridge.run_job(request(workbook, tmp_path / "out", values=values), threading.Event(), lambda _: None)


def test_old_core_callers_keep_default_all_visible_behavior(workbook, tmp_path):
    cfg = {"header_mode": "auto", "split_column": "部门", "input_path": str(workbook),
           "output_path": str(tmp_path / "legacy"), "make_zip": False}
    assert list_values(str(workbook), cfg, "部门") == ["研发部", "设计部"]
    cfg["selected_sheets"] = []
    assert list_values(str(workbook), cfg, "部门") == []
    cfg["selected_sheets"] = ["十一月"]
    result = run_split(cfg)
    wb = openpyxl.load_workbook(next(Path(result).glob("*.xlsx")))
    assert wb.sheetnames == ["十一月"]
    wb.close()


def test_folder_selection_is_per_workbook_and_does_not_reinclude_other_files(workbook, tmp_path):
    folder = tmp_path / "input"
    folder.mkdir()
    a, b = folder / "A.xlsx", folder / "B.xlsx"
    a.write_bytes(workbook.read_bytes())
    wb = openpyxl.load_workbook(workbook)
    wb["十一月"]["A3"] = "市场部"
    wb.save(b)
    wb.close()
    (folder / "C.xlsx").write_bytes(workbook.read_bytes())
    cfg = {"input_path": str(folder), "output_path": str(tmp_path / "out"), "header_mode": "auto",
           "split_column": "部门", "make_zip": False,
           "sheet_selection": {os.path.normcase(str(a)): ["十月"], os.path.normcase(str(b)): ["十一月"]}}
    result = run_split(cfg)
    rows = []
    for path in Path(result).rglob("*.xlsx"):
        out = openpyxl.load_workbook(path)
        rows.extend(row for ws in out for row in ws.values)
        out.close()
    assert sum("研发部" in row for row in rows) == 1
    assert sum("市场部" in row for row in rows) == 1
    assert not any("设计部" in row for row in rows)
    cfg["output_path"] = str(tmp_path / "merged")
    merged = openpyxl.load_workbook(merge_tables(cfg))
    values = [row for ws in merged if ws.title != "汇总说明" for row in ws.values]
    assert sum("研发部" in row for row in values) == 1
    assert sum("市场部" in row for row in values) == 1
    assert not any("设计部" in row for row in values)
    merged.close()


def test_cancellation_before_run_creates_no_results(workbook, tmp_path):
    cancelled = threading.Event()
    cancelled.set()
    result = bridge.run_job(request(workbook, tmp_path / "out"), cancelled, lambda _: None)
    assert result["stopped"] is True
    assert not (tmp_path / "out").exists()


def test_explicit_selection_does_not_lose_workbooks_with_result_like_names(workbook, tmp_path):
    folder = tmp_path / "input"
    folder.mkdir()
    selected_file = folder / "汇总_原始业务.xlsx"
    selected_file.write_bytes(workbook.read_bytes())
    req = request(selected_file, tmp_path / "merged", inputPath=str(folder), mode="merge")
    result = bridge.run_job(req, threading.Event(), lambda _: None)
    assert result["summary"]["rows"] == 1


def test_xls_sheet_names_and_dates_are_read_by_python():
    path = Path(__file__).parent / "golden/fixtures/monthly.xls"
    result = bridge.inspect_source({"path": str(path)})
    assert result["workbooks"][0]["sheets"][0]["name"] == "数据"
    preview = bridge.preview_sheet({"path": str(path), "sheet": "数据"})
    assert any(cell["text"].startswith("2026-10-01") for row in preview["rows"] for cell in row["cells"])


def test_protocol_accepts_unicode_paths_and_returns_structured_errors(workbook):
    script = Path(bridge.__file__)
    for request_data, expected in [({"command": "inspect", "path": str(workbook)}, "result"),
                                   ({"command": "inspect", "path": str(workbook.parent / "missing.xlsx")}, "error")]:
        process = subprocess.run([sys.executable, "-u", str(script)], input=json.dumps(request_data),
                                 capture_output=True, text=True, encoding="utf-8", timeout=20)
        event = json.loads(process.stdout)
        assert event["type"] == expected


def test_manual_header_and_output_guard(workbook, tmp_path):
    result = bridge.inspect_source({"path": str(workbook), "headerRow": 2})
    assert result["workbooks"][0]["sheets"][0]["headerRow"] == 2
    with pytest.raises(ValueError, match="1 到 1000"):
        bridge.header_config(1001)
    req = request(workbook, tmp_path, inputPath=str(tmp_path))
    with pytest.raises(ValueError, match="输入目录"):
        bridge.run_job(req, threading.Event(), lambda _: None)
