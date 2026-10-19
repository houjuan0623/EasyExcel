"""预览需要保留原始布局、显示格式，并严格限定原表前 100 行。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
import pytest

from desktop.bridge import preview_sheet
from desktop.preview_format import display


def at(preview, row, column):
    return next(cell for record in preview["rows"] if record["index"] == row for cell in record["cells"] if cell["column"] == column)


@pytest.fixture
def styled(tmp_path):
    path = tmp_path / "原样预览.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报表"
    ws.merge_cells("A1:F2")
    ws["A1"] = "十月销售报表"
    ws["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="21845E")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].border = Border(bottom=Side(style="medium", color="112233"))
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 24
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["D"].hidden = True
    ws.row_dimensions[7].hidden = True
    ws.append([])
    for c, text in enumerate(["编号", "日期", "金额", "内部信息", "比例", "备注"], 1):
        ws.cell(4, c, text)
    for r in range(5, 151):
        ws.cell(r, 1, r)
        ws.cell(r, 2, dt.datetime(2026, 10, 19))
        ws.cell(r, 2).number_format = 'yyyy"年"m"月"d"日"'
        ws.cell(r, 3, 1234.5).number_format = '"¥"#,##0.00'
        ws.cell(r, 4, "隐藏信息不传给界面")
        ws.cell(r, 5, 0.125).number_format = "0.0%"
    ws["A5"].number_format = "00000"
    ws["F5"] = "换行内容\n第二行"
    ws["F5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["F5"].font = Font(color=Color(theme=4, tint=0.4))
    ws["F5"].border = Border(left=Side(style="dashed", color="FF0000"))
    ws.row_dimensions[5].height = 42
    ws["A7"] = "隐藏行不传给界面"
    ws.merge_cells("A99:B105")
    ws["A99"] = "合并区域跨越预览边界"
    ws["F100"] = "第100行仍然显示"
    ws["F101"] = "第101行不进入预览"
    wb.save(path)
    wb.close()
    return path


def test_original_top_100_rows_styles_merges_and_dimensions(styled):
    before = hashlib.sha256(styled.read_bytes()).digest()
    preview = preview_sheet({"path": str(styled), "sheet": "报表"})
    assert preview["endRow"] == 100
    assert preview["rows"][-1]["index"] == 100
    assert preview["hiddenRows"] == 1 and preview["hiddenColumns"] == 1
    assert len(preview["rows"]) == 99
    assert preview["columns"][0] == {"index": 1, "label": "A", "width": 159}
    assert [c["label"] for c in preview["columns"]] == ["A", "B", "C", "E", "F"]
    assert preview["rows"][0]["height"] == 40
    title = at(preview, 1, 1)
    assert title["text"] == "十月销售报表"
    assert title["rowSpan"] == 2 and title["colSpan"] == 5
    assert preview["rows"][1]["cells"] == []
    style = preview["styles"][title["style"]]
    assert style["fontWeight"] == "bold" and style["fontSize"] == 26.67
    assert style["backgroundColor"] == "#21845E" and style["color"] == "#FFFFFF"
    assert style["textAlign"] == "center" and style["verticalAlign"] == "middle"
    assert style["borderBottom"] == "2px solid #112233"
    assert at(preview, 3, 1)["text"] == ""  # 空白行不删、不改为破折号
    assert at(preview, 5, 1)["text"] == "00005"
    assert at(preview, 5, 2)["text"] == "2026年10月19日"
    assert at(preview, 5, 3)["text"] == "¥1,234.50"
    assert at(preview, 5, 5)["text"] == "12.5%"
    note = at(preview, 5, 6)
    assert note["text"] == "换行内容\n第二行"
    assert preview["styles"][note["style"]]["wrapText"] is True
    assert re.fullmatch(r"#[0-9A-F]{6}", preview["styles"][note["style"]]["color"])
    assert at(preview, 99, 1)["rowSpan"] == 2
    assert at(preview, 100, 6)["text"] == "第100行仍然显示"
    payload = json.dumps(preview, ensure_ascii=False, allow_nan=False)
    for private in ("隐藏信息不传给界面", "隐藏行不传给界面", "第101行不进入预览"):
        assert private not in payload
    assert hashlib.sha256(styled.read_bytes()).digest() == before


def rewrite_sheet(path, transform):
    with ZipFile(path) as source:
        items = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name, content in items:
            target.writestr(name, transform(content.decode()).encode() if name == "xl/worksheets/sheet1.xml" else content)


def test_formula_uses_saved_cache_and_uncached_formula_is_explicit(tmp_path):
    path = tmp_path / "公式.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=1+2"
    ws["A1"].number_format = "0.00"
    ws["B1"] = "=1+3"
    wb.save(path)
    rewrite_sheet(path, lambda xml: xml.replace("<f>1+2</f><v></v>", "<f>1+2</f><v>3</v>"))
    preview = preview_sheet({"path": str(path), "sheet": "Sheet"})
    assert at(preview, 1, 1)["text"] == "3.00"
    assert at(preview, 1, 2)["text"] == "=1+3"
    assert any("缓存" in note for note in preview["warnings"])


def test_no_header_required_and_wrong_dimensions_do_not_truncate(styled):
    rewrite_sheet(styled, lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A1"', xml))
    preview = preview_sheet({"path": str(styled), "sheet": "报表", "headerRow": 130})
    assert preview["rows"][0]["index"] == 1
    assert at(preview, 100, 6)["text"] == "第100行仍然显示"


def test_hidden_column_groups_and_column_limit(tmp_path):
    path = tmp_path / "宽表.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions.group("B", "D", hidden=True)
    ws["A1"] = "标题"
    ws["B1"] = "隐藏组"
    ws.cell(1, 257, "超出列窗口")
    wb.save(path)
    preview = preview_sheet({"path": str(path), "sheet": "Sheet"})
    assert len(preview["columns"]) == 253
    assert preview["columns"][-1]["label"] == "IV"
    assert preview["truncatedColumns"] is True
    assert "隐藏组" not in json.dumps(preview, ensure_ascii=False)


def test_xls_layout_and_number_formats(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = tmp_path / "旧表.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("数据")
    ws.write_merge(0, 1, 0, 2, "原表标题", xlwt.easyxf('font: bold on, height 320; align: horiz center; pattern: pattern solid, fore_colour yellow;'))
    ws.col(0).width = 20 * 256
    ws.row(0).height = 600
    ws.row(0).height_mismatch = True
    for row in range(2, 105):
        ws.write(row, 0, row / 100, xlwt.easyxf(num_format_str="0.00%"))
    wb.save(str(path))
    preview = preview_sheet({"path": str(path), "sheet": "数据"})
    assert preview["endRow"] == 100 and len(preview["rows"]) == 100
    title = at(preview, 1, 1)
    assert title["rowSpan"] == 2 and title["colSpan"] == 3
    assert preview["styles"][title["style"]]["backgroundColor"] == "#FFFF00"
    assert preview["columns"][0]["width"] == 145
    assert at(preview, 3, 1)["text"] == "2.00%"


@pytest.mark.parametrize(("value", "fmt", "text"), [
    (1234.5, '#,##0.00', '1,234.50'), (0.125, '0.0%', '12.5%'),
    (12, '00000', '00012'), (-1234.5, '#,##0.00;[Red](#,##0.00)', '(1,234.50)'),
    (12.5, '0', '13'), (0, '0;-0;"—"', '—'), (1234, '0.00E+00', '1.23E+03'),
    (1.2, '0.00##', '1.20'), (1234000, '0.0,,"百万"', '1.2百万'),
    (dt.datetime(2026, 10, 19, 14, 5, 9), 'yyyy/mm/dd hh:mm:ss', '2026/10/19 14:05:09'),
    (dt.time(14, 5, 9), 'h:mm AM/PM', '2:05 PM'), (dt.timedelta(hours=49, minutes=3), '[h]:mm:ss', '49:03:00'),
    ('<script>alert(1)</script>', 'General', '<script>alert(1)</script>'),
])
def test_common_excel_display_formats(value, fmt, text):
    assert display(value, fmt) == (text, False)
