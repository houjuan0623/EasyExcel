"""统计图与复杂样式的预览回归，所有合成文件由 pytest 写入临时目录。

Copyright (c) 2026 AbeLin · MIT License
"""
import base64
from io import BytesIO
import json
from zipfile import ZipFile

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import AreaChart, BarChart, BarChart3D, DoughnutChart, LineChart, PieChart, RadarChart, Reference, ScatterChart, Series
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule
from openpyxl.styles import Border, Font, GradientFill, PatternFill, Side
from PIL import Image
import pytest

from desktop.preview import read_preview
from desktop.preview_drawings import relationships


def at(preview, row, column):
    """按 Excel 的行列坐标查找输出单元格。"""
    return next(cell for record in preview['rows'] if record['index'] == row for cell in record['cells'] if cell['column'] == column)


def style(preview, row, column):
    """解析输出中去重后的单元格样式。"""
    return preview['styles'][at(preview, row, column)['style']]


def save_preview(wb, tmp_path):
    """保存再读取，覆盖实际 OOXML 与只读解析链路。"""
    path = tmp_path / 'complex.xlsx'
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    result = read_preview(path, 'Sheet')
    assert before == path.read_bytes()
    json.dumps(result, allow_nan=False)
    return result


def test_gradients_patterns_diagonals_rich_text_and_negative_color(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = '渐变标题'
    ws['A1'].fill = GradientFill(stop=['156B50', 'DFF5E9'], degree=45)
    ws['A1'].border = Border(diagonal=Side(style='thin', color='FF0000'), diagonalUp=True, diagonalDown=True)
    ws['A2'] = '图案'
    ws['A2'].fill = PatternFill('darkGrid', fgColor='112233', bgColor='EFEFEF')
    ws['A3'] = CellRichText('普通', TextBlock(InlineFont(b=True, color='CC0000'), '重点'), TextBlock(InlineFont(vertAlign='superscript'), '2'))
    ws['B1'] = -1234.5
    ws['B1'].number_format = '#,##0.00;[Red](#,##0.00)'
    preview = save_preview(wb, tmp_path)
    assert style(preview, 1, 1)['backgroundImage'].startswith('linear-gradient(135.0deg')
    assert style(preview, 1, 1)['diagonalUp'] and style(preview, 1, 1)['diagonalDown']
    assert style(preview, 1, 1)['diagonalColor'] == '#FF0000'
    assert 'repeating-linear-gradient' in style(preview, 2, 1)['backgroundImage']
    assert style(preview, 2, 1)['backgroundColor'] == '#EFEFEF'
    runs = at(preview, 3, 1)['runs']
    assert ''.join(run['text'] for run in runs) == '普通重点2'
    assert runs[1]['style'] == {'fontWeight': 'bold', 'color': '#CC0000'}
    assert runs[2]['style']['verticalAlign'] == 'super'
    assert at(preview, 1, 2)['text'] == '(1,234.50)'
    assert style(preview, 1, 2)['color'] == '#FF0000'


def test_shared_string_rich_text_from_excel_style_workbooks(tmp_path):
    import xlsxwriter
    path = tmp_path / 'shared.xlsx'
    with xlsxwriter.Workbook(path) as book:
        ws = book.add_worksheet('Sheet')
        ws.write_rich_string('A1', '普通', book.add_format({'bold': True, 'font_color': '#CC0000'}), '加粗红色', '结尾')
    preview = read_preview(path, 'Sheet')
    assert at(preview, 1, 1)['text'] == '普通加粗红色结尾'
    assert ''.join(run['text'] for run in at(preview, 1, 1)['runs']) == '普通加粗红色结尾'
    assert at(preview, 1, 1)['runs'][1]['style']['color'] == '#CC0000'


def test_conditional_styles_priority_full_range_and_simple_formulas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 121):
        ws.cell(r, 1, r)
        ws.cell(r, 2, r)
        ws.cell(r, 6, r)
    for r, value in enumerate([-50, 0, 50], 1):
        ws.cell(r, 3, value)
        ws.cell(r, 4, value)
    ws['E1'] = 15
    ws['E1'].font = Font(bold=True, color='000000', size=16)
    ws['E1'].fill = PatternFill('solid', fgColor='FFFFFF')
    ws.conditional_formatting.add('E1', CellIsRule(operator='greaterThan', formula=['10'], stopIfTrue=True,
        font=Font(color='CC0000'), fill=PatternFill('solid', fgColor='FFFF00')))
    ws.conditional_formatting.add('E1', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='0000FF')))
    ws.conditional_formatting.add('B1:B120', ColorScaleRule(start_type='min', start_color='000000', end_type='max', end_color='FFFFFF'))
    ws.conditional_formatting.add('C1:C3', DataBarRule(start_type='min', end_type='max', color='638EC6'))
    ws.conditional_formatting.add('D1:D3', IconSetRule('3Arrows', 'num', [-50, 0, 50], showValue=False))
    ws['G1'] = 90
    ws.conditional_formatting.add('F1:F120', FormulaRule(formula=['A1>$G$1'], fill=PatternFill('solid', fgColor='FFDDDD')))
    # 最高优先级显式指定与原样式相同的白色，也应阻止后面的蓝色填充。
    ws['H1'] = 1
    ws.conditional_formatting.add('H1', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='FFFFFF')))
    ws.conditional_formatting.add('H1', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='0000FF')))
    ws.conditional_formatting.add('G1', FormulaRule(formula=['SUM(A1:A3)>0'], font=Font(bold=True)))
    for r, value in enumerate([0, 50, 100], 1):
        ws.cell(r, 9, value)
    ws.conditional_formatting.add('I1:I3', DataBarRule(start_type='min', end_type='max', color='638EC6', minLength=5, maxLength=95))
    preview = save_preview(wb, tmp_path)
    assert style(preview, 1, 5)['color'] == '#CC0000'
    assert style(preview, 1, 5)['fontWeight'] == 'bold'
    assert style(preview, 1, 5)['fontSize'] == 21.33
    assert style(preview, 1, 5)['backgroundColor'] == '#FFFF00'
    assert style(preview, 100, 2)['backgroundColor'] == '#D4D4D4'  # 阈值读取到了第 120 行
    assert at(preview, 1, 3)['dataBar'] == {'start': 0.0, 'width': 50.0, 'color': '#D65B5B'}
    assert at(preview, 3, 3)['dataBar']['start'] == 50
    assert at(preview, 3, 4)['icon']['text'] == '↑' and at(preview, 3, 4)['hideText']
    assert style(preview, 1, 6)['backgroundColor'] == '#FFFFFF'
    assert style(preview, 100, 6)['backgroundColor'] == '#FFDDDD'
    assert style(preview, 1, 8)['backgroundColor'] == '#FFFFFF'
    assert at(preview, 1, 9)['dataBar']['width'] == 5 and at(preview, 3, 9)['dataBar']['width'] == 95
    assert any('复杂条件格式' in warning for warning in preview['warnings'])
    assert preview['rows'][-1]['index'] == 100


def test_charts_read_beyond_100_and_exclude_hidden_rows(tmp_path, monkeypatch):
    from matplotlib.axes import Axes
    captured = []
    original = Axes.bar
    def bar(self, x, height, *args, **kwargs):
        captured.append(list(height))
        return original(self, x, height, *args, **kwargs)
    monkeypatch.setattr(Axes, 'bar', bar)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['月份', '金额'])
    for i in range(1, 121):
        ws.append([f'M{i}', i * 10])
    ws.row_dimensions[2].hidden = True
    chart = BarChart()
    chart.title = '完整范围统计'
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=121), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=121))
    ws.add_chart(chart, 'D110')
    preview = save_preview(wb, tmp_path)
    assert len(captured[0]) == 119 and captured[0][-1] == 1200 and 10 not in captured[0]
    asset = preview['drawings'][0]
    assert asset['kind'] == 'chart' and asset['inline'] is False and asset['title'] == '完整范围统计'
    with Image.open(BytesIO(base64.b64decode(asset['src'].split(',')[1]))) as image:
        assert image.size == (1000, 575)


@pytest.mark.parametrize('kind', [BarChart, LineChart, AreaChart, PieChart, DoughnutChart, RadarChart, ScatterChart])
def test_supported_chart_types_produce_png(tmp_path, kind):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [('类别', '金额'), (1, 20), (2, 30), (3, 10)]:
        ws.append(row)
    chart = kind()
    data = Reference(ws, min_col=2, min_row=1, max_row=4)
    categories = Reference(ws, min_col=1, min_row=2, max_row=4)
    if kind is ScatterChart:
        chart.series.append(Series(data, categories, title_from_data=True))
    else:
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
    ws.add_chart(chart, 'D2')
    preview = save_preview(wb, tmp_path)
    asset = preview['drawings'][0]
    assert asset['src'].startswith('data:image/png;base64,') and asset['inline'] is True
    assert asset['x'] > 0 and asset['y'] == 20


def test_combo_retains_secondary_axis_and_percent_scale(tmp_path, monkeypatch):
    from matplotlib.figure import Figure
    captured = []
    original = Figure.savefig
    def save(self, *args, **kwargs):
        captured.append([(ax.get_ylabel(), ax.get_ylim(), len(ax.lines), len(ax.patches)) for ax in self.axes])
        return original(self, *args, **kwargs)
    monkeypatch.setattr(Figure, 'savefig', save)
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in [('月份', '销售额', '完成率'), ('1月', 12000, .7), ('2月', 20000, .8), ('3月', 18000, .9)]:
        ws.append(row)
    bar, line = BarChart(), LineChart()
    bar.y_axis.title = '销售额'
    line.y_axis.title = '完成率'
    line.y_axis.axId = 200
    line.y_axis.numFmt = '0%'
    line.y_axis.scaling.min, line.y_axis.scaling.max = 0, 1
    for chart, col in [(bar, 2), (line, 3)]:
        chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    bar += line
    ws.add_chart(bar, 'E2')
    preview = save_preview(wb, tmp_path)
    assert 'src' in preview['drawings'][0]
    assert len(captured[0]) == 2
    assert captured[0][0][0] == '销售额' and captured[0][0][3] == 3
    assert captured[0][1] == ('完成率', (0.0, 1.0), 1, 0)


def test_images_unsupported_charts_and_hidden_sheet_refs_are_isolated(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = '单元格始终可用'
    private = wb.create_sheet('隐藏数据')
    private.append(['Secret', 123456])
    private.sheet_state = 'hidden'
    chart = BarChart()
    chart.add_data(Reference(private, min_col=2, min_row=1, max_row=1))
    ws.add_chart(chart, 'D2')
    unsupported = BarChart3D()
    unsupported.add_data(Reference(private, min_col=2, min_row=1, max_row=1))
    ws.add_chart(unsupported, 'D20')
    png = BytesIO()
    Image.new('RGB', (80, 40), '#21845E').save(png, format='PNG')
    png.seek(0)
    ws.add_image(ExcelImage(png), 'B4')
    preview = save_preview(wb, tmp_path)
    assert at(preview, 1, 1)['text'] == '单元格始终可用'
    assert len(preview['drawings']) == 3
    charts = [asset for asset in preview['drawings'] if asset['kind'] == 'chart']
    assert all('error' in asset and 'src' not in asset for asset in charts)
    assert '隐藏' in charts[0]['error'] and '暂不支持' in charts[1]['error']
    image = next(asset for asset in preview['drawings'] if asset['kind'] == 'image')
    assert image['src'].startswith('data:image/png;base64,') and image['width'] == 80
    assert '123456' not in json.dumps(preview)


def test_external_drawing_relationships_never_resolve(tmp_path):
    path = tmp_path / 'external.zip'
    with ZipFile(path, 'w') as archive:
        archive.writestr('xl/drawings/_rels/drawing1.xml.rels', '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="external" Target="https://example.invalid/private.png" TargetMode="External" />
          <Relationship Id="local" Target="../media/image1.png" />
        </Relationships>''')
        archive.writestr('xl/media/image1.png', b'local')
    with ZipFile(path) as archive:
        assert relationships(archive, 'xl/drawings/drawing1.xml') == {'local': 'xl/media/image1.png'}
