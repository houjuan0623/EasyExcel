"""多 Sheet 整理的真实文件回归；所有产物写入 pytest 临时目录。

Copyright (c) 2026 AbeLin · MIT License
"""
import hashlib
from copy import copy
import json
from pathlib import Path
import threading

import openpyxl
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

from core.sheet_bundle import bundle_workbooks
from core.splitter import run_split
from desktop.bridge import run_job


def make_book(path, sheets=('数据',), amount=123.45):
    """包含多行标题、格式、合并区域、隐藏列和超过预览上限的数据。"""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in sheets:
        ws = book.create_sheet(name)
        ws.append(['统计表'])
        ws.merge_cells('A1:D1')
        ws.append(['部门', '金额', '编号', '备注', '隐藏列'])
        for i in range(140):
            ws.append(['甲组' if i % 2 == 0 else '乙组', amount, f'{i:05d}', '=原样文字', '不带出'])
            ws.cell(i + 3, 4).data_type = 's'
        ws['B3'].number_format = '#,##0.00'
        ws['A2'].fill = PatternFill('solid', fgColor='21845E')
        ws['C3'].font = Font(color='FF0000', bold=True)
        ws['C3'].protection = Protection(locked=False)
        ws.row_dimensions[3].height = 35
        ws.row_dimensions[4].hidden = True
        ws.column_dimensions['C'].width = 24
        ws.column_dimensions['E'].hidden = True
        ws.freeze_panes = 'B3'
        ws.auto_filter.ref = 'A2:D142'
        ws.print_area = 'A1:D142'
        ws.print_title_rows = '1:2'
        validation = DataValidation(type='list', formula1='"是,否"')
        validation.add('D3:D142')
        ws.add_data_validation(validation)
    book.save(path)
    book.close()
    return path


def split_config(path, out, **kwargs):
    return {'input_path': str(path), 'output_path': str(out), 'header_mode': 'row', 'header_row': 2,
            'split_column': '部门', 'preserve_format': True, 'make_zip': False, **kwargs}


def read_summary(logs):
    return json.loads(next(line[len('[SUMMARY] '):] for line in reversed(logs) if line.startswith('[SUMMARY] ')))


def test_split_also_bundles_each_result_and_keeps_original_files(tmp_path):
    source = make_book(tmp_path / '原始.xlsx', ('十月', '十一月'))
    before = hashlib.sha256(source.read_bytes()).digest()
    logs, progress = [], []
    folder = Path(run_split(split_config(source, tmp_path / 'out', bundle_sheets=True), logs.append, progress.append))
    result = read_summary(logs)
    assert result['files'] == 2
    assert result['rows'] == 280
    assert result['bundle_sheets'] == 4
    assert sorted(progress) == progress
    bundled = openpyxl.load_workbook(result['bundle_file'])
    assert set(bundled.sheetnames) == {'甲组_十月', '甲组_十一月', '乙组_十月', '乙组_十一月'}
    assert sum(ws.max_row - 2 for ws in bundled) == 280
    for path in folder.glob('*.xlsx'):
        original = openpyxl.load_workbook(path)
        for ws in original:
            dst = bundled[f'{path.stem}_{ws.title}']
            assert list(ws.values) == list(dst.values)
            assert copy(ws['A2'].fill) == copy(dst['A2'].fill)
            assert str(ws.merged_cells) == str(dst.merged_cells)
        original.close()
    bundled.close()
    assert hashlib.sha256(source.read_bytes()).digest() == before


def test_bundle_preserves_layout_text_styles_and_all_rows(tmp_path):
    source = make_book(tmp_path / '甲组.xlsx', ('数据', '归档'))
    wb = openpyxl.load_workbook(source)
    wb['归档'].sheet_state = 'hidden'
    wb.save(source)
    wb.close()
    result = bundle_workbooks([(source, None)], tmp_path / '合集.xlsx', {'header_mode': 'row', 'header_row': 2})
    assert result['files'] == 1 and result['sheets'] == 1 and result['skipped_sheets'] == 1
    wb = openpyxl.load_workbook(result['output'])
    ws = wb['甲组']
    assert ws.max_row == 142
    assert ws['B142'].value == 123.45
    assert ws['D3'].value == '=原样文字' and ws['D3'].data_type == 's'
    assert ws['C3'].value == '00000' and not ws['C3'].protection.locked
    assert ws['B3'].number_format == '#,##0.00'
    assert ws['A2'].fill.fgColor.rgb == '0021845E'
    assert ws['C3'].font.color.rgb == '00FF0000'
    assert str(ws.merged_cells) == 'A1:D1'
    assert ws.row_dimensions[3].height == 35 and ws.row_dimensions[4].hidden
    assert ws.column_dimensions['C'].width == 24 and ws.column_dimensions['E'].hidden
    assert ws['E3'].value is None
    assert ws.freeze_panes == 'B3' and ws.auto_filter.ref == 'A2:D142'
    assert 'A1:$D$142' in str(ws.print_area) or '$A$1:$D$142' in str(ws.print_area)
    assert ws.print_title_rows == '$1:$2'
    assert ws.data_validations.count == 1
    wb.close()


def test_repeated_names_long_names_and_reserved_titles_do_not_overwrite(tmp_path):
    paths = []
    for index, name in enumerate(['同名', '同名', 'A' * 35, 'A' * 34 + 'B', 'History']):
        folder = tmp_path / str(index)
        folder.mkdir()
        paths.append(make_book(folder / (name + '.xlsx'), amount=index + 1))
    result = bundle_workbooks([(p, ['数据']) for p in paths], tmp_path / '合集.xlsx')
    wb = openpyxl.load_workbook(result['output'])
    assert len(wb.sheetnames) == len(set(n.lower() for n in wb.sheetnames)) == 5
    assert all(len(n) <= 31 and n.lower() != 'history' for n in wb.sheetnames)
    assert [ws['B3'].value for ws in wb] == [1, 2, 3, 4, 5]
    wb.close()


def test_failure_and_stop_do_not_leave_a_misleading_total_or_overwrite_source(tmp_path):
    source = make_book(tmp_path / '原表.xlsx')
    content = source.read_bytes()
    with pytest.raises(FileExistsError):
        bundle_workbooks([(source, None)], source)
    assert source.read_bytes() == content
    cancelled = threading.Event()
    result = bundle_workbooks([(source, None)], tmp_path / '已停止.xlsx',
                              progress_fn=lambda _: cancelled.set(), stop_flag=cancelled.is_set)
    assert result['stopped'] and result['files'] == 0
    assert not (tmp_path / '已停止.xlsx').exists()
    assert not list(tmp_path.glob('*.tmp'))
    bad = tmp_path / '坏文件.xlsx'
    bad.write_bytes(b'not an excel file')
    result = bundle_workbooks([(bad, None), (source, ['数据'])], tmp_path / '部分.xlsx')
    assert result['failed_files'] == 1 and result['files'] == 1 and result['sheets'] == 1


def test_split_bundle_keeps_safe_formulas_and_excludes_review_workbook(tmp_path):
    source = make_book(tmp_path / '原始.xlsx')
    wb = openpyxl.load_workbook(source)
    wb.active['D3'] = '=B3*2'
    wb.active['A4'] = None
    wb.save(source)
    wb.close()
    logs = []
    run_split(split_config(source, tmp_path / 'out', bundle_sheets=True, keep_formulas=True), logs.append)
    result = read_summary(logs)
    assert result['unassigned'] == 1 and result['unassigned_file']
    wb = openpyxl.load_workbook(result['bundle_file'])
    assert set(wb.sheetnames) == {'甲组', '乙组'}
    assert wb['甲组']['D3'].value == '=B3*2'
    assert sum(ws.max_row - 2 for ws in wb) == 139
    wb.close()


def test_desktop_sheet_mode_keeps_different_structures_and_exact_selection(tmp_path):
    source = make_book(tmp_path / '多表.xlsx', ('数据', '其他', '不选'))
    wb = openpyxl.load_workbook(source)
    wb['其他'].delete_rows(1, 200)
    wb['其他'].append(['物料', '库存'])
    wb['其他'].append(['示例', 7])
    wb.save(source)
    wb.close()
    request = {'inputPath': str(source), 'outputPath': str(tmp_path / 'out'), 'mode': 'merge',
               'selection': [{'path': str(source), 'sheets': ['数据', '其他']}], 'headerRow': None,
               'preserveFormat': True, 'keepFormulas': False, 'makeZip': False, 'mergeAcrossFiles': False,
               'mergeMode': 'sheets', 'bundleSheets': False}
    result = run_job(request, threading.Event(), lambda _: None)
    assert result['summary']['sheets'] == 2
    wb = openpyxl.load_workbook(result['output'])
    assert wb.sheetnames == ['多表_数据', '多表_其他']
    assert wb['多表_其他']['B2'].value == 7
    wb.close()
    again = run_job(request, threading.Event(), lambda _: None)
    assert again['output'] != result['output'] and Path(result['output']).exists()
    with pytest.raises(ValueError, match='总工作簿设置'):
        run_job({**request, 'bundleSheets': 'false'}, threading.Event(), lambda _: None)
    with pytest.raises(ValueError, match='汇总方式'):
        run_job({**request, 'mergeMode': 'invalid'}, threading.Event(), lambda _: None)


def test_split_default_does_not_create_bundle(tmp_path):
    source = make_book(tmp_path / '原始.xlsx')
    logs = []
    output = Path(run_split(split_config(source, tmp_path / 'out'), logs.append))
    assert 'bundle_file' not in read_summary(logs)
    assert len(list(output.rglob('*.xlsx'))) == 2


def test_old_xls_can_be_bundled_with_named_sheet(tmp_path):
    source = Path(__file__).parent / 'golden' / 'fixtures' / 'monthly.xls'
    result = bundle_workbooks([(source, ['数据'])], tmp_path / '旧表合集.xlsx')
    assert result['files'] == 1 and result['sheets'] == 1
    wb = openpyxl.load_workbook(result['output'])
    assert wb.sheetnames == ['monthly'] and wb.active.max_row > 1
    wb.close()


def test_batch_split_bundle_does_not_repeat_zip_copies(tmp_path):
    folder = tmp_path / '输入'
    folder.mkdir()
    make_book(folder / '一月.xlsx')
    make_book(folder / '二月.xlsx')
    logs = []
    run_split(split_config(folder, tmp_path / 'out', bundle_sheets=True, make_zip=True,
                           merge_across_files=False), logs.append)
    summary = read_summary(logs)
    assert summary['files'] == 4 and summary['bundle_sheets'] == 4 and summary['zips'] == 2
    wb = openpyxl.load_workbook(summary['bundle_file'])
    assert sum(ws.max_row - 2 for ws in wb) == 280
    wb.close()
