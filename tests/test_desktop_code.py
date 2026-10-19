"""Real offline-worker tests, with outputs confined to pytest's external scratch.

Copyright (c) 2026 AbeLin · MIT License
"""
import base64
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

import openpyxl
import pytest

from desktop.bridge import selection_config, source_files


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "代码测试 ' 数据.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws.append(["原始标题"])
    ws.append(["部门", "金额", "隐藏字段"])
    for i in range(125):
        ws.append([["研发/一", "研发:一", "CON", None][i % 4], i + 1, "不载入"])
    ws.column_dimensions["C"].hidden = True
    ws.row_dimensions[5].hidden = True
    other = wb.create_sheet("补充")
    other.append(["标题"])
    other.append(["部门", "金额"])
    other.append(["销售", 1000])
    secret = wb.create_sheet("隐藏表")
    secret.append(["不应读取"])
    secret.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return path


def request(path, tmp_path, code, sheets=None):
    return {"command": "code", "inputPath": str(path), "outputPath": str(tmp_path / "输出"),
            "selection": [{"path": str(path), "sheets": ["明细"] if sheets is None else sheets}],
            "headerRow": 2, "code": code}


def worker(req):
    proc = subprocess.run([sys.executable, "-u", str(Path(__file__).resolve().parents[1] / "desktop/bridge.py")],
                          input=json.dumps(req, ensure_ascii=False) + "\n", capture_output=True,
                          text=True, encoding="utf-8", timeout=60,
                          env={**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"})
    assert proc.returncode == 0, proc.stdout + proc.stderr
    events = [json.loads(line) for line in proc.stdout.splitlines()]
    assert events[-1]["type"] == "result", events
    return events[-1]["result"], events


def test_full_selected_data_analysis_plot_split_merge_and_source_preserved(workbook, tmp_path):
    original = hashlib.sha256(workbook.read_bytes()).digest()
    code = '''
assert df.shape == (125, 2)
assert len(tables) == 2
assert "隐藏字段" not in df.columns
assert sum(df["金额"]) == sum(range(1, 126))
assert inputs[0]["headerRow"] == 2
assert inputs[0]["columns"] == ["部门", "金额"]
print("中文日志\\n第二行日志")
display(df, "完整数据")
display({"合计": int(df["金额"].sum())}, "分析")
merged = merge_tables()
assert len(merged) == 126
save_excel(merged, "合并.xlsx")
save_excel(df, "原数据.xlsx")
save_excel(df.head(1), "原数据.xlsx")
paths = split_table(df, "部门")
assert len(paths) == 4
assert len(set(paths)) == 4
fig, ax = plt.subplots()
ax.plot(df["金额"])
ax.set_title("数据趋势")
plt.show()
save_excel({"明细": df.head(2), "补充": merged.tail(1)}, "多Sheet.xlsx")
'''
    result, events = worker(request(workbook, tmp_path, code, ["明细", "补充"]))
    assert result["status"] == "success", result.get("error")
    assert result["rows"] == 126
    directory = Path(result["output"])
    assert (directory / "script.py").read_text(encoding="utf-8") == code
    assert "第二行日志" in (directory / "运行日志.txt").read_text(encoding="utf-8")
    assert "pandas" in json.loads((directory / "运行环境.json").read_text(encoding="utf-8"))["libraries"]
    tables = [e["data"] for e in events if e["type"] == "output" and e["data"]["kind"] == "table"]
    assert tables[0]["totalRows"] == 125
    assert len(tables[0]["rows"]) == 100
    charts = [e["data"] for e in events if e["type"] == "output" and e["data"]["kind"] == "image"]
    assert len(charts) == 1
    assert base64.b64decode(charts[0]["src"].split(",", 1)[1]).startswith(b"\x89PNG\r\n\x1a\n")
    merged = openpyxl.load_workbook(directory / "合并.xlsx")
    assert merged.active.max_row == 127
    assert merged.active.cell(1, 3).value == "来源"
    merged.close()
    counts = []
    for path in directory.glob("拆分_*.xlsx"):
        wb = openpyxl.load_workbook(path)
        counts.append(wb.active.max_row - 1)
        wb.close()
    assert sum(counts) == 125
    assert (directory / "原数据_2.xlsx").exists()
    assert hashlib.sha256(workbook.read_bytes()).digest() == original


@pytest.mark.parametrize("code,detail", [
    ('print("先完成的输出")\nraise ValueError("计算口径不明确")', 'ValueError: 计算口径不明确'),
    ('if True\n    pass', 'SyntaxError'),
    ('input("请输入")', 'EOFError'),
    ('raise SystemExit(3)', 'SystemExit: 3'),
])
def test_script_errors_return_tracebacks_and_keep_run_files(workbook, tmp_path, code, detail):
    result, events = worker(request(workbook, tmp_path, code))
    assert result["status"] == "error"
    assert detail in result["error"]
    assert "script.py" in result["error"]
    assert (Path(result["output"]) / "script.py").exists()
    assert any(e["type"] == "log" and detail in e["message"] for e in events)


def test_output_limits_and_collision_helpers(workbook, tmp_path):
    code = '''
print("x" * 250000)
display(pd.DataFrame(np.ones((120, 45))), "截断预览")
save_excel(df.head(1), "../../CON.xlsx")
save_excel({"A/B": df.head(1), "A:B": df.head(2)}, "工作簿.xlsx")
'''
    result, events = worker(request(workbook, tmp_path, code))
    assert result["status"] == "success", result.get("error")
    table = next(e["data"] for e in events if e["type"] == "output")
    assert len(table["rows"]) == 100
    assert len(table["columns"]) == 40
    assert table["totalColumns"] == 45
    assert sum(len(e.get("message", "")) for e in events if e["type"] == "log") < 202000
    wb = openpyxl.load_workbook(Path(result["output"]) / "工作簿.xlsx")
    assert wb.sheetnames == ["A_B", "A_B_2"]
    wb.close()


@pytest.mark.parametrize("sheets", [[], ["隐藏表"], ["不存在"]])
def test_bad_selections_fail_before_user_code(workbook, tmp_path, sheets):
    req = request(workbook, tmp_path, "raise AssertionError('不能运行')", sheets)
    with pytest.raises(ValueError):
        selection_config(req)
    assert not (tmp_path / "输出").exists()


def test_code_outputs_are_not_reimported_as_folder_inputs(workbook, tmp_path):
    output = tmp_path / "Python_20261019_100000_abcdef"
    output.mkdir()
    (output / "结果.xlsx").write_bytes(workbook.read_bytes())
    _, files = source_files(str(tmp_path))
    assert files == [workbook]


def test_offline_finance_libraries_do_real_calculations_and_exports(workbook, tmp_path):
    script = Path(__file__).resolve().parents[1] / "desktop/scripts/finance_checks.py"
    result, events = worker(request(workbook, tmp_path, script.read_text(encoding="utf-8")))
    assert result["status"] == "success", result.get("error")
    assert any("FINANCE_CHECKS_OK" in e.get("message", "") for e in events)
    directory = Path(result["output"])
    wb = openpyxl.load_workbook(directory / "财务库报表.xlsx")
    assert wb.active.max_row == 13
    assert len(wb.active._charts) == 1
    assert wb.active["B2"].number_format == "#,##0.00"
    wb.close()
    assert (directory / "财务分析图.png").read_bytes().startswith(b"\x89PNG")
    env = json.loads((directory / "运行环境.json").read_text(encoding="utf-8"))
    for name in ["numpy-financial", "pyxirr", "scipy", "statsmodels", "seaborn", "XlsxWriter", "Babel", "QuantLib", "decimal"]:
        assert name in env["libraries"]
