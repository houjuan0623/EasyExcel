"""统计口径、跨表关联和完整区域读取的独立数值回归。

Copyright (c) 2026 AbeLin · MIT License
"""

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path

import openpyxl
import pytest

from desktop import analysis


@pytest.fixture
def book(tmp_path):
    path = tmp_path / "analysis-fixture.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明细"
    ws["A1"] = "多层表头"
    ws.merge_cells("A1:G1")
    for cell, value in {
        "A3": "编号",
        "B3": "部门",
        "C3": "金额",
        "C4": "预算",
        "D4": "实际",
        "E3": "日期",
        "F3": "备注",
        "G3": "公式",
    }.items():
        ws[cell] = value
    for span in ("A3:A4", "B3:B4", "C3:D3", "E3:E4", "F3:F4", "G3:G4"):
        ws.merge_cells(span)
    data = [
        ["A01", "东部", 10, 12, datetime(2026, 1, 3), None, "=C5*2"],
        ["A02", "西部", 20, None, datetime(2026, 1, 9), "text", None],
        ["A03", "东部", 30, 36, datetime(2026, 2, 3), None, None],
        ["A04", "东部", None, 48, datetime(2026, 2, 7), None, None],
        ["A05", "西部", "50", 60, datetime(2026, 2, 9), None, None],
        ["合计", None, 9999, 9999, None, None, None],
        ["HIDDEN", "隐藏", 5000, 5000, None, None, None],
    ]
    for r, row in enumerate(data, 5):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    ws.row_dimensions[11].hidden = True
    ws["A121"] = "AFTER_100"
    ws["C121"] = 777
    ws["D121"] = 888
    other = wb.create_sheet("关联表")
    for row in [["编号", "利润"], ["A03", 300], ["A01", 100], ["A05", 500], ["Z99", 999]]:
        other.append(row)
    transpose = wb.create_sheet("横向数据")
    for row in [["项目", "a", "b", "c"], ["预算", 10, 20, 30], ["实际", 12, 24, 36]]:
        transpose.append(row)
    appended = wb.create_sheet("追加表")
    appended.append(["金额 / 实际", "编号", "部门", "金额 / 预算", "日期", "备注", "公式"])
    appended.append([6, "A06", "东部", 5, datetime(2026, 3, 1), None, None])
    hidden = wb.create_sheet("不分析的隐藏表")
    hidden.append(["value", 999])
    hidden.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return path


def block(path, identifier="a", sheet="明细", area="A3:G9", headers=2, **changes):
    return {
        "id": identifier,
        "name": identifier,
        "path": str(path),
        "sheet": sheet,
        "range": area,
        "headerRows": headers,
        "orientation": "columns",
        "skipTotals": True,
        "includeHidden": False,
        **changes,
    }


def request(path, kind="bar", blocks=None, **changes):
    specs = blocks or [block(path)]
    info = analysis.inspect({"inputPath": str(path), "blocks": specs})
    return {
        "inputPath": str(path),
        "blocks": specs,
        "token": info["token"],
        "kind": kind,
        "title": "测试分析",
        "x": "a:f1",
        "y": ["a:f2"],
        "combine": "single",
        "joinKeys": {},
        "joinType": "inner",
        "aggregation": "sum",
        "sort": "source",
        "timeBucket": "none",
        "bins": 5,
        "limit": 20,
        "unit": "元",
        "note": "测试口径",
        **changes,
    }


def test_region_supports_rows_columns_and_full_workbook_extent(book):
    meta = {"rows": 121, "columns": 7}
    assert analysis.region("$C:$D", meta) == (3, 1, 4, 121)
    assert analysis.region("5:9", meta) == (1, 5, 7, 9)
    assert analysis.region("B5:D121", meta) == (2, 5, 4, 121)
    data = analysis.grid(
        {"inputPath": str(book), "path": str(book), "sheet": "明细", "startRow": 120, "startColumn": 1, "headerRow": 4}
    )
    assert data["rows"][1][0] == "AFTER_100"
    assert data["totalRows"] == 121
    assert data["suggestedRange"] == "A3:G121" and data["suggestedHeaderRows"] == 2
    r = analysis.render(request(book, "metric", [block(book, area="C:C", headers=4)], y=["a:f0"]))
    # 整列明确包含合计值，因只选 C 列没有合计标签，清理不猜测最后一行。
    assert r["metric"] == 10 + 20 + 30 + 50 + 9999 + 777


def test_merged_headers_and_missing_formula_cache_are_reported(book):
    r = analysis.inspect({"inputPath": str(book), "blocks": [block(book)]})["blocks"][0]
    assert [f["name"] for f in r["fields"]] == ["编号", "部门", "金额 / 预算", "金额 / 实际", "日期", "备注", "公式"]
    assert r["rows"] == 5 and r["audit"]["missingFormulaCaches"] == 1
    assert r["sample"][0][-1] is None
    assert any("缺少已保存结果" in w for w in r["warnings"])


def test_hidden_and_totals_are_explicit_and_not_double_counted(book):
    r = analysis.render(request(book, "metric", [block(book, area="A3:G11")]))
    assert r["metric"] == 110 and r["counts"]["inputRows"] == 5
    assert any("排除 1 条" in w for w in r["warnings"])
    include = analysis.render(request(book, "metric", [block(book, area="A3:G11", includeHidden=True)]))
    assert include["metric"] == 5110
    grid = analysis.grid({"inputPath": str(book), "path": str(book), "sheet": "明细", "startRow": 1, "startColumn": 1})
    assert grid["rows"][10][0] == ""
    grid = analysis.grid(
        {
            "inputPath": str(book),
            "path": str(book),
            "sheet": "明细",
            "startRow": 1,
            "startColumn": 1,
            "includeHidden": True,
        }
    )
    assert grid["rows"][10][0] == "HIDDEN"


def test_statistics_and_bar_aggregation_match_hand_calculation(book):
    result = analysis.render(request(book))
    assert result["table"]["rows"] == [["东部", 40.0], ["西部", 70.0]]
    stats = result["statistics"]["rows"][0]
    assert stats[1:6] == [4, 1, 0, 110.0, 27.5]
    assert stats[6:9] == [25.0, 10.0, 50.0]
    assert stats[9] == pytest.approx(17.0782512766)
    assert stats[10:] == [17.5, 35.0]
    assert result["image"].startswith("data:image/png;base64,")
    json.dumps(result, allow_nan=False)


@pytest.mark.parametrize("kind", ["metric", "bar", "line", "pie", "histogram", "scatter", "box", "correlation"])
def test_all_chart_types_generate_with_verified_numeric_results(book, kind):
    spec = block(book, sheet="横向数据", area="A1:D3", headers=1, orientation="rows")
    req = request(book, kind, [spec], x="a:f0", y=["a:f1"])
    if kind in ("box", "correlation"):
        req["y"] = ["a:f1", "a:f2"]
    if kind == "scatter":
        req.update(x="a:f1", y=["a:f2"])
    result = analysis.render(req)
    assert result["counts"]["inputRows"] == 3
    if kind == "metric":
        assert result["metric"] == 60
    elif kind == "histogram":
        assert sum(row[-1] for row in result["table"]["rows"]) == 3
    elif kind == "correlation":
        assert result["table"]["rows"][1][2] == pytest.approx(1)
        assert result["table"]["rows"][1][3] == 3
    elif kind == "scatter":
        assert result["table"]["rows"][0] == [3, 0, pytest.approx(1)]
    else:
        assert result["image"].startswith("data:image/png;base64,")
    json.dumps(result, allow_nan=False)


def test_append_aligns_by_header_names_even_when_columns_are_reordered(book):
    req = request(book, "metric", [block(book), block(book, "b", "追加表", "A1:G2", 1)], combine="append")
    result = analysis.render(req)
    assert result["metric"] == 115 and result["counts"]["inputRows"] == 6
    req["blocks"][1]["headerRows"] = 0
    info = analysis.inspect(req)
    req["token"] = info["token"]
    with pytest.raises(ValueError, match="字段名必须一致"):
        analysis.render(req)


def test_join_uses_unique_keys_not_row_position_and_reports_unmatched(book):
    specs = [block(book), block(book, "b", "关联表", "A1:B5", 1)]
    req = request(book, "scatter", specs, combine="join", joinKeys={"a": "a:f0", "b": "b:f0"}, x="a:f2", y=["b:f1"])
    result = analysis.render(req)
    assert result["counts"]["inputRows"] == 3
    assert result["table"]["rows"][0] == [3, 0, pytest.approx(1)]
    assert any("未匹配 2 个键" in w and "独有 1 个键" in w for w in result["warnings"])
    req.update(joinType="left")
    result = analysis.render(req)
    assert result["counts"]["inputRows"] == 5
    assert result["table"]["rows"][0][:2] == [3, 2]


def test_duplicate_join_keys_block_cartesian_expansion(book):
    wb = openpyxl.load_workbook(book)
    wb["关联表"]["A5"] = "A01"
    wb.save(book)
    wb.close()
    req = request(
        book,
        "metric",
        [block(book), block(book, "b", "关联表", "A1:B5", 1)],
        combine="join",
        joinKeys={"a": "a:f0", "b": "b:f0"},
    )
    with pytest.raises(ValueError, match="重复值"):
        analysis.render(req)


def test_numeric_and_text_join_keys_are_not_silently_conflated(book):
    wb = openpyxl.load_workbook(book)
    wb["明细"]["A5"] = "001"
    wb["关联表"]["A3"] = 1
    wb.save(book)
    wb.close()
    req = request(
        book,
        "metric",
        [block(book), block(book, "b", "关联表", "A1:B5", 1)],
        combine="join",
        joinKeys={"a": "a:f0", "b": "b:f0"},
    )
    assert analysis.render(req)["metric"] == 80


def test_date_bucketing_groups_all_rows_and_preserves_missing_values(book):
    req = request(book, "line", x="a:f4", timeBucket="month", y=["a:f2", "a:f3"])
    result = analysis.render(req)
    assert result["table"]["rows"] == [["2026-01", 30.0, 12.0], ["2026-02", 80.0, 144.0]]
    req.update(x="a:f2")
    with pytest.raises(ValueError, match="普通数字"):
        analysis.render(req)


def test_invalid_pie_inputs_and_truncation_are_rejected(book):
    req = request(book, "pie", aggregation="mean")
    with pytest.raises(ValueError, match="仅支持"):
        analysis.render(req)
    req.update(aggregation="sum", limit=1)
    with pytest.raises(ValueError, match="分类数超过"):
        analysis.render(req)
    wb = openpyxl.load_workbook(book)
    wb["明细"]["C5"] = -100
    wb.save(book)
    wb.close()
    req = request(book, "pie")
    with pytest.raises(ValueError, match="非负"):
        analysis.render(req)


def test_stale_files_and_changed_headers_require_review(book):
    req = request(book, "metric")
    wb = openpyxl.load_workbook(book)
    wb["明细"]["C5"] = 200
    wb.save(book)
    wb.close()
    with pytest.raises(ValueError, match="已改变"):
        analysis.render(req)
    req = request(book, "metric", fieldLabels={"a:f2": "旧字段"})
    with pytest.raises(ValueError, match="字段已改变"):
        analysis.render(req)


@pytest.mark.parametrize(
    "area", ["A0:B2", "D5:A1", "1:100001", "A:XFD", "A1:B2;D1:E2", "A1:ZZ100000", "../A1", "A1048576"]
)
def test_unsafe_or_unbounded_ranges_fail(book, area):
    with pytest.raises(ValueError):
        analysis.inspect({"inputPath": str(book), "blocks": [block(book, area=area)]})


def test_foreign_files_and_hidden_sheets_are_not_read(book, tmp_path):
    other = tmp_path / "foreign.xlsx"
    other.write_bytes(book.read_bytes())
    with pytest.raises(ValueError, match="不属于"):
        analysis.inspect({"inputPath": str(book), "blocks": [block(other)]})
    with pytest.raises(ValueError, match="可见"):
        analysis.inspect({"inputPath": str(book), "blocks": [block(book, sheet="不分析的隐藏表")]})


def test_local_report_is_self_contained_and_escapes_html_and_csv(book, tmp_path):
    original = sha256(book.read_bytes()).hexdigest()
    req = request(book, "metric", title="<script>bad()</script>", note="=danger")
    plan = {"version": 1, "title": "<report>", "cards": [{"id": "card", "config": req}], "outputPath": str(tmp_path)}
    result = analysis.export_dashboard(plan)
    folder = Path(result["output"])
    html = (folder / "dashboard.html").read_text(encoding="utf-8")
    assert "&lt;script&gt;" in html and "<script>" not in html
    assert "default-src" in html and "http://" not in html and "https://" not in html
    assert (folder / "chart_01_statistics.csv").exists()
    assert json.loads((folder / "dashboard-plan.json").read_text(encoding="utf-8"))["version"] == 2
    assert analysis.csv_text({"columns": ["=cmd"], "rows": [["@evil"]]}) == "\ufeff'=cmd\r\n'@evil\r\n"
    assert sha256(book.read_bytes()).hexdigest() == original
    stale = {**req, "token": "changed"}
    before = set(tmp_path.iterdir())
    with pytest.raises(ValueError):
        analysis.export_dashboard({**plan, "cards": [{"id": "x", "config": stale}]})
    assert set(tmp_path.iterdir()) == before


def test_dashboard_export_preserves_custom_layout_and_layer_order(book, tmp_path):
    req = request(book, "metric")
    layout = {"x": 85, "y": 145, "width": 700, "height": 350}
    plan = {"version": 2, "canvas": {"width": 1920, "height": 1080}, "title": "可布局看板",
            "cards": [{"id": "metric", "config": req, "layout": layout},
                      {"id": "above", "config": req, "layout": {**layout, "x": 200}}], "outputPath": str(tmp_path)}
    output = Path(analysis.export_dashboard(plan)["output"])
    stored = json.loads((output / "dashboard-plan.json").read_text(encoding="utf-8"))
    assert stored["cards"] == plan["cards"]
    assert stored["canvas"] == plan["canvas"]
    document = (output / "dashboard.html").read_text(encoding="utf-8")
    assert 'viewBox="0 0 1920 1080"' in document
    assert 'left:85px;top:145px;width:700px;height:350px;z-index:1' in document
    assert document.index('data-card-id="metric"') < document.index('data-card-id="above"')
    assert '<script' not in document and 'object-fit: contain' in document


@pytest.mark.parametrize("layout", [None, {}, {"x": True, "y": 120, "width": 600, "height": 400},
                                   {"x": 24, "y": 120, "width": 1900, "height": 400},
                                   {"x": "24px", "y": 120, "width": 600, "height": 400},
                                   {"x": 24, "y": 0, "width": 600, "height": 400}])
def test_invalid_dashboard_layout_fails_without_creating_report(book, tmp_path, layout):
    before = set(tmp_path.iterdir())
    with pytest.raises(ValueError):
        analysis.export_dashboard({"version": 2, "canvas": {"width": 1920, "height": 1080}, "title": "布局",
                                   "cards": [{"id": "test", "config": {}, "layout": layout}], "outputPath": str(tmp_path)})
    assert set(tmp_path.iterdir()) == before


def test_legacy_dashboard_arrangement_matches_canvas_contract():
    from desktop.dashboard import dashboard_cards
    for count in range(1, 13):
        cards = dashboard_cards({"version": 1, "cards": [{"id": str(i), "config": {}} for i in range(count)]})
        for card in cards:
            r = card["layout"]
            assert r["x"] >= 24 and r["y"] >= 120
            assert r["width"] >= 240 and r["height"] >= 168
            assert r["x"] + r["width"] <= 1896 and r["y"] + r["height"] <= 1056
    assert cards[-1]["layout"] == {"x": 1446, "y": 760, "width": 450, "height": 296}


def test_legacy_xls_dates_booleans_and_hidden_rows(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = tmp_path / "legacy.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Data")
    for c, v in enumerate(["ID", "Amount", "Date", "Flag"]):
        ws.write(0, c, v)
    for r in range(1, 4):
        ws.write(r, 0, f"ID{r}")
        ws.write(r, 1, r * 10)
        ws.write(r, 2, datetime(2026, r, 1), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
        ws.write(r, 3, True)
    ws.row(3).hidden = True
    wb.save(str(path))
    spec = block(path, sheet="Data", area="A1:D4", headers=1)
    result = analysis.render(request(path, "line", [spec], x="a:f2", y=["a:f1"], timeBucket="month"))
    assert result["table"]["rows"] == [["2026-01", 10.0], ["2026-02", 20.0]]
    assert result["counts"]["inputRows"] == 2
    with pytest.raises(ValueError, match="没有有效数字"):
        analysis.render(request(path, "metric", [spec], y=["a:f3"]))


def test_missing_and_constant_correlation_is_null_not_fabricated(book):
    wb = openpyxl.load_workbook(book)
    for row in range(5, 10):
        wb["明细"].cell(row, 4, 7)
    wb.save(book)
    wb.close()
    result = analysis.render(request(book, "correlation", y=["a:f2", "a:f3"]))
    assert result["table"]["rows"][1][2] is None
    assert result["table"]["rows"][1][3] == 4
    json.dumps(result, allow_nan=False)


@pytest.mark.parametrize("value,expected", [("1,200.50", 1200.5), ("25%", 0.25), ("￥1,200", 1200), ("-3.5", -3.5)])
def test_explicit_numeric_text_formats(value, expected):
    assert analysis.numeric(value) == expected


def test_large_numeric_join_keys_stay_distinct(book):
    wb = openpyxl.load_workbook(book)
    wb["明细"]["A5"] = 9007199254740992
    wb["关联表"]["A3"] = 9007199254740993
    wb.save(book)
    wb.close()
    # Excel 数字存储精度由文件决定；这里额外测试精确整数键组合，不强制转 float。
    import pandas as pd

    frames = [
        pd.DataFrame({"a:f0": [9007199254740992], "a:f1": [10]}),
        pd.DataFrame({"b:f0": [9007199254740993], "b:f1": [99]}),
    ]
    blocks = [{"id": "a", "name": "a", "fields": []}, {"id": "b", "name": "b", "fields": []}]
    with pytest.raises(ValueError, match="没有匹配"):
        analysis.combine(
            frames, blocks, {"combine": "join", "joinKeys": {"a": "a:f0", "b": "b:f0"}, "joinType": "inner"}
        )
