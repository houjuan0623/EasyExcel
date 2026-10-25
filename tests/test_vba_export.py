"""VBA 结构副本保留代码接口和公式，数据常量仅允许通过表头白名单。

Copyright (c) 2026 AbeLin · MIT License
"""
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableFormula
import pytest

from desktop.bridge import selection_config
from desktop.privacy_export import inspect_privacy, export_privacy
from desktop.vba_export import NS, RID, R, P, C, parse_ranges

DATA = 'PRIVATE_DATA_92837465'
FORMULA_TEXT = 'FORMULA_CRITERION_75319'
AMOUNT = 987654321.75


def rewrite_zip(path, transform):
    with ZipFile(BytesIO(path.read_bytes())) as src:
        parts = {name: src.read(name) for name in src.namelist()}
    transform(parts)
    with ZipFile(path, 'w', ZIP_DEFLATED) as dst:
        for name, content in parts.items():
            dst.writestr(name, content)


@pytest.fixture
def source(tmp_path):
    path = tmp_path / '财务模板.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '财务 汇总'
    ws.sheet_properties.codeName = 'BudgetSheet'
    ws.append([DATA])
    ws.merge_cells('A1:F1')
    ws.append([DATA, AMOUNT])
    ws.append(['部门', '金额', '计算值', '跨表值', '条件值', '备注'])
    for r in range(4,9):
        ws.append([DATA, AMOUNT, f'=B{r}*1.25+42', "='参数区'!$B$2", f'=IF(A{r}="{FORMULA_TEXT}",C{r},0)', DATA])
    ws['B4'].number_format = f'#,##0.00"{DATA}"'
    ws['A4'].font = Font(name=DATA)
    ws['A4'].comment = Comment(DATA, DATA)
    ws['A4'].hyperlink = 'https://example.invalid/' + DATA
    ws.column_dimensions['F'].hidden = True
    ws.row_dimensions[6].hidden = True
    ws.freeze_panes = 'C4'
    ws.conditional_formatting.add('B4:B8', CellIsRule(operator='greaterThan', formula=[str(AMOUNT)]))
    table = Table(displayName='BudgetData', ref='A3:F8')
    table._initialise_columns()
    for col, name in zip(table.tableColumns, ['部门', '金额', '计算值', '跨表值', '条件值', '备注']):
        col.name = name
    table.tableColumns[2].calculatedColumnFormula = TableFormula(attr_text='[金额]*1.25+42')
    ws.add_table(table)
    chart = BarChart()
    chart.title = DATA
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=8), titles_from_data=True)
    ws.add_chart(chart, 'H3')
    ws.protection.sheet = True
    ws.protection.password = 'not-an-exported-secret'
    hidden = wb.create_sheet('参数区')
    hidden.append(['项目', '参数'])
    hidden.append([DATA, AMOUNT])
    hidden.sheet_state = 'veryHidden'
    hidden['C2'] = '=B2*2'
    wb.defined_names.add(DefinedName('AmountInput', attr_text="'财务 汇总'!$B$4:$B$8"))
    wb.defined_names.add(DefinedName('RateConstant', attr_text='0.125'))
    wb.properties.creator = DATA
    wb.properties.description = DATA
    wb.save(path)
    wb.close()
    def inject(parts):
        parts['customXml/item1.xml'] = f'<private>{DATA}</private>'.encode()
        parts['xl/pivotCache/pivotCacheRecords1.xml'] = f'<cache>{DATA}{AMOUNT}</cache>'.encode()
        for name in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml'):
            sheet = ET.fromstring(parts[name])
            for cell in sheet.iter(NS + 'c'):
                if cell.find(NS + 'f') is not None:
                    cell.find(NS + 'v').text = str(AMOUNT)
            parts[name] = ET.tostring(sheet)
    rewrite_zip(path, inject)
    return path


def prepare(source):
    config = selection_config({'inputPath': str(source), 'selection': [{'path': str(source), 'sheets': ['财务 汇总']}], 'headerRow': 3})
    request = {'mode': 'vba', 'headerRanges': {'workbook_001_sheet_001': 'A3:F3', 'workbook_001_sheet_002': 'A1:B1'}}
    plan = inspect_privacy(config, request)
    request.update(token=plan['token'], reviewConfirmed=True, exportId=str(uuid4()), outputPath=str(source.parent/'share.zip'))
    return config, request, plan


def exported(source):
    config, request, plan = prepare(source)
    result = export_privacy(request, config)
    with ZipFile(result['output']) as archive:
        schema = json.loads(archive.read('结构说明.json'))
        payload = archive.read(schema['workbooks'][0]['template'])
        outer = {name: archive.read(name) for name in archive.namelist() if not name.endswith('.xlsx')}
    with ZipFile(BytesIO(payload)) as archive:
        inner = {name: archive.read(name) for name in archive.namelist()}
    return result, plan, schema, payload, outer, inner


def raw_formulas(parts):
    return {(part, cell.get('r')): (dict(f.attrib), f.text) for part, content in parts.items()
            if part.startswith('xl/worksheets/sheet') and part.endswith('.xml')
            for cell in ET.fromstring(content).iter(NS+'c')
            if (f := cell.find(NS+'f')) is not None}


def test_exact_formulas_names_fields_coordinates_and_hidden_dependencies(source):
    before = source.read_bytes()
    result, plan, schema, payload, outer, inner = exported(source)
    with ZipFile(source) as archive:
        original = {name: archive.read(name) for name in archive.namelist()}
    assert raw_formulas(inner) == raw_formulas(original)
    assert result['formulaCount'] == 16
    assert result['sheets'] == 2
    assert plan['vba']['workbooks'][0]['sheets'][1]['state'] == 'veryHidden'
    assert any(FORMULA_TEXT in item['value'] for item in plan['vba']['review'])
    assert any('0.125' in item['value'] for item in plan['vba']['review'])
    wb = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    assert wb.sheetnames == ['财务 汇总', '参数区']
    ws = wb['财务 汇总']
    assert ws['A3'].value == '部门' and ws['F3'].value == '备注'
    assert ws['A4'].value is None and ws['B4'].value is None and ws['F4'].value is None
    assert ws['A1'].value is None and ws['B2'].value is None
    assert ws['C4'].value == '=B4*1.25+42'
    assert ws['D4'].value == "='参数区'!$B$2"
    assert ws['E4'].value == f'=IF(A4="{FORMULA_TEXT}",C4,0)'
    assert ws.column_dimensions['F'].hidden and ws.row_dimensions[6].hidden
    assert str(ws.merged_cells) == 'A1:F1' and ws.freeze_panes == 'C4'
    assert not ws.protection.sheet
    assert schema['workbooks'][0]['sheets'][0]['protected']
    assert ws.tables['BudgetData'].ref == 'A3:F8'
    assert ws.tables['BudgetData'].tableColumns[2].calculatedColumnFormula.text == '[金额]*1.25+42'
    assert wb.defined_names['AmountInput'].attr_text == "'财务 汇总'!$B$4:$B$8"
    assert wb['参数区']['B2'].value is None and wb['参数区']['C2'].value == '=B2*2'
    assert wb.calculation.calcMode == 'manual' and not wb.calculation.fullCalcOnLoad
    wb.close()
    for content in [*outer.values(), *inner.values()]:
        text = content.decode('utf-8')
        assert DATA not in text
        assert str(AMOUNT) not in text
        assert str(source.parent) not in text
    assert FORMULA_TEXT in outer['结构说明.json'].decode()
    assert not any(any(word in name for word in ('chart', 'pivotCache', 'comments', 'sharedStrings', 'customXml', 'vbaProject')) for name in inner)
    for name, content in inner.items():
        if name.startswith('xl/worksheets/sheet'):
            for cell in ET.fromstring(content).iter(NS+'c'):
                if cell.find(NS+'f') is not None:
                    assert cell.find(NS+'v') is None and cell.find(NS+'is') is None
    assert b'ThisWorkbook' in outer['workbook_001/运行前结构检查.bas']
    assert outer['workbook_001/运行前结构检查.bas'].isascii()
    assert not (source.parent/'privacy-mappings').exists()
    assert source.read_bytes() == before


def test_shared_and_array_formulas_preserve_xml_attributes(source):
    def inject(parts):
        root = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        cells = {c.get('r'): c for c in root.iter(NS+'c')}
        cells['C4'].find(NS+'f').attrib.update(t='shared', si='0', ref='C4:C5')
        cells['C5'].find(NS+'f').attrib.update(t='shared', si='0')
        cells['C5'].find(NS+'f').text = None
        formula = cells['D7'].find(NS+'f')
        formula.text = '_xlfn.UNIQUE(A4:A6)'
        formula.attrib.update(t='array', ref='D7:D8', aca='1')
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(root)
    rewrite_zip(source, inject)
    _, _, _, _, _, actual = exported(source)
    with ZipFile(source) as z:
        expected = {name: z.read(name) for name in z.namelist()}
    assert raw_formulas(actual) == raw_formulas(expected)


def test_external_formula_skeleton_retains_target_but_no_saved_values(source):
    external_target = 'file:///C:/LocalBudget/reference.xlsx'
    def inject(parts):
        workbook = ET.fromstring(parts['xl/workbook.xml'])
        external = ET.SubElement(workbook, NS+'externalReferences')
        ET.SubElement(external, NS+'externalReference', {RID:'rIdExt'})
        parts['xl/workbook.xml'] = ET.tostring(workbook)
        rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        ET.SubElement(rels, '{'+P+'}Relationship', {'Id':'rIdExt','Type':R+'/externalLink','Target':'externalLinks/externalLink1.xml'})
        parts['xl/_rels/workbook.xml.rels'] = ET.tostring(rels)
        types = ET.fromstring(parts['[Content_Types].xml'])
        ET.SubElement(types, '{'+C+'}Override', {'PartName':'/xl/externalLinks/externalLink1.xml','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml'})
        parts['[Content_Types].xml'] = ET.tostring(types)
        root = ET.Element(NS+'externalLink')
        book = ET.SubElement(root,NS+'externalBook',{RID:'rId1'})
        names = ET.SubElement(book,NS+'sheetNames')
        ET.SubElement(names,NS+'sheetName',{'val':'外部表'})
        data = ET.SubElement(ET.SubElement(book,NS+'sheetDataSet'),NS+'sheetData',{'sheetId':'0'})
        row = ET.SubElement(data,NS+'row',{'r':'1'})
        cell = ET.SubElement(row,NS+'cell',{'r':'A1','t':'str'})
        ET.SubElement(cell,NS+'v').text = DATA
        parts['xl/externalLinks/externalLink1.xml'] = ET.tostring(root)
        rels = ET.Element('{'+P+'}Relationships')
        ET.SubElement(rels,'{'+P+'}Relationship',{'Id':'rId1','Type':R+'/externalLinkPath','Target':external_target,'TargetMode':'External'})
        parts['xl/externalLinks/_rels/externalLink1.xml.rels'] = ET.tostring(rels)
        sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        next(c for c in sheet.iter(NS+'c') if c.get('r')=='C8').find(NS+'f').text = "'[1]外部表'!A1+10"
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
    rewrite_zip(source, inject)
    _, plan, schema, payload, _, parts = exported(source)
    assert any(item['value']==external_target for item in plan['vba']['review'])
    assert schema['workbooks'][0]['externalLinks'][0]['target'] == external_target
    assert b'sheetDataSet' not in parts['xl/externalLinks/externalLink1.xml']
    assert DATA.encode() not in b''.join(parts.values())
    assert ET.fromstring(parts['xl/workbook.xml']).find(NS+'workbookPr').get('updateLinks') == 'never'
    wb = openpyxl.load_workbook(BytesIO(payload), keep_links=True)
    assert wb['财务 汇总']['C8'].value == "='[1]外部表'!A1+10"
    assert wb._external_links[0].externalBook.sheetDataSet is None
    wb.close()


def test_review_and_source_hash_are_required(source):
    config, request, _ = prepare(source)
    with pytest.raises(ValueError, match='检查'):
        export_privacy({**request,'reviewConfirmed':False},config)
    assert not Path(request['outputPath']).exists()
    with ZipFile(source,'a') as archive:
        archive.writestr('customXml/changed.xml','<changed/>')
    with pytest.raises(ValueError, match='已变化'):
        export_privacy(request,config)
    assert not Path(request['outputPath']).exists()


def test_header_range_change_requires_new_review_and_keeps_only_whitelist(source):
    config, request, _ = prepare(source)
    changed = {**request,'headerRanges':{**request['headerRanges'],'workbook_001_sheet_002':''}}
    with pytest.raises(ValueError, match='已变化'):
        export_privacy(changed,config)
    plan = inspect_privacy(config,changed)
    changed['token'] = plan['token']
    result = export_privacy(changed,config)
    with ZipFile(result['output']) as outer:
        wb = openpyxl.load_workbook(BytesIO(outer.read('workbook_001/VBA结构模板.xlsx')))
        assert wb['参数区']['A1'].value is None
        assert wb['参数区']['C2'].value == '=B2*2'
        wb.close()
    with pytest.raises(ValueError,match='表对象的表头'):
        inspect_privacy(config,{**request,'headerRanges':{'workbook_001_sheet_001':'A3:E3'}})


@pytest.mark.parametrize('value',['A:A','1:5','A0:B2','B4:A3','A1:XFD100','A1:A1048577','Other!A1:B2'])
def test_header_regions_must_be_explicit_and_bounded(value):
    with pytest.raises(ValueError):
        parse_ranges(value)


def test_multiple_files_preserve_same_sheet_names_without_combining(source):
    second = source.with_name('第二本.xlsx')
    second.write_bytes(source.read_bytes())
    config = selection_config({'inputPath':str(source.parent),'selection':[{'path':str(p),'sheets':['财务 汇总']} for p in (source,second)],'headerRow':3})
    request = {'mode':'vba','headerRanges':{f'workbook_{n:03d}_sheet_00{s}': 'A3:F3' if s==1 else 'A1:B1' for n in (1,2) for s in (1,2)}}
    plan = inspect_privacy(config,request)
    result = export_privacy({**request,'token':plan['token'],'reviewConfirmed':True,'exportId':str(uuid4()),'outputPath':str(source.parent/'both.zip')},config)
    assert result['sheets'] == 4 and result['workbooks'] == 2
    with ZipFile(result['output']) as archive:
        for n in (1,2):
            wb = openpyxl.load_workbook(BytesIO(archive.read(f'workbook_{n:03d}/VBA结构模板.xlsx')))
            assert wb.sheetnames == ['财务 汇总','参数区']
            wb.close()


def test_xlsm_exports_clean_xlsx_without_existing_macro_project(source):
    macro = source.with_suffix('.xlsm')
    macro.write_bytes(source.read_bytes())
    with ZipFile(macro,'a') as archive:
        archive.writestr('xl/vbaProject.bin',DATA)
    _,_,_,_,_,parts = exported(macro)
    assert 'xl/vbaProject.bin' not in parts
    assert b'macroEnabled' not in parts['[Content_Types].xml']


def test_missing_dimension_still_describes_original_cell_extent(source):
    def remove_dimension(parts):
        sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        sheet.remove(sheet.find(NS+'dimension'))
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
    rewrite_zip(source, remove_dimension)
    _, _, schema, _, _, _ = exported(source)
    assert schema['workbooks'][0]['sheets'][0]['dataExtent'] == 'A1:F8'
