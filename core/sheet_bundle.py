"""将工作簿中的表逐张整理为一个多 Sheet 工作簿，不拼接数据行。

Copyright (c) 2026 AbeLin · MIT License
"""
from copy import copy, deepcopy
import datetime
import json
from pathlib import Path
import re
from uuid import uuid4

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.worksheet.print_settings import PrintArea

from core.splitter import _copy_cell_style, _safe_sheet_title, detect_uncalculated_formulas
from core.xlio import _head_limit, hidden_info, open_book, resolve_header_row


class _Stopped(Exception):
    """协作式停止，不将未完成的工作簿当作完整结果保存。"""


def _check_stop(stop):
    if stop and stop():
        raise _Stopped


def _local_formula(formula):
    """只保留无需其他 Sheet、工作簿或定义名称的公式引用。"""
    if not isinstance(formula, str):
        return False
    try:
        tokens = Tokenizer(formula if formula.startswith('=') else '=' + formula).items
        return all(re.fullmatch(r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?|"
                               r"\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d+:\$?\d+", t.value)
                   for t in tokens if t.type == 'OPERAND' and t.subtype == 'RANGE')
    except Exception:
        return False


def _copy_sheet(src, dst, hidden_cols, header, stop, log):
    """跨工作簿复制实际存在的单元格及布局，样式编号按目标工作簿重新注册。"""
    cache = {}
    data_rows = set()
    for index, cell in enumerate(src._cells.values()):
        if index % 1024 == 0:
            _check_stop(stop)
        out = dst.cell(cell.row, cell.column)
        omitted = cell.column in hidden_cols and cell.row > header
        if not omitted:
            out.value, out.data_type = cell.value, cell.data_type
            if cell.row > header and cell.value is not None:
                data_rows.add(cell.row)
            if cell.hyperlink:
                out.hyperlink = copy(cell.hyperlink)
            if cell.comment:
                out.comment = copy(cell.comment)
        if cell.has_style:
            _copy_cell_style(cell, out, cache)
            out.protection = copy(cell.protection)
    for attr in ('row_dimensions', 'column_dimensions'):
        for key, dim in getattr(src, attr).items():
            target = copy(dim)
            target.parent = dst
            target._style = None
            if dim.has_style:
                _copy_cell_style(dim, target, cache)
                target.protection = copy(dim.protection)
            getattr(dst, attr)[key] = target
    for merged in src.merged_cells.ranges:
        dst.merge_cells(str(merged))
    for attr in ('sheet_format', 'sheet_properties', 'views', 'auto_filter', 'page_margins',
                 'page_setup', 'print_options', 'protection', 'HeaderFooter', 'row_breaks', 'col_breaks'):
        setattr(dst, attr, copy(getattr(src, attr)))
    dst.page_setup._parent = dst
    # 打印机二进制设置不随跨工作簿复制，避免残留失效的关系编号。
    dst.page_setup.id = None
    dst.freeze_panes = src.freeze_panes
    if src.print_area:
        dst.print_area = [str(r) for r in PrintArea.from_string(src.print_area).ranges]
    dst.print_title_rows = src.print_title_rows
    dst.print_title_cols = src.print_title_cols
    dropped = 0
    for validation in src.data_validations.dataValidation:
        if all(not f or _local_formula(f) for f in (validation.formula1, validation.formula2)):
            dst.add_data_validation(deepcopy(validation))
        else:
            dropped += 1
    for region, rules in src.conditional_formatting._cf_rules.items():
        for rule in rules:
            if all(_local_formula(f) for f in rule.formula):
                dst.conditional_formatting.add(str(region.sqref), deepcopy(rule))
            else:
                dropped += 1
    if dropped:
        log(f"  ⚠️ {src.title}：{dropped} 项引用其他表或名称的验证/条件格式未复制")
    if src._images or src._charts or src.tables:
        log(f"  ⚠️ {src.title}：保留单元格与布局；图片、图表和 Excel 表对象不复制")
    return len(data_rows)


def bundle_workbooks(sources, output, config=None, log_fn=None, progress_fn=None, stop_flag=None,
                     split_results=False):
    """sources 为 (路径, Sheet 名列表或 None)；仅整理明确提供的文件，不重新扫描输出目录。

    split_results 只供拆分内核调用，保留分表中已安全平移的同行公式；普通整理读保存的计算值。
    停止或保存失败时不留下看似完整的总表，已存在的文件绝不覆盖。
    """
    cfg, log = config or {}, log_fn or (lambda _: None)
    output = Path(output)
    if output.exists():
        raise FileExistsError(f"结果已存在，请换一个文件名：{output.name}")
    sources = list(sources)
    summary = dict(mode="sheet_bundle", files=0, rows=0, sheets=0, failed_files=0, failed_saves=0,
                   skipped_sheets=0, stopped=False, output="")
    book = openpyxl.Workbook()
    book.remove(book.active)
    used = set()
    temporary = output.with_name(f".{output.stem}-{uuid4().hex}.tmp")
    try:
        for index, (path, names) in enumerate(sources):
            _check_stop(stop_flag)
            source = None
            added, rows = [], 0
            try:
                path = str(path)
                if split_results:
                    source = openpyxl.load_workbook(path, data_only=False)
                else:
                    source, is_xls = open_book(path)
                    if is_xls:
                        log(f"  ℹ️ {Path(path).name}：旧版 .xls 按数据值整理，原格式无法完整保留")
                    elif detect_uncalculated_formulas(path):
                        log(f"  ⚠️ {Path(path).name}：发现没有已保存计算值的公式，对应结果可能为空；请先用 Excel/WPS 计算并保存后再整理")
                hidden = hidden_info(path)
                sheets = [ws for ws in source.worksheets if names is None or ws.title in names]
                missing = set(names or ()) - set(source.sheetnames)
                if missing:
                    raise ValueError("工作表不存在：" + "、".join(sorted(missing)))
                visible = [ws for ws in sheets if hidden.get(ws.title, (ws.sheet_state, set()))[0] == 'visible']
                summary['skipped_sheets'] += len(sheets) - len(visible)
                for ws in visible:
                    _check_stop(stop_flag)
                    # 同一文件选多张表时保留文件名与原 Sheet 名，避免默默丢掉第二张表。
                    name = Path(path).stem if len(visible) == 1 else f"{Path(path).stem}_{ws.title}"
                    name = re.sub(r"[\x00-\x1f]", "_", name).strip(" '") or "Sheet"
                    title = _safe_sheet_title(name, set()).strip("'") or "Sheet"
                    if title.lower() == 'history':
                        title = 'History_'
                    title = _safe_sheet_title(title, used)
                    target = book.create_sheet(title)
                    added.append(target)
                    head = list(ws.iter_rows(max_row=min(ws.max_row, _head_limit(cfg)), values_only=True))
                    header = resolve_header_row(head, cfg)
                    hidden_cols = set() if split_results or cfg.get('keep_hidden') else hidden.get(ws.title, ('visible', set()))[1]
                    rows += _copy_sheet(ws, target, hidden_cols, max(0, header), stop_flag, log)
                    log(f"  {Path(path).name} / {ws.title} → {target.title}")
                summary['rows'] += rows
                summary['sheets'] += len(added)
            except _Stopped:
                raise
            except Exception as exc:
                for ws in added:
                    book.remove(ws)
                summary['failed_files'] += 1
                log(f"  ❌ 整理失败 {Path(path).name}：{exc}")
            finally:
                if source is not None:
                    source.close()
            if progress_fn:
                progress_fn(0.9 * (index + 1) / max(1, len(sources)))
        _check_stop(stop_flag)
        if book.worksheets:
            output.parent.mkdir(parents=True, exist_ok=True)
            book.save(temporary)
            _check_stop(stop_flag)
            # 'xb' 在最终写入时仍检查同名冲突；失败时只清理本次创建的文件。
            with temporary.open('rb') as src, output.open('xb') as dst:
                try:
                    for chunk in iter(lambda: src.read(1 << 20), b''):
                        _check_stop(stop_flag)
                        dst.write(chunk)
                except BaseException:
                    dst.close()
                    output.unlink(missing_ok=True)
                    raise
            summary.update(files=1, output=str(output))
            log(f"✅ 已整理为 1 个工作簿，包含 {summary['sheets']} 个 Sheet：{output.name}")
    except _Stopped:
        summary['stopped'] = True
        log("⛔ 已停止整理，总工作簿未生成；原文件保留")
    except Exception as exc:
        summary['failed_saves'] += 1
        log(f"❌ 总工作簿保存失败：{exc}")
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    if progress_fn and not summary['stopped']:
        progress_fn(1.0)
    return summary


def merge_as_sheets(config, log_fn=None, progress_fn=None, stop_flag=None):
    """桌面「每个表独立一个 Sheet」入口，与按列名拼接数据的汇总分别处理。"""
    import os

    log = log_fn or (lambda _: None)
    selected = config['sheet_selection']
    sources = [(path, selected[os.path.normcase(os.path.abspath(path))])
               for path in config['selected_input_files']]
    root = Path(config['output_path'])
    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    output = root / f"工作表合集_{stamp}.xlsx"
    index = 2
    while output.exists():
        output = root / f"工作表合集_{stamp}_{index}.xlsx"
        index += 1
    log("将每张勾选的表保存在独立 Sheet 中；单元格公式使用源文件已保存的计算值。")
    summary = bundle_workbooks(sources, output, config, log, progress_fn, stop_flag)
    log(f"[SUMMARY] {json.dumps(summary, ensure_ascii=False)}")
    return summary['output']
