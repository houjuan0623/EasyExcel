"""读取内嵌图片和 Excel 图表并离线绘制；不读取外部文件或联网资源。

Copyright (c) 2026 AbeLin · MIT License
"""
import base64
import datetime
from io import BytesIO
import math
import posixpath
import re
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils.cell import range_boundaries

REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
MAX_CELLS = 100000
MAX_POINTS = 5000
MAX_OBJECTS = 16
MAX_MEDIA = 8 * 1024 * 1024
CHART_COLORS = ['#4472C4', '#ED7D31', '#70AD47', '#A5A5A5', '#FFC000', '#5B9BD5']


def xml(archive, path):
    """仅解析大小受限的包内 XML 部件。"""
    if archive.getinfo(path).file_size > MAX_MEDIA:
        raise ValueError('XML 部件过大')
    return ET.fromstring(archive.read(path))


def relationships(archive, path):
    """只接受 ZIP 包内关系，不访问 TargetMode=External。"""
    directory, name = posixpath.split(path)
    part = posixpath.join(directory, '_rels', name + '.rels')
    if part not in archive.namelist():
        return {}
    result = {}
    for rel in xml(archive, part):
        if rel.get('TargetMode') == 'External':
            continue
        target = rel.get('Target', '')
        resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(directory, target))
        if not resolved.startswith('../') and resolved in archive.namelist():
            result[rel.get('Id')] = resolved
    return result


class Values:
    """按受限引用范围读取已保存的值，隐藏行列和 Sheet 不进入图表。"""
    def __init__(self, path, book, layout_fn, layouts):
        self.path, self.book, self.layout_fn = path, book, layout_fn
        self.layouts, self.cache, self.prime = layouts, {}, {}
        self.data_book = None
        self.count = 0

    def layout(self, sheet):
        if sheet not in self.layouts:
            self.layouts[sheet] = self.layout_fn(self.book[sheet])
        return self.layouts[sheet]

    def read(self, reference, current_sheet):
        reference = reference.strip().lstrip('=')
        name, span = current_sheet, reference
        if '!' in reference:
            name, span = reference.rsplit('!', 1)
            name = name[1:-1].replace("''", "'") if name.startswith("'") and name.endswith("'") else name
        if '[' in name or name not in self.book.sheetnames or self.book[name].sheet_state != 'visible':
            raise ValueError('图表或条件格式引用了外部文件、隐藏或不存在的工作表')
        if not re.fullmatch(r'\$?[A-Za-z]{1,3}\$?\d*(?::\$?[A-Za-z]{1,3}\$?\d*)?', span):
            raise ValueError('不支持命名区域或复杂引用')
        key = name, span.upper().replace('$', '')
        if key in self.cache:
            return self.cache[key]
        left, top, right, bottom = range_boundaries(span)
        layout = self.layout(name)
        top, bottom = top or 1, bottom or max(1, layout['lastRow'])
        size = (right-left+1) * (bottom-top+1)
        if size < 1 or size > MAX_CELLS or self.count + size > MAX_CELLS or bottom > 100000:
            raise ValueError('引用范围超过预览计算上限')
        hidden_columns = set()
        for col in layout['columns']:
            if col.get('hidden') in ('1', 'true') or float(col.get('width', layout['defaultWidth'])) == 0:
                hidden_columns.update(range(int(col['min']), int(col['max'])+1))
        if size == 1 and (name, top, left) in self.prime:
            values = [self.prime[name, top, left]]
        else:
            if self.data_book is None:
                self.data_book = openpyxl.load_workbook(self.path, read_only=True, data_only=True, keep_links=False)
            ws = self.data_book[name]
            ws.reset_dimensions()
            values = [value for row in ws.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right, values_only=True) for value in row]
        flags = [r in layout['hiddenRowIndices'] or c in hidden_columns for r in range(top, bottom+1) for c in range(left, right+1)]
        values = [None if hidden else value for value, hidden in zip(values, flags)]
        self.count += size
        self.cache[key] = values, flags
        return values, flags

    def close(self):
        if self.data_book:
            self.data_book.close()


def child_value(node, name, default=None):
    item = node.find('{*}' + name) if node is not None else None
    return item.get('val', default) if item is not None else default


def title(node, fallback):
    text = ''.join(node.itertext()) if node is not None and node.tag.endswith('}v') else ''
    if node is not None:
        text = ''.join(t.text or '' for t in node.findall('.//{*}t')) or text
    return text[:200] or fallback


def series_values(node, reader, sheet, warnings):
    """优先读取源单元格；无公式缓存时只使用文件自带的图表缓存并提示。"""
    if node is None:
        return [], []
    points = {int(pt.get('idx', '0')): pt.findtext('{*}v') for pt in node.findall('.//{*}pt')}
    if points and max(points) >= MAX_POINTS:
        raise ValueError('图表超过 5000 个数据点')
    cached = [points.get(i) for i in range(max(points, default=-1)+1)]
    ref = node.find('.//{*}f')
    if ref is not None and ref.text:
        values, hidden = reader.read(ref.text, sheet)
        if len(values) > MAX_POINTS:
            raise ValueError('图表超过 5000 个数据点')
        if cached and all(v is None for v, h in zip(values, hidden) if not h):
            values = [None if h else cached[i] if i < len(cached) else None for i, h in enumerate(hidden)]
            warnings.add('部分图表使用工作簿已保存的图表缓存；预览不重新计算 Excel 公式')
        return values, hidden
    return cached, [False] * len(cached)


def _numeric(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        import matplotlib.dates as dates
        return dates.date2num(value)
    try:
        return float(value) if value is not None and math.isfinite(float(value)) else math.nan
    except (TypeError, ValueError):
        return math.nan


def configure_axis(axis, node, direction, percent=False):
    """还原常见数值轴的标题、百分比、边界、对数与反向设置。"""
    from matplotlib.ticker import FuncFormatter, PercentFormatter
    label = title(node.find('{*}title'), '')
    getattr(axis, f'set_{direction}label')(label)
    target = axis.xaxis if direction == 'x' else axis.yaxis
    fmt = node.find('{*}numFmt')
    code = fmt.get('formatCode', '') if fmt is not None else ''
    if percent or '%' in code:
        target.set_major_formatter(PercentFormatter(1))
    elif '#,##' in code:
        target.set_major_formatter(FuncFormatter(lambda x, _pos: f'{x:,.0f}'))
    scaling = node.find('{*}scaling')
    if scaling is not None:
        log = child_value(scaling, 'logBase')
        if log:
            getattr(axis, f'set_{direction}scale')('log', base=float(log))
        low, high = child_value(scaling, 'min'), child_value(scaling, 'max')
        if low is not None or high is not None:
            getattr(axis, f'set_{direction}lim')(float(low) if low is not None else None, float(high) if high is not None else None)
        if child_value(scaling, 'orientation') == 'maxMin':
            getattr(axis, f'invert_{direction}axis')()


def render_chart(root, reader, sheet, warnings):
    """重绘常见二维统计图，保留标题、系列、数据、图例及堆积关系。"""
    import matplotlib
    matplotlib.use('Agg', force=True)
    from matplotlib import font_manager, rc_context
    from matplotlib.figure import Figure

    plot = root.find('.//{*}plotArea')
    if plot is None:
        raise ValueError('图表缺少绘图区')
    supported = {'barChart', 'lineChart', 'pieChart', 'doughnutChart', 'areaChart', 'scatterChart', 'radarChart'}
    groups = [node for node in plot if node.tag.rsplit('}', 1)[-1].endswith('Chart')]
    if not groups or any(node.tag.rsplit('}', 1)[-1] not in supported for node in groups):
        raise ValueError('暂不支持此图表类型（如三维、透视或扩展图表）')
    kinds = {node.tag.rsplit('}', 1)[-1] for node in groups}
    if len(groups) > 1 and (not kinds <= {'barChart', 'lineChart', 'areaChart'} or
                            any(child_value(g, 'barDir') == 'bar' for g in groups)):
        raise ValueError('此组合图暂不支持，支持二维柱形、折线和面积的组合')
    value_axes = {child_value(node, 'axId'): node for node in plot if node.tag.endswith('}valAx')}
    chart_title = title(root.find('.//{*}chart/{*}title'), '统计图')
    installed = {font.name for font in font_manager.fontManager.ttflist}
    fonts = [n for n in ['Microsoft YaHei', 'SimHei', 'Noto Sans CJK SC', 'WenQuanYi Micro Hei'] if n in installed] + ['DejaVu Sans']
    with rc_context({'font.sans-serif': fonts, 'axes.unicode_minus': False, 'font.size': 10}):
        figure = Figure(figsize=(8, 4.6), dpi=125, layout='constrained', facecolor='white')
        radar = groups[0].tag.endswith('}radarChart')
        primary = figure.add_subplot(111, projection='polar' if radar else None)
        axes, configured = {}, set()
        series_count, plotted = 0, 0
        for group in groups:
            kind = group.tag.rsplit('}', 1)[-1]
            ids = [node.get('val') for node in group.findall('{*}axId')]
            axis_id = ids[-1] if ids else 'primary'
            if axis_id not in axes:
                if len(axes) >= 2:
                    raise ValueError('超过两组坐标轴的图表暂不支持')
                axes[axis_id] = primary if not axes else primary.twinx()
            axis = axes[axis_id]
            series = group.findall('{*}ser')
            if not series or len(series) > 20:
                raise ValueError('图表系列数量为空或超过 20 组')
            grouping = child_value(group, 'grouping', 'standard')
            stacked = grouping in ('stacked', 'percentStacked')
            horizontal = kind == 'barChart' and child_value(group, 'barDir') == 'bar'
            stacks_pos, stacks_neg = {}, {}
            prepared = []
            for i, ser in enumerate(series):
                vals, hidden = series_values(ser.find('{*}yVal') if kind == 'scatterChart' else ser.find('{*}val'), reader, sheet, warnings)
                cats, cat_hidden = series_values(ser.find('{*}xVal') if kind == 'scatterChart' else ser.find('{*}cat'), reader, sheet, warnings)
                keep = [j for j in range(len(vals)) if not hidden[j] and (j >= len(cat_hidden) or not cat_hidden[j])]
                ys = [_numeric(vals[j]) for j in keep]
                categories = [cats[j] if j < len(cats) and cats[j] is not None else j+1 for j in keep]
                labels = [str(value) for value in categories]
                tx = ser.find('{*}tx')
                named, _ = series_values(tx, reader, sheet, warnings)
                label = str(named[0]) if named and named[0] is not None else (tx.findtext('{*}v') if tx is not None else None) or f'系列 {i+1}'
                rgb = ser.find('.//{*}spPr/{*}solidFill/{*}srgbClr')
                if rgb is None:
                    rgb = ser.find('.//{*}spPr/{*}ln/{*}solidFill/{*}srgbClr')
                color = '#' + rgb.get('val') if rgb is not None and re.fullmatch('[0-9A-Fa-f]{6}', rgb.get('val', '')) else CHART_COLORS[series_count % len(CHART_COLORS)]
                prepared.append((ys, labels, label, color, categories))
                series_count += 1
            for i, (ys, labels, label, color, categories) in enumerate(prepared):
                if not any(math.isfinite(y) for y in ys):
                    warnings.add(f'图表「{chart_title}」的部分系列没有可用计算值')
                    continue
                plotted += 1
                xs = list(range(len(ys)))
                if grouping == 'percentStacked':
                    ys = [y / (sum(abs(p[0][j]) for p in prepared if j < len(p[0]) and math.isfinite(p[0][j])) or 1) for j, y in enumerate(ys)]
                if kind == 'barChart':
                    bottoms = [stacks_pos.get(j, 0) if y >= 0 else stacks_neg.get(j, 0) for j, y in enumerate(ys)] if stacked else [0] * len(ys)
                    thickness = .75 if stacked else .75 / len(series)
                    positions = xs if stacked else [x + (i-(len(series)-1)/2) * thickness for x in xs]
                    if horizontal:
                        axis.barh(positions, ys, thickness, left=bottoms, label=label, color=color)
                    else:
                        axis.bar(positions, ys, thickness, bottom=bottoms, label=label, color=color)
                    for j, y in enumerate(ys):
                        if math.isfinite(y):
                            store = stacks_pos if y >= 0 else stacks_neg
                            store[j] = store.get(j, 0) + y
                elif kind in ('pieChart', 'doughnutChart'):
                    if i > 0 or any(y < 0 for y in ys if math.isfinite(y)):
                        raise ValueError('多系列圆环或含负值的饼图暂不支持')
                    valid = [(text, y) for text, y in zip(labels, ys) if math.isfinite(y) and y > 0]
                    if not valid:
                        raise ValueError('饼图没有可用的正数')
                    axis.pie([y for _, y in valid], labels=[text for text, _ in valid], colors=CHART_COLORS,
                             autopct='%1.1f%%', wedgeprops={'width': 1-float(child_value(group, 'holeSize', 50))/100} if kind == 'doughnutChart' else None)
                elif kind == 'scatterChart':
                    xs = [_numeric(x) for x in categories]
                    scatter_style = child_value(group, 'scatterStyle', 'marker')
                    if scatter_style in ('line', 'lineMarker', 'smooth', 'smoothMarker'):
                        axis.plot(xs, ys, label=label, color=color, marker='o' if 'Marker' in scatter_style else None)
                    else:
                        axis.scatter(xs, ys, label=label, color=color, s=24)
                elif kind == 'radarChart':
                    angles = [j * 2 * math.pi / len(ys) for j in range(len(ys))]
                    axis.plot(angles + angles[:1], ys + ys[:1], label=label, color=color)
                    axis.set_xticks(angles, labels)
                elif kind == 'areaChart':
                    if stacked and any(y < 0 for y in ys):
                        raise ValueError('含负值的堆积面积图暂不支持')
                    bottoms = [stacks_pos.get(j, 0) for j in xs] if stacked else [0] * len(ys)
                    tops = [a+b for a, b in zip(bottoms, ys)]
                    axis.fill_between(xs, bottoms, tops, label=label, color=color, alpha=.7)
                    stacks_pos.update(enumerate(tops))
                else:
                    if stacked:
                        ys = [y + stacks_pos.get(j, 0) for j, y in enumerate(ys)]
                        stacks_pos.update(enumerate(ys))
                    axis.plot(xs, ys, label=label, color=color, marker='o' if len(ys) <= 30 else None, linewidth=2)
                if kind not in ('pieChart', 'doughnutChart', 'scatterChart', 'radarChart'):
                    step = max(1, math.ceil(len(labels)/12))
                    ticks = list(range(0, len(labels), step))
                    if horizontal:
                        axis.set_yticks(ticks, [labels[j] for j in ticks])
                    else:
                        axis.set_xticks(ticks, [labels[j] for j in ticks], rotation=25 if max(map(len, labels), default=0) > 6 else 0)
            if axis_id in value_axes and axis_id not in configured and not radar:
                configure_axis(axis, value_axes[axis_id], 'x' if horizontal else 'y', grouping == 'percentStacked')
                configured.add(axis_id)
            for node in plot:
                if node.tag.endswith(('}catAx', '}dateAx')) and child_value(node, 'axId') in ids:
                    getattr(axis, 'set_ylabel' if horizontal else 'set_xlabel')(title(node.find('{*}title'), ''))
                elif kind == 'scatterChart' and node.tag.endswith('}valAx') and child_value(node, 'axId') == ids[0]:
                    configure_axis(axis, node, 'x')
        if not plotted:
            raise ValueError('没有可绘制的数据；公式需先保存计算结果')
        primary.set_title(chart_title, pad=14, fontweight='bold')
        if not kinds & {'pieChart', 'doughnutChart', 'radarChart'}:
            primary.grid(axis='x' if horizontal else 'y', alpha=.2)
            primary.set_axisbelow(True)
            primary.spines['top'].set_visible(False)
            if len(axes) == 1:
                primary.spines['right'].set_visible(False)
        if root.find('.//{*}legend') is not None:
            handles, labels = [], []
            for axis in axes.values():
                h, texts = axis.get_legend_handles_labels()
                handles.extend(h)
                labels.extend(texts)
            if handles:
                primary.legend(handles, labels, frameon=False, fontsize=9, loc='best')
        output = BytesIO()
        figure.savefig(output, format='png')
        figure.clear()
        return chart_title, output.getvalue()


def _anchor(node):
    """读取绘图锚点与 EMU 尺寸；最终位置由网格的可见行列计算。"""
    start = node.find('{*}from')
    end = node.find('{*}to')
    ext = node.find('{*}ext')
    if start is None:
        return None
    def marker(item):
        return {key: int(item.findtext('{*}' + key, '0')) for key in ('col', 'row', 'colOff', 'rowOff')}
    result = {'from': marker(start)}
    if end is not None:
        result['to'] = marker(end)
    elif ext is not None:
        result['width'], result['height'] = int(ext.get('cx', '0')) / 9525, int(ext.get('cy', '0')) / 9525
    return result


def drawings(book, ws, reader, warnings):
    """逐对象容错，返回 PNG 数据与原锚点；不把不支持的图表画成另一种图。"""
    from PIL import Image
    archive, assets, total = book._archive, [], 0
    rels = relationships(archive, ws._worksheet_path)
    for rid in reader.layout(ws.title)['drawingIds']:
        drawing_path = rels.get(rid)
        if not drawing_path:
            warnings.add('部分绘图引用缺失或属于外部资源，未加载')
            continue
        try:
            root = xml(archive, drawing_path)
            children = relationships(archive, drawing_path)
            for node in root:
                if len(assets) >= MAX_OBJECTS:
                    warnings.add('本次最多预览 16 个图表/图片')
                    return assets
                entry = {'id': f'drawing-{len(assets)+1}', 'anchor': _anchor(node), 'kind': 'image', 'title': '嵌入图片'}
                try:
                    chart = node.find('.//{*}chart')
                    blip = node.find('.//{*}blip')
                    if chart is not None:
                        entry['kind'] = 'chart'
                        path = children.get(chart.get(REL + 'id'))
                        if not path:
                            raise ValueError('图表引用缺失或是外部资源')
                        chart_root = xml(archive, path)
                        entry['title'] = title(chart_root.find('.//{*}chart/{*}title'), '统计图')
                        entry['title'], data = render_chart(chart_root, reader, ws.title, warnings)
                    elif blip is not None:
                        path = children.get(blip.get(REL + 'embed'))
                        if not path or archive.getinfo(path).file_size > MAX_MEDIA:
                            raise ValueError('图片是外部资源或超过 8 MB')
                        with Image.open(BytesIO(archive.read(path))) as picture:
                            if picture.width * picture.height > 16000000:
                                raise ValueError('图片分辨率过大')
                            picture.thumbnail((1600, 1200))
                            encoded = BytesIO()
                            picture.convert('RGBA').save(encoded, format='PNG')
                            data = encoded.getvalue()
                    else:
                        warnings.add('部分形状或 SmartArt 暂不支持预览')
                        continue
                    total += len(data)
                    if total > MAX_MEDIA:
                        raise ValueError('图表/图片合计超过 8 MB 预览上限')
                    entry['src'] = 'data:image/png;base64,' + base64.b64encode(data).decode('ascii')
                except Exception as exc:
                    entry['error'] = str(exc)
                    warnings.add(f'{entry["title"]}：{exc}')
                assets.append(entry)
        except Exception:
            warnings.add('部分绘图部件无法读取，已保留单元格预览')
    return assets
