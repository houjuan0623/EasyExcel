"""本地统计分析：显式区域、跨表组合、可解释图表与离线 dashboard。

Copyright (c) 2026 AbeLin · MIT License
"""

import base64
from contextlib import ExitStack
from datetime import date, datetime
from hashlib import sha256
from html import escape
from io import BytesIO, StringIO
import csv
import json
import math
import os
from pathlib import Path
import re
import shutil
from uuid import uuid4
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_ROWS = 100_000
MAX_CELLS = 500_000
MAX_FIELDS = 256
KINDS = {"metric", "bar", "line", "pie", "histogram", "scatter", "box", "correlation"}
AGGS = {"sum": "合计", "mean": "平均值", "median": "中位数", "min": "最小值", "max": "最大值", "count": "有效值计数"}
COLORS = ["#188263", "#548bc0", "#e0a547", "#9c7dc4", "#df806f", "#73a8a1"]


def integer(value, low, high, name):
    if type(value) is not int or not low <= value <= high:
        raise ValueError(f"{name}须为 {low}～{high} 的整数")
    return value


def text(value, name, limit=200):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError(f"{name}不能为空，且最多 {limit} 个字符")
    return value.strip()


def plain(value):
    """协议和统计表不包含 NaN/Infinity；日期转换为可读 ISO 文本。"""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (str, bool, int)):
        return value
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (ValueError, TypeError):
        return str(value)


def numeric(value):
    """数值字符串允许千分位、货币符号和百分号；日期、布尔和不明确单位不转数值。"""
    if value is None or isinstance(value, (bool, datetime, date)):
        return math.nan
    if isinstance(value, str):
        value = value.strip().replace("，", ",")
        if value[:1] in ("¥", "￥", "$"):
            value = value[1:].strip()
        percent = value.endswith("%")
        value = value.removesuffix("%").strip()
        if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?", value):
            return math.nan
        value = float(value.replace(",", "")) / (100 if percent else 1)
    try:
        result = float(value)
        return result if math.isfinite(result) else math.nan
    except (ValueError, TypeError, OverflowError):
        return math.nan


def layout(ws):
    """流式扫描坐标及隐藏/合并信息，避免把预览的 100 行误作分析上限。"""
    result = {"rows": 1, "columns": 1, "hiddenRows": set(), "hiddenCols": set(), "merges": []}
    sheet_data = None
    with ws.parent._archive.open(ws._worksheet_path) as stream:
        for event, node in ET.iterparse(stream, events=("start", "end")):
            tag = node.tag.rsplit("}", 1)[-1]
            if event == "start":
                if tag == "sheetData":
                    sheet_data = node
                continue
            if tag == "row":
                row = int(node.get("r", "1"))
                result["rows"] = max(row, result["rows"])
                if node.get("hidden") in ("1", "true") or node.get("ht") == "0":
                    result["hiddenRows"].add(row)
                for cell in node:
                    match = re.match("[A-Z]+", cell.get("r", ""))
                    if match:
                        result["columns"] = max(result["columns"], column_index_from_string(match[0]))
                node.clear()
                if sheet_data is not None:
                    sheet_data.clear()
            elif tag == "col":
                if node.get("hidden") in ("1", "true") or float(node.get("width", "8")) < 0.5:
                    result["hiddenCols"].update(range(int(node.get("min")), min(16384, int(node.get("max"))) + 1))
                node.clear()
            elif tag == "mergeCell":
                if len(result["merges"]) < 20_000:
                    result["merges"].append(range_boundaries(node.get("ref")))
                node.clear()
    return result


class Reader:
    """一次请求复用工作簿，只读取当前数据源内的显式工作表；不运行宏或重算公式。"""

    def __init__(self, input_path, stack):
        from desktop.bridge import source_files

        _, files = source_files(input_path)
        self.available = {os.path.normcase(str(p.resolve())) for p in files}
        self.stack, self.books, self.layouts, self.hashes = stack, {}, {}, {}
        self.used = 0

    def sheet(self, filename, name):
        path = Path(text(filename, "文件路径", 3000)).resolve()
        key = os.path.normcase(str(path))
        if key not in self.available:
            raise ValueError("分析区域的文件不属于当前数据源，请重新选择文件或文件夹")
        if key not in self.books:
            if path.stat().st_size > 200 * 1024 * 1024:
                raise ValueError("分析文件超过 200 MB，请先分批处理")
            self.hashes[key] = sha256(path.read_bytes()).hexdigest()
            if path.suffix.lower() == ".xls":
                import xlrd

                book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
                self.stack.callback(book.release_resources)
            else:
                with ZipFile(path) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 200 * 1024 * 1024:
                        raise ValueError("工作簿解压后超过 200 MB，请先分批处理")
                book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
                self.stack.callback(book.close)
            self.books[key] = book
        book = self.books[key]
        sheet_key = key, name
        legacy = path.suffix.lower() == ".xls"
        if legacy:
            if name not in book.sheet_names():
                raise ValueError("工作表不存在，请重新读取")
            ws = book.sheet_by_name(name)
            if ws.visibility != 0:
                raise ValueError("统计分析不读取隐藏 Sheet，请在本机确认并取消隐藏后重试")
        else:
            if name not in book.sheetnames:
                raise ValueError("工作表不存在，请重新读取")
            ws = book[name]
            if ws.sheet_state != "visible" or not hasattr(ws, "iter_rows"):
                raise ValueError("请选择可见的数据工作表")
        if sheet_key not in self.layouts:
            self.layouts[sheet_key] = (
                {
                    "rows": ws.nrows,
                    "columns": ws.ncols,
                    "hiddenRows": {r + 1 for r, info in ws.rowinfo_map.items() if info.hidden},
                    "hiddenCols": {c + 1 for c, info in ws.colinfo_map.items() if info.hidden},
                    "merges": [(cl + 1, rl + 1, ch, rh) for rl, rh, cl, ch in ws.merged_cells],
                }
                if legacy
                else layout(ws)
            )
        return path, ws, self.layouts[sheet_key], legacy

    def values(self, filename, name, bounds):
        path, ws, meta, legacy = self.sheet(filename, name)
        left, top, right, bottom = bounds
        size = (right - left + 1) * (bottom - top + 1)
        self.used += size
        if self.used > MAX_CELLS:
            raise ValueError("一次分析最多读取 500,000 个区域单元格，请缩小范围或分批")
        rows, missing_formulas, errors = [], 0, 0
        if legacy:
            import xlrd

            for r in range(top - 1, bottom):
                values = []
                for c in range(left - 1, right):
                    cell = ws.cell(r, c) if r < ws.nrows and c < ws.ncols else None
                    value = cell.value if cell is not None else None
                    if cell is not None and cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(value, ws.book.datemode)
                    elif cell is not None and cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell is not None and cell.ctype == xlrd.XL_CELL_ERROR:
                        errors += 1
                        value = None
                    values.append(value)
                rows.append(values)
        else:
            for row in ws.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "e":
                        errors += 1
                        value = None
                    values.append(value)
                rows.append(values)
            # 单独扫描原 XML 判定无缓存公式，避免把缺失公式结果当成零。
            with ws.parent._archive.open(ws._worksheet_path) as stream:
                sheet_data = None
                for event, node in ET.iterparse(stream, events=("start", "end")):
                    tag = node.tag.rsplit("}", 1)[-1]
                    if event == "start":
                        if tag == "sheetData":
                            sheet_data = node
                        continue
                    if tag == "row":
                        r = int(node.get("r", "0"))
                        if top <= r <= bottom:
                            for cell in node:
                                ref = cell.get("r", "")
                                col = re.match("[A-Z]+", ref)
                                if (
                                    col
                                    and left <= column_index_from_string(col[0]) <= right
                                    and cell.find(NS + "f") is not None
                                ):
                                    v = cell.find(NS + "v")
                                    missing_formulas += int(v is None or (v.text is None and cell.get("t") != "str"))
                        node.clear()
                        if sheet_data is not None:
                            sheet_data.clear()
                        if r > bottom:
                            break
        return rows, meta, {"missingFormulaCaches": missing_formulas, "errorCells": errors}


def region(value, meta):
    """支持矩形、整列、整行；整列/行按实际坐标边界展开，不静默截成前 100 行。"""
    value = text(value, "区域", 80).upper().replace("$", "").replace(" ", "")
    if re.fullmatch("[A-Z]{1,3}:[A-Z]{1,3}", value):
        a, b = value.split(":")
        bounds = column_index_from_string(a), 1, column_index_from_string(b), max(1, meta["rows"])
    elif re.fullmatch("[1-9][0-9]*:[1-9][0-9]*", value):
        a, b = value.split(":")
        bounds = 1, int(a), max(1, meta["columns"]), int(b)
    elif re.fullmatch(r"[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?", value):
        bounds = range_boundaries(value)
    else:
        raise ValueError("区域格式应为 A6:K246、B:D 或 6:20；多个不连续区域请分别添加")
    left, top, right, bottom = bounds
    if left > right or top > bottom or right > 16384 or bottom > MAX_ROWS or right - left + 1 > MAX_FIELDS:
        raise ValueError("区域须正向选择，每块最多 256 列，最远第 100,000 行；整行/列过大时请改用具体矩形")
    return bounds


def grid(request):
    """分页选区视图，100 行以后的区域也能跳转选择。"""
    with ExitStack() as stack:
        reader = Reader(request["inputPath"], stack)
        _, _, meta, _ = reader.sheet(request["path"], request["sheet"])
        top = integer(request.get("startRow", 1), 1, MAX_ROWS, "预览起始行")
        left = integer(request.get("startColumn", 1), 1, 16384, "预览起始列")
        right, bottom = min(16384, left + 19), min(MAX_ROWS, top + 39)
        rows, _, _ = reader.values(request["path"], request["sheet"], (left, top, right, bottom))
        header = integer(request.get("headerRow", 1), 1, 1000, "检测到的表头行")
        first = min([header] + [b for a, b, c, d in meta["merges"] if b <= header <= d])
        if request.get("includeHidden") is not True:
            rows = [
                [
                    v if top + r not in meta["hiddenRows"] and left + c not in meta["hiddenCols"] else None
                    for c, v in enumerate(row)
                ]
                for r, row in enumerate(rows)
            ]
        return {
            "startRow": top,
            "startColumn": left,
            "totalRows": meta["rows"],
            "totalColumns": meta["columns"],
            "suggestedRange": f"A{first}:{get_column_letter(meta['columns'])}{max(header + 1, meta['rows'])}",
            "suggestedHeaderRows": header - first + 1,
            "columns": [get_column_letter(c) for c in range(left, right + 1)],
            "rows": [[str(plain(v) if v is not None else "")[:300] for v in row] for row in rows],
            "hiddenRows": sorted(r for r in meta["hiddenRows"] if top <= r <= bottom),
            "hiddenColumns": sorted(c for c in meta["hiddenCols"] if left <= c <= right),
        }


def read_block(reader, spec):
    """合并仅在表头区域展开，数据区不填充，防止把合并值重复计入。"""
    import pandas as pd

    identifier = text(spec.get("id"), "数据块标识", 80)
    if not re.fullmatch("[a-zA-Z0-9_-]+", identifier):
        raise ValueError("数据块标识无效")
    name = text(spec.get("name"), "数据块名称", 80)
    _, _, meta, _ = reader.sheet(spec["path"], spec["sheet"])
    bounds = region(spec.get("range"), meta)
    left, top, right, bottom = bounds
    header = integer(spec.get("headerRows", 1), 0, 10, "表头层数")
    orientation = spec.get("orientation", "columns")
    if orientation not in ("columns", "rows"):
        raise ValueError("请选择按列或按行排布")
    rows, meta, audit = reader.values(spec["path"], spec["sheet"], bounds)
    for a, b, c, d in meta["merges"]:
        if left <= a <= right and top <= b <= bottom:
            for r in range(max(top, b), min(bottom, d) + 1):
                for col in range(max(left, a), min(right, c) + 1):
                    is_header = r < top + header if orientation == "columns" else col < left + header
                    if is_header:
                        rows[r - top][col - left] = rows[b - top][a - left]
    included_rows = [
        r for r in range(top, bottom + 1) if spec.get("includeHidden") is True or r not in meta["hiddenRows"]
    ]
    included_cols = [
        c for c in range(left, right + 1) if spec.get("includeHidden") is True or c not in meta["hiddenCols"]
    ]
    # 表头按原位置计算，隐藏字段/记录的过滤不能将首条数据误提升为表头。
    if orientation == "columns":
        raw = [[rows[r - top][c - left] for c in included_cols] for r in range(top, min(bottom + 1, top + header))]
        raw += [[rows[r - top][c - left] for c in included_cols] for r in included_rows if r >= top + header]
        locations = [get_column_letter(c) for c in included_cols]
    else:
        raw = [[rows[r - top][c - left] for r in included_rows] for c in range(left, min(right + 1, left + header))]
        raw += [[rows[r - top][c - left] for r in included_rows] for c in included_cols if c >= left + header]
        locations = [str(r) for r in included_rows]
    if len(locations) > MAX_FIELDS or not locations or len(raw) <= header:
        raise ValueError(f"「{name}」没有数据记录，或转置后超过 256 个字段，请检查表头层数和区域")
    labels = []
    for index, location in enumerate(locations):
        parts = [
            str(plain(row[index])).strip() for row in raw[:header] if row[index] is not None and str(row[index]).strip()
        ]
        labels.append(" / ".join(dict.fromkeys(parts)) or f"字段 {location}")
    body, empty_rows, totals = [], 0, 0
    for row in raw[header:]:
        if all(v is None or v == "" for v in row):
            empty_rows += 1
            continue
        if spec.get("skipTotals", True) and any(
            str(v).strip() in ("合计", "小计", "总计", "平均", "平均值") for v in row[:2] if v is not None
        ):
            totals += 1
            continue
        body.append(row)
    if not body:
        raise ValueError(f"「{name}」清理后没有数据记录")
    frame = pd.DataFrame(body, columns=[f"{identifier}:f{i}" for i in range(len(locations))], dtype=object)
    fields = []
    for index, col in enumerate(frame):
        values = [v for v in frame[col] if v is not None and v != ""]
        numbers = sum(math.isfinite(numeric(v)) for v in values)
        dates = sum(isinstance(v, (datetime, date)) for v in values)
        kind = (
            "number"
            if numbers and numbers == len(values)
            else "date"
            if dates == len(values) and dates
            else "mixed"
            if numbers
            else "text"
        )
        fields.append({"id": col, "name": labels[index], "location": locations[index], "kind": kind})
    resolved = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
    audit.update(
        emptyRows=empty_rows,
        totalsSkipped=totals,
        hiddenRows=bottom - top + 1 - len(included_rows),
        hiddenColumns=right - left + 1 - len(included_cols),
    )
    warnings = []
    if audit["missingFormulaCaches"]:
        warnings.append(
            f"{name}：{audit['missingFormulaCaches']} 个公式缺少已保存结果，按缺失值处理；可在 Excel 重算并保存后刷新"
        )
    if audit["errorCells"]:
        warnings.append(f"{name}：{audit['errorCells']} 个错误单元格按缺失值处理")
    if totals:
        warnings.append(f"{name}：排除 {totals} 条前两字段为合计/小计/总计/平均的记录")
    if audit["hiddenRows"] or audit["hiddenColumns"]:
        warnings.append(f"{name}：排除 {audit['hiddenRows']} 行、{audit['hiddenColumns']} 列隐藏内容")
    return frame, {
        "id": identifier,
        "name": name,
        "fields": fields,
        "rows": len(frame),
        "resolvedRange": resolved,
        "sample": [[plain(v) for v in row] for row in frame.head(6).values.tolist()],
        "audit": audit,
        "warnings": warnings,
    }


def read_blocks(request):
    specs = request.get("blocks")
    if not isinstance(specs, list) or not 1 <= len(specs) <= 4 or any(not isinstance(b, dict) for b in specs):
        raise ValueError("请添加 1～4 个数据区域")
    if len({b.get("id") for b in specs}) != len(specs):
        raise ValueError("数据块标识不能重复")
    with ExitStack() as stack:
        reader = Reader(request["inputPath"], stack)
        frames, blocks = [], []
        for spec in specs:
            frame, block = read_block(reader, spec)
            frames.append(frame)
            blocks.append(block)
        token = sha256(
            json.dumps({"files": reader.hashes, "blocks": specs}, ensure_ascii=False, sort_keys=True).encode()
        ).hexdigest()
    return frames, {"token": token, "blocks": blocks}


def inspect(request):
    return read_blocks(request)[1]


def combine(frames, blocks, request):
    """追加按字段名对齐；关联只接受一对一唯一键，拒绝笛卡尔积放大金额。"""
    import pandas as pd

    mode = request.get("combine", "single")
    warnings = []
    if mode == "single":
        if len(frames) != 1:
            raise ValueError("单区域模式只能有一个数据块，请选择追加或关联")
        return frames[0], blocks[0]["fields"], warnings
    if mode == "append":
        labels = [f["name"] for f in blocks[0]["fields"]]
        if len(set(labels)) != len(labels):
            raise ValueError("追加需要唯一字段名；当前存在重复表头，请调整多层表头或选区")
        all_frames = [frames[0]]
        for frame, block in zip(frames[1:], blocks[1:]):
            names = [f["name"] for f in block["fields"]]
            if len(names) != len(set(names)) or set(names) != set(labels):
                raise ValueError("追加数据块的字段名必须一致，可调整区域和表头层数；不同指标请使用按字段关联")
            cols = [next(f["id"] for f in block["fields"] if f["name"] == label) for label in labels]
            copied = frame[cols].copy()
            copied.columns = frames[0].columns
            all_frames.append(copied)
        return pd.concat(all_frames, ignore_index=True), blocks[0]["fields"], warnings
    if mode != "join" or len(frames) < 2:
        raise ValueError("请选择单区域、追加同类记录或按字段关联")
    keys = request.get("joinKeys", {})
    join_type = request.get("joinType", "inner")
    if join_type not in ("inner", "left"):
        raise ValueError("请选择交集或保留首表的关联方式")
    prepared, sets = [], []
    for frame, block in zip(frames, blocks):
        key = keys.get(block["id"])
        if key not in frame:
            raise ValueError(f"请为「{block['name']}」选择关联字段")

        # 类型参与键匹配，保留文本编号前导零，不把文本 001 和数字 1 合并。
        def canonical(v):
            if v is None or v == "":
                return None
            if isinstance(v, str) and not v.strip():
                return None
            if isinstance(v, (datetime, date)):
                return ("date", plain(v))
            if isinstance(v, (float, int)) and not isinstance(v, bool):
                return ("number", v) if math.isfinite(v) else None
            return ("text", str(v).strip())

        copied = frame.copy()
        copied["__key"] = frame[key].map(canonical)
        nulls = int(copied["__key"].isna().sum())
        copied = copied[copied["__key"].notna()]
        if copied["__key"].duplicated().any():
            raise ValueError(f"「{block['name']}」关联字段存在重复值；请先汇总为每键一行，避免重复计算")
        if nulls:
            warnings.append(f"{block['name']}：{nulls} 条空关联键被排除")
        sets.append(set(copied["__key"]))
        prepared.append(copied)
    result = prepared[0]
    for i, frame in enumerate(prepared[1:], 1):
        current = set(result["__key"])
        missing = len(current - sets[i])
        other = len(sets[i] - current)
        warnings.append(
            f"关联 {blocks[i]['name']}：当前主表未匹配 {missing} 个键，对方独有 {other} 个键；数值键和文本键分开匹配"
        )
        result = result.merge(frame, on="__key", how=join_type, validate="one_to_one", sort=False)
    if result.empty:
        raise ValueError("关联后没有匹配记录，请核对键字段及数字/文本类型")
    return (
        result.drop(columns="__key"),
        [{**f, "name": b["name"] + " / " + f["name"]} for b in blocks for f in b["fields"]],
        warnings,
    )


def summary(frame, fields):
    """描述统计使用有效数值；标准差采用样本口径 ddof=1，缺失值不补零。"""
    rows = []
    for field in fields:
        values = frame[field["id"]].map(numeric)
        valid = values.dropna()
        blanks = sum(v is None or v == "" or (isinstance(v, float) and math.isnan(v)) for v in frame[field["id"]])
        rows.append(
            [
                field["name"],
                len(valid),
                blanks,
                len(frame) - len(valid) - blanks,
                valid.sum() if len(valid) else None,
                valid.mean(),
                valid.median(),
                valid.min(),
                valid.max(),
                valid.std(ddof=1),
                valid.quantile(0.25),
                valid.quantile(0.75),
            ]
        )
    return {
        "columns": [
            "字段",
            "有效数值",
            "缺失",
            "非数值",
            "合计",
            "平均值",
            "中位数",
            "最小值",
            "最大值",
            "样本标准差",
            "25%分位",
            "75%分位",
        ],
        "rows": [[plain(v) for v in row] for row in rows],
    }


def render(request):
    """计算与绘图均在 Python；返回真实口径、来源和清理数量供界面解释。"""
    import numpy as np
    import pandas as pd

    kind = request.get("kind")
    if kind not in KINDS:
        raise ValueError("请选择支持的统计图类型")
    title = text(request.get("title"), "图表标题", 120)
    frames, schema = read_blocks(request)
    if request.get("token") != schema["token"]:
        raise ValueError("数据文件或选区已改变，请重新读取数据后生成")
    frame, fields, warnings = combine(frames, schema["blocks"], request)
    for block in schema["blocks"]:
        warnings.extend(block["warnings"])
    names = {f["id"]: f["name"] for f in fields}
    expected = request.get("fieldLabels", {})
    if expected and any(names.get(key) != value for key, value in expected.items()):
        raise ValueError("原表字段已改变，请编辑图表并重新选择字段")
    y = request.get("y", [])
    x = request.get("x", "")
    if not isinstance(y, list) or not 1 <= len(y) <= 6 or len(set(y)) != len(y) or any(c not in names for c in y):
        raise ValueError("请选择 1～6 个数值字段")
    if kind in ("bar", "line", "pie", "scatter") and x not in names:
        raise ValueError("请选择分类、时间或 X 轴字段")
    if kind in ("pie", "histogram", "scatter", "metric") and len(y) != 1:
        raise ValueError("这类图表需要且只接受一个数值字段")
    if kind == "correlation" and len(y) < 2:
        raise ValueError("相关性热力图至少需要两个数值字段")
    agg = request.get("aggregation", "sum")
    if agg not in AGGS:
        raise ValueError("请选择有效的汇总方式")
    unit = request.get("unit", "")
    note = request.get("note", "")
    if not isinstance(unit, str) or len(unit) > 30 or not isinstance(note, str) or len(note) > 1000:
        raise ValueError("单位或注解过长")
    selected_fields = [next(f for f in fields if f["id"] == c) for c in y]
    numeric_frame = frame[y].apply(lambda col: col.map(numeric)).astype(float)
    if any(numeric_frame[c].notna().sum() == 0 for c in y):
        raise ValueError("某个数值字段没有有效数字；请检查表头、选区或公式缓存")
    statistics = summary(frame, selected_fields)
    invalid = sum(row[3] for row in statistics["rows"])
    if invalid:
        warnings.append(f"{invalid} 个非数值项未参与数值统计；未将缺失值或文本当成零")
    table, metric, image = {"columns": [], "rows": []}, None, None
    method, counts = "", {"inputRows": len(frame)}
    if kind != "metric":
        import matplotlib

        matplotlib.use("Agg", force=True)
        from matplotlib import pyplot as plt, font_manager

        installed = {font.name for font in font_manager.fontManager.ttflist}
        plt.rcParams["font.sans-serif"] = [
            f for f in ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC") if f in installed
        ] + ["DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        fig, ax = plt.subplots(figsize=(8, 4.5), dpi=130)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=9, colors="#52655c")
        ax.set_title(title, loc="left", fontsize=14, pad=20, color="#254a39")
    try:
        if kind == "metric":
            s = numeric_frame[y[0]].dropna()
            metric = float(len(s)) if agg == "count" else float(getattr(s, agg)())
            table = {
                "columns": ["指标", "统计方式", "数值", "有效数值数"],
                "rows": [[names[y[0]], AGGS[agg], metric, len(s)]],
            }
            method = f"对 {names[y[0]]} 的全部有效数值计算{AGGS[agg]}；缺失和非数值项不参与。"
        elif kind in ("bar", "line", "pie"):
            category = frame[x].map(
                lambda v: None if v is None or v == "" or (isinstance(v, float) and math.isnan(v)) else plain(v)
            )
            bucket = request.get("timeBucket", "none")
            if bucket not in ("none", "day", "month", "year"):
                raise ValueError("时间粒度无效")
            if kind == "line" and bucket != "none":
                if any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in category.dropna()):
                    raise ValueError(
                        "时间字段含普通数字，无法推断其日期含义；请选择 Excel 日期或日期文字，或取消日期分组"
                    )
                parsed = pd.to_datetime(category, errors="coerce", format="mixed")
                if parsed.notna().sum() == 0:
                    raise ValueError("时间字段无法解析为日期，请选择 Excel 日期或日期文字字段")
                category = parsed.dt.strftime({"day": "%Y-%m-%d", "month": "%Y-%m", "year": "%Y"}[bucket])
                warnings.append(
                    f"按{ {'day': '日', 'month': '月', 'year': '年'}[bucket] }汇总；{int(parsed.isna().sum())} 条无效日期被排除"
                )
            work = numeric_frame.copy()
            work["__category"] = category
            missing = int(category.isna().sum())
            if missing:
                warnings.append(f"{missing} 条分类/时间为空的记录未进入图表；下方描述统计仍针对全部输入数值")
            groups = work.dropna(subset=["__category"]).groupby("__category", sort=False, dropna=True)[y]
            values = groups.sum(min_count=1) if agg == "sum" else getattr(groups, agg)()
            if values.empty:
                raise ValueError("没有有效分类或时间记录")
            order = request.get("sort", "source")
            if order == "label" or (kind == "line" and bucket != "none"):
                values = values.iloc[sorted(range(len(values)), key=lambda i: str(values.index[i]))]
            elif order == "value_desc":
                values = values.sort_values(y[0], ascending=False, kind="stable")
            elif order != "source":
                raise ValueError("排序方式无效")
            groups_total = len(values)
            limit = integer(request.get("limit", 20), 1, 60, "图上分类数")
            if kind == "pie":
                if agg not in ("sum", "count"):
                    raise ValueError("饼图仅支持合计或有效值计数，平均值不能作为整体占比")
                if (values[y[0]].dropna() < 0).any() or values[y[0]].sum() <= 0:
                    raise ValueError("饼图需要非负数值且总值大于零；含负数请改用柱形图")
                if values[y[0]].isna().any():
                    raise ValueError("部分饼图分类没有有效数值，请缩小范围或使用有效值计数")
                if len(values) > limit:
                    raise ValueError("饼图分类数超过图上分类数，请提高上限或缩小区域，避免遗漏分类导致占比失真")
            elif len(values) > limit:
                warnings.append(
                    f"共有 {len(values)} 个分类，图表和图表数据展示当前排序的前 {limit} 个；描述统计包含全部输入数值"
                )
                values = values.head(limit)
            labels = [str(v) for v in values.index]
            table = {
                "columns": [names[x]] + [names[c] for c in y],
                "rows": [[labels[i]] + [plain(v) for v in row] for i, row in enumerate(values.values)],
            }
            counts.update(categories=groups_total, displayedCategories=len(values))
            if kind == "pie":
                wedges, _, _ = ax.pie(
                    values[y[0]],
                    autopct="%1.1f%%",
                    colors=COLORS,
                    startangle=90,
                    pctdistance=0.78,
                    wedgeprops={"width": 0.46, "edgecolor": "white"},
                )
                ax.legend(wedges, labels, loc="center left", bbox_to_anchor=(1, 0.5), fontsize=8, frameon=False)
            else:
                positions = np.arange(len(values))
                for i, col in enumerate(y):
                    if kind == "bar":
                        width = 0.78 / len(y)
                        ax.bar(
                            positions + (i - (len(y) - 1) / 2) * width,
                            values[col],
                            width=width,
                            color=COLORS[i],
                            label=names[col],
                        )
                    else:
                        ax.plot(
                            positions,
                            values[col],
                            marker="o",
                            markersize=3,
                            linewidth=2,
                            color=COLORS[i],
                            label=names[col],
                        )
                ax.set_xticks(
                    positions,
                    labels,
                    rotation=35 if len(labels) > 6 else 0,
                    ha="right" if len(labels) > 6 else "center",
                )
                ax.set_ylabel(unit or AGGS[agg])
                ax.grid(axis="y", alpha=0.15)
                ax.set_axisbelow(True)
                if len(y) > 1:
                    ax.legend(fontsize=8, frameon=False)
            method = f"按 {names[x]} 分组，对所选数值计算{AGGS[agg]}。同名分类合并，缺失数值不补零。"
        elif kind == "histogram":
            bins = integer(request.get("bins", 12), 3, 50, "直方图组数")
            s = numeric_frame[y[0]].dropna()
            freq, edges, _ = ax.hist(s, bins=bins, color=COLORS[0], edgecolor="white")
            ax.set_xlabel(names[y[0]] + (f"（{unit}）" if unit else ""))
            ax.set_ylabel("频数")
            table = {
                "columns": ["区间下限", "区间上限", "频数"],
                "rows": [[float(edges[i]), float(edges[i + 1]), int(v)] for i, v in enumerate(freq)],
            }
            counts["validRows"] = len(s)
            method = f"全部 {len(s)} 个有效数值分为 {bins} 个等宽区间。左闭右开，最后一组包含最大值；纵轴是频数。"
        elif kind == "scatter":
            pairs = pd.DataFrame({"x": frame[x].map(numeric), "y": numeric_frame[y[0]]}).dropna()
            if len(pairs) < 2:
                raise ValueError("散点图至少需要 2 条 X、Y 均有效的配对记录")
            points = pairs.iloc[np.linspace(0, len(pairs) - 1, min(5000, len(pairs)), dtype=int)]
            ax.scatter(points.x, points.y, s=18, color=COLORS[0], alpha=0.65)
            ax.set_xlabel(names[x])
            ax.set_ylabel(names[y[0]] + (f"（{unit}）" if unit else ""))
            corr = (
                pairs.x.corr(pairs.y)
                if len(pairs) >= 3 and pairs.x.nunique() > 1 and pairs.y.nunique() > 1
                else math.nan
            )
            counts.update(validPairs=len(pairs), displayedPoints=len(points))
            table = {
                "columns": ["有效配对", "缺失配对", "Pearson r"],
                "rows": [[len(pairs), len(frame) - len(pairs), plain(corr)]],
            }
            method = "每个点表示同一条记录的 X/Y 数值。Pearson 相关性基于全部有效配对（至少 3 对且两列有变化），相关不代表因果。"
            if len(points) < len(pairs):
                warnings.append("散点均匀抽取最多 5,000 个点用于显示，统计仍基于全部有效配对")
        elif kind == "box":
            ax.boxplot(
                [numeric_frame[c].dropna() for c in y],
                tick_labels=[names[c] for c in y],
                patch_artist=True,
                boxprops={"facecolor": "#c8e7da", "edgecolor": COLORS[0]},
                medianprops={"color": COLORS[0]},
                whis=1.5,
            )
            ax.set_ylabel(unit or "数值")
            ax.tick_params(axis="x", rotation=15)
            table = statistics
            method = "每个字段单独计算：箱体为第 25%～75% 分位数，中线为中位数，须延伸到 1.5×四分位距内最远观测；外侧点为候选异常值，不自动删除。"
        else:
            corr = numeric_frame.corr(method="pearson", min_periods=3)
            valid = numeric_frame.notna().astype(int)
            pairs = valid.T.dot(valid)
            plot = ax.imshow(np.ma.masked_invalid(corr.values), vmin=-1, vmax=1, cmap="BrBG")
            fig.colorbar(plot, ax=ax, fraction=0.04, pad=0.04)
            labels = [names[c] for c in y]
            ax.set_xticks(range(len(y)), labels, rotation=30, ha="right")
            ax.set_yticks(range(len(y)), labels)
            for i in range(len(y)):
                for j in range(len(y)):
                    v = corr.iloc[i, j]
                    ax.text(
                        j,
                        i,
                        "—" if pd.isna(v) else f"{v:.2f}",
                        ha="center",
                        va="center",
                        color="white" if pd.notna(v) and abs(v) > 0.65 else "#304e43",
                    )
            table = {
                "columns": ["字段一", "字段二", "Pearson r", "有效配对数"],
                "rows": [
                    [labels[i], labels[j], plain(corr.iloc[i, j]), int(pairs.iloc[i, j])]
                    for i in range(len(y))
                    for j in range(i, len(y))
                ],
            }
            method = "Pearson 相关系数在 -1～1 之间。每对字段分别剔除缺失值，至少需要 3 对有效观测；常量字段或不足配对显示“—”。相关不代表因果。"
        if kind != "metric":
            fig.tight_layout()
            stream = BytesIO()
            fig.savefig(stream, format="png", facecolor="white", bbox_inches="tight")
            image = "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode()
    finally:
        if kind != "metric":
            plt.close(fig)
    table["rows"] = [[plain(v) for v in row] for row in table["rows"]]
    if kind == "metric" and plain(metric) is None:
        raise ValueError("统计结果超出可表示的数值范围，请缩小范围或检查原始数值")
    return {
        "title": title,
        "kind": kind,
        "image": image,
        "metric": plain(metric),
        "unit": unit,
        "note": note,
        "method": method,
        "statistics": statistics,
        "table": table,
        "counts": counts,
        "warnings": warnings,
        "sources": [
            {
                "name": s["name"],
                "file": Path(s["path"]).name,
                "sheet": s["sheet"],
                "range": b["resolvedRange"],
                "headerRows": s["headerRows"],
                "orientation": s["orientation"],
                "rows": b["rows"],
            }
            for s, b in zip(request["blocks"], schema["blocks"])
        ],
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "token": schema["token"],
    }


def csv_text(table):
    stream = StringIO(newline="")
    writer = csv.writer(stream)

    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    writer.writerow([safe(v) for v in table["columns"]])
    writer.writerows([[safe(v) for v in row] for row in table["rows"]])
    return "\ufeff" + stream.getvalue()


def html_table(table):
    head = "".join("<th>" + escape(str(c)) + "</th>" for c in table["columns"])
    rows = "".join(
        "<tr>" + "".join("<td>" + escape("—" if v is None else str(v)) + "</td>" for v in row) + "</tr>"
        for row in table["rows"]
    )
    return '<div class="scroll"><table><thead><tr>' + head + "</tr></thead><tbody>" + rows + "</tbody></table></div>"


def export_dashboard(request):
    """重新核对源文件后原子写入报告，导出真实统计结果而非脱敏资料。"""
    title = text(request.get("title"), "看板标题", 120)
    cards = request.get("cards")
    if not isinstance(cards, list) or not 1 <= len(cards) <= 12:
        raise ValueError("一个 dashboard 可包含 1～12 个图表")
    root = Path(request["outputPath"]).resolve(strict=True)
    if not root.is_dir():
        raise ValueError("请选择报告保存目录")
    output = root / f"Dashboard_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex[:6]}"
    results = [render(card["config"]) for card in cards]
    temporary = root / ("." + output.name)
    temporary.mkdir()
    sections = []
    try:
        for index, result in enumerate(results, 1):
            if result["image"]:
                (temporary / f"chart_{index:02d}.png").write_bytes(base64.b64decode(result["image"].split(",", 1)[1]))
            (temporary / f"chart_{index:02d}_data.csv").write_text(csv_text(result["table"]), encoding="utf-8")
            (temporary / f"chart_{index:02d}_statistics.csv").write_text(
                csv_text(result["statistics"]), encoding="utf-8"
            )
            visual = (
                '<img src="' + result["image"] + '" alt="' + escape(result["title"], quote=True) + '">'
                if result["image"]
                else '<div class="metric">'
                + escape(str(result["metric"]))
                + " <small>"
                + escape(result["unit"])
                + "</small></div>"
            )
            warnings = "".join("<li>" + escape(w) + "</li>" for w in result["warnings"])
            source = "；".join(
                f"{s['file']} / {s['sheet']}!{s['range']}（{s['rows']} 条记录）" for s in result["sources"]
            )
            sections.append(
                "<section><h2>"
                + escape(result["title"])
                + "</h2>"
                + visual
                + "<p>"
                + escape(result["method"])
                + "</p><p>"
                + escape(result["note"])
                + '</p><p class="source">'
                + escape(source)
                + "</p><ul>"
                + warnings
                + "</ul><details><summary>图表数据与描述统计</summary>"
                + html_table(result["table"])
                + html_table(result["statistics"])
                + "</details></section>"
            )
        document = (
            """<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; img-src data:; style-src 'unsafe-inline'"><title>"""
            + escape(title)
            + """</title><style>body{margin:0;padding:32px;background:#f2f6f3;color:#284c3a;font:14px 'Microsoft YaHei',sans-serif}header{max-width:1200px;margin:auto}main{max-width:1200px;margin:24px auto;display:grid;grid-template-columns:repeat(auto-fit,minmax(440px,1fr));gap:20px}section{padding:22px;background:white;border:1px solid #d7e3db;border-radius:14px;min-width:0}h1{font-size:30px}h2{font-size:18px}img{width:100%}.metric{font-size:48px;padding:28px 0;color:#188263}small{font-size:16px}.source,li{font-size:12px;color:#6b7a70}p{line-height:1.8}.scroll{overflow:auto}table{border-collapse:collapse;white-space:nowrap;font-size:12px;margin:14px 0}th,td{padding:7px;border-bottom:1px solid #e0e8e2}summary{cursor:pointer}@media(max-width:600px){body{padding:12px}main{display:block}section{margin-bottom:16px}}@media print{body{background:white;padding:0}main{display:block}section{break-inside:avoid}details{display:block}}</style><header><p>EASYEXCEL · 本地统计分析</p><h1>"""
            + escape(title)
            + "</h1><p>此报告包含真实统计结果。公式使用工作簿已保存的缓存，未进行重算；数据来源与统计口径见每张图。</p></header><main>"
            + "".join(sections)
            + "</main></html>"
        )
        (temporary / "dashboard.html").write_text(document, encoding="utf-8")
        (temporary / "dashboard-plan.json").write_text(
            json.dumps({"version": 1, "title": title, "cards": cards}, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            shutil.rmtree(temporary)
    return {"output": str(output), "charts": len(cards)}
