"""Electron 的本地 JSON Lines 入口；所有 Excel 读写均调用 Python 核心。

Copyright (c) 2026 AbeLin · MIT License
"""
import json
import os
from pathlib import Path
import re
import sys
import threading
from itertools import islice

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import splitter  # noqa: E402
from core.xlio import _head_limit, col_names, resolve_header_row  # noqa: E402
from desktop.preview import read_preview  # noqa: E402
import openpyxl  # noqa: E402

EXTENSIONS = {".xlsx", ".xls"}
SKIP_DIRS = {"拆分结果", "汇总结果", "node_modules", ".git", ".venv"}
PROTOCOL_OUT = sys.stdout


def emit(event):
    """一行一个 UTF-8 JSON，避免普通日志污染进程协议。"""
    print(json.dumps(event, ensure_ascii=False, allow_nan=False), file=PROTOCOL_OUT, flush=True)


def header_config(row):
    """限制手动表头范围，避免错误输入导致无界扫描。"""
    if row is not None and (type(row) is not int or not 1 <= row <= 1000):
        raise ValueError("表头行号请输入 1 到 1000 的整数")
    return {"header_mode": "row" if row is not None else "auto", "header_row": row or 1}


def source_files(source):
    """枚举真实输入，跳过 Excel 锁文件、临时表和历史结果目录。"""
    path = Path(source).expanduser().resolve(strict=True)
    if path.is_file():
        if path.suffix.lower() not in EXTENSIONS or path.name.startswith("~$"):
            raise ValueError("请选择 .xlsx 或 .xls 文件")
        return path, [path]
    files = []
    for root, dirs, names in os.walk(path):
        dirs[:] = sorted(d for d in dirs if d not in SKIP_DIRS and not re.fullmatch(r"\d{8}结果(?:_\d+)?|Python_\d{8}_\d{6}_[0-9a-f]{6}", d))
        files.extend(Path(root) / name for name in sorted(names)
                     if Path(name).suffix.lower() in EXTENSIONS and not name.startswith("~$")
                     and not name.lower().endswith("__tmp__.xlsx"))
    if not files:
        raise ValueError("这个文件夹里没有可读取的 Excel 文件")
    return path, files


def sheet_info(name, state, rows, total, config):
    """统一新旧 Excel 格式的工作表元数据；行数仅来自文件声明的尺寸。"""
    h = resolve_header_row(rows, config)
    columns = col_names(rows, h) if 0 < h <= len(rows) else []
    selectable = state == "visible" and any(columns)
    reason = "隐藏工作表" if state != "visible" else ("" if selectable else "未识别到表头，可手动指定表头行")
    return {"name": name, "state": state, "headerRow": h, "columns": columns,
            "estimatedRows": max(0, total - h) if total is not None and h > 0 else None,
            "selectable": selectable, "reason": reason}


def inspect_book(path, config):
    """只扫描表头附近，不在界面打开文件时加载整本工作簿。"""
    limit = _head_limit(config)
    sheets = []
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for index in range(book.nsheets):
                sh = book.sheet_by_index(index)
                state = "visible" if sh.visibility == 0 else "hidden"
                rows = [sh.row_values(r) for r in range(min(sh.nrows, limit))] if state == "visible" else []
                sheets.append(sheet_info(sh.name, state, rows, sh.nrows, config))
                book.unload_sheet(index)
        finally:
            book.release_resources()
    else:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in book.worksheets:
                total = ws.max_row
                ws.reset_dimensions()
                rows = list(islice(ws.iter_rows(values_only=True), limit)) if ws.sheet_state == "visible" else []
                sheets.append(sheet_info(ws.title, ws.sheet_state, rows, total, config))
        finally:
            book.close()
    return {"path": str(path), "name": path.name, "size": path.stat().st_size, "sheets": sheets}


def inspect_source(request):
    """目录里单个文件损坏时仍展示其余文件，并把失败原因交给界面。"""
    path, files = source_files(request["path"])
    config = header_config(request.get("headerRow"))
    books = []
    for file in files:
        try:
            books.append(inspect_book(file, config))
        except Exception as exc:
            books.append({"path": str(file), "name": file.name, "size": file.stat().st_size,
                          "sheets": [], "error": friendly_error(exc)})
    return {"path": str(path), "name": path.name, "isDirectory": path.is_dir(), "workbooks": books,
            "suggestedOutput": str((path if path.is_dir() else path.parent) / "拆分结果")}


def preview_sheet(request):
    """预览原表第 1～100 行（含标题和表头），样式及格式化均由 Python 读取。"""
    path, _ = source_files(request["path"])
    if path.is_dir():
        raise ValueError("预览需要选择一个 Excel 文件")
    header_config(request.get("headerRow"))
    return read_preview(path, request["sheet"])


def selection_config(request):
    """在 Python 边界再次校验文件和 Sheet，空选择绝不回退成处理全部。"""
    path, files = source_files(request["inputPath"])
    available = {os.path.normcase(str(f)): f for f in files}
    choices = request.get("selection")
    if not isinstance(choices, list) or not choices:
        raise ValueError("请至少选择一个工作表")
    config = header_config(request.get("headerRow"))
    selected = {}
    for choice in choices:
        key = os.path.normcase(str(Path(choice["path"]).resolve()))
        if key not in available:
            raise ValueError("所选文件不在当前数据源中，请重新选择")
        names = choice.get("sheets")
        if not isinstance(names, list) or any(not isinstance(n, str) for n in names):
            raise ValueError("工作表选择格式无效")
        if not names:
            continue
        actual = {s["name"]: s for s in inspect_book(available[key], config)["sheets"]}
        for name in names:
            if name not in actual or not actual[name]["selectable"]:
                raise ValueError(f"工作表「{name}」不存在、被隐藏或没有可识别表头，请重新读取")
        selected[key] = list(dict.fromkeys(selected.get(key, []) + names))
    if not selected:
        raise ValueError("请至少选择一个可处理的工作表")
    return {**config, "input_path": str(path), "sheet_selection": selected,
            "selected_input_files": [str(available[key]) for key in selected],
            "skip_values": ["合计", "小计", "总计", "平均", ""], "keep_hidden": False}


def groups(request):
    """分组取值完全由 Python 计算，与实际拆分的筛选口径一致。"""
    config = selection_config(request)
    column = request.get("column", "")
    if not isinstance(column, str) or not column.strip():
        raise ValueError("请选择拆分字段")
    values, missing = set(), []
    for path in config["selected_input_files"]:
        for name, _h, columns in splitter.sheet_headers(path, config):
            if splitter._find_col(columns, column) is None:
                missing.append(f"{Path(path).name} / {name}")
        values.update(splitter.list_values(path, config, column))
    return {"values": sorted(values), "missing": missing}


def run_job(request, cancelled, report):
    """调用现有引擎；UI 参数白名单映射，不允许覆盖内部配置。"""
    config = selection_config(request)
    mode = request.get("mode")
    if mode not in ("split", "merge"):
        raise ValueError("不支持的处理模式")
    output = request.get("outputPath")
    if not isinstance(output, str) or not output.strip():
        raise ValueError("请选择输出目录")
    out = Path(output).expanduser().resolve()
    source = Path(config["input_path"])
    if source.is_dir() and (out == source or out in source.parents):
        raise ValueError("输出目录不能是输入目录本身或它的上级目录，请另选一个子目录")
    if out.exists() and not out.is_dir():
        raise ValueError("输出位置必须是文件夹")
    config["output_path"] = str(out)
    for ui, key in (("preserveFormat", "preserve_format"), ("keepFormulas", "keep_formulas"),
                    ("makeZip", "make_zip"), ("mergeAcrossFiles", "merge_across_files")):
        if type(request.get(ui)) is not bool:
            raise ValueError("处理设置格式无效")
        config[key] = request[ui]
    config["keep_formulas"] = config["keep_formulas"] and config["preserve_format"]
    if type(request.get("bundleSheets", False)) is not bool:
        raise ValueError("总工作簿设置格式无效")
    config["bundle_sheets"] = request.get("bundleSheets", False)
    merge_mode = request.get("mergeMode", "rows")
    if merge_mode not in ("rows", "sheets"):
        raise ValueError("不支持的汇总方式")
    if mode == "split":
        group_data = groups(request)
        if group_data["missing"]:
            raise ValueError("以下工作表缺少拆分字段，请取消勾选或更换字段：" + "、".join(group_data["missing"]))
        if not group_data["values"]:
            raise ValueError("所选字段没有可拆分的数据")
        values = request.get("values")
        if values is not None and (not isinstance(values, list) or not values
                                   or any(v not in group_data["values"] for v in values)):
            raise ValueError("分组选择为空或已经变化，请重新选择分组")
        config["split_column"] = request["column"]
        config["selected_values"] = values or []
    if cancelled.is_set():
        return {"output": "", "summary": {}, "stopped": True}
    summary = {}

    def log(message):
        nonlocal summary
        if message.startswith("[SUMMARY] "):
            summary = json.loads(message[len("[SUMMARY] "):])
        else:
            report({"type": "log", "message": message})

    def progress(value):
        report({"type": "progress", "value": max(0, min(1, value))})

    if mode == "merge":
        if merge_mode == "sheets":
            from core.sheet_bundle import merge_as_sheets
            output = merge_as_sheets(config, log_fn=log, progress_fn=progress, stop_flag=cancelled.is_set)
        else:
            from core.collect import merge_tables
            output = merge_tables(config, log_fn=log, progress_fn=progress, stop_flag=cancelled.is_set)
    else:
        output = splitter.run_split(config, log_fn=log, progress_fn=progress, stop_flag=cancelled.is_set)
    stopped = bool(summary.get("stopped")) or cancelled.is_set()
    if not output and not stopped:
        raise ValueError("没有生成结果，请查看处理日志并检查所选工作表")
    return {"output": output or "", "summary": summary, "stopped": stopped}


def friendly_error(exc):
    """把常见文件错误转成人话，详细路径仅留在本地。"""
    if isinstance(exc, FileNotFoundError):
        return "文件已移动或不存在，请重新选择"
    if isinstance(exc, PermissionError):
        return "无法访问文件，请关闭占用它的 Excel，或选择有写入权限的目录"
    if isinstance(exc, ModuleNotFoundError):
        return f"Python 缺少依赖 {exc.name}，请安装 desktop/requirements.txt 中的依赖"
    if "not a zip file" in str(exc).lower():
        return "无法读取这个 Excel，文件可能损坏或设有打开口令"
    return str(exc) or "处理失败，请查看日志"


def main():
    """每个请求独立子进程；运行任务时 stdin 继续接收协作式停止信号。"""
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    try:
        request = json.loads(sys.stdin.readline())
        command = request.get("command")
        if command == "health":
            result = {"python": sys.executable, "version": sys.version.split()[0],
                      "frozen": bool(getattr(sys, "frozen", False))}
        elif command == "inspect":
            result = inspect_source(request)
        elif command == "preview":
            result = preview_sheet(request)
        elif command == "groups":
            result = groups(request)
        elif command == "code_environment":
            from desktop.code_runner import environment
            result = environment()
        elif command == "code":
            from desktop.code_runner import run_code
            result = run_code(request, selection_config(request), emit)
        elif command == "run":
            cancelled = threading.Event()

            def listen():
                for line in sys.stdin:
                    try:
                        if json.loads(line).get("cancel"):
                            cancelled.set()
                    except (ValueError, AttributeError):
                        continue

            threading.Thread(target=listen, daemon=True).start()
            result = run_job(request, cancelled, emit)
        else:
            raise ValueError("未知的本地处理请求")
        emit({"type": "result", "result": result})
    except Exception as exc:
        emit({"type": "error", "message": friendly_error(exc)})
        sys.exit(1)


if __name__ == "__main__":
    main()
