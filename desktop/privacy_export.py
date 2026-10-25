"""严格脱敏导出：从空白文件重建演示包，真实映射只写入应用本机目录。

Copyright (c) 2026 AbeLin · MIT License
"""
from contextlib import ExitStack
from copy import copy
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import TranslatorError
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.styles.fills import Stop
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from desktop.preview import bounded, color, metadata, themes
from desktop.privacy_formula import FormulaSanitizer, demo_parameter

KINDS = {'text': '文本', 'id': '编号', 'name': '姓名', 'category': '类别', 'integer': '整数',
         'money': '金额', 'number': '小数', 'percent': '比例', 'date': '日期', 'boolean': '布尔',
         'phone': '电话', 'email': '邮箱'}
MODES = {'synthetic', 'blank', 'schema'}
MAX_COLUMNS = 256
MAX_SHEETS = 32


def inferred_kind(label, cell):
    """仅在本机由字段语义/单元格格式推断类型，不输出真实值或分布。"""
    for pattern, kind in [(r'邮箱|email', 'email'), (r'电话|手机|phone', 'phone'), (r'姓名|联系人', 'name'),
                          (r'编号|编码|序号|账号|单号|\bid\b', 'id'), (r'日期|时间|date', 'date'),
                          (r'比例|占比|利率|完成率|税率', 'percent'), (r'金额|收入|成本|利润|价格|单价|余额|费用|销售额', 'money'),
                          (r'数量|人数|次数', 'integer'), (r'类别|类型|部门|区域|状态', 'category')]:
        if re.search(pattern, str(label), re.I):
            return kind
    if cell is not None:
        if cell.is_date:
            return 'date'
        if '%' in cell.number_format:
            return 'percent'
        if cell.data_type == 'b':
            return 'boolean'
        if cell.data_type in ('n', 'f') and cell.value is not None:
            return 'number'
    return 'text'


def _prepare(config, stack):
    """在本机建立列映射，只读取表头和固定的少量模板行。"""
    from desktop.bridge import inspect_book
    import pandas as pd
    tables = []
    for filename in config['selected_input_files']:
        path = Path(filename)
        if path.suffix.lower() != '.xlsx':
            raise ValueError('严格脱敏导出目前支持 .xlsx；请先在本机将旧 .xls 另存为 .xlsx，以保留公式')
        with ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError('工作簿解压后超过 200 MB，请先在本机拆成较小工作簿')
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        stack.callback(book.close)
        info = {s['name']: s for s in inspect_book(path, config)['sheets']}
        for name in config['sheet_selection'][os.path.normcase(str(path))]:
            header = info[name]['headerRow']
            if not 1 <= header <= 100:
                raise ValueError('脱敏模板目前支持前 100 行内的表头，请在运行设置中确认表头行')
            ws = book[name]
            layout = metadata(ws)
            names = list(pd.read_excel(path, sheet_name=name, header=header-1, nrows=0, engine='openpyxl').columns)
            width = len(names)
            if not 1 <= width <= MAX_COLUMNS:
                raise ValueError('脱敏导出每张表支持 1～256 列，请先调整表头或在本机拆表')
            hidden = set()
            for column in layout['columns']:
                # 与 code_runner 使用的 core.xlio.hidden_info 保持同一列选择规则。
                if column.get('hidden') in ('1', 'true') or (
                    column.get('customWidth') in ('1', 'true') and float(column.get('width', 8.43)) < 0.5
                ):
                    hidden.update(range(int(column['min']), int(column['max'])+1))
            visible = [c for c in range(1, width+1) if c not in hidden]
            if not visible:
                raise ValueError('没有可导出的可见字段')
            ws.reset_dimensions()
            # 固定读取 50 个模板行；不导出源总行数，也不按真实数据分布取样。
            rows = list(ws.iter_rows(min_row=1, max_row=header+50, max_col=width))
            fields = []
            for index, col in enumerate(visible):
                cell = rows[header][col-1] if header < len(rows) else None
                fields.append({'id': f'字段_{index+1:03d}', 'name': str(names[col-1]), 'column': col,
                               'index': index, 'kind': inferred_kind(names[col-1], cell)})
            tables.append({'id': f'Sheet_{len(tables)+1:03d}', 'path': str(path), 'sheet': name,
                           'header': header, 'width': width, 'visible': visible, 'fields': fields,
                           'rows': rows, 'layout': layout, 'palette': themes(book)})
            if len(tables) > MAX_SHEETS:
                raise ValueError('一次最多导出 32 个 Sheet')
    return tables


def plan_public_local(tables):
    """仅返回给本机 UI；原名称绝不会写入对外 ZIP。"""
    plans = [{'id': t['id'], 'source': f"{Path(t['path']).name} / {t['sheet']}", 'fields': t['fields']} for t in tables]
    identity = [{key: t[key] for key in ('id', 'path', 'sheet', 'header', 'width', 'fields')} for t in tables]
    token = sha256(json.dumps(identity, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return {'token': token, 'tables': plans, 'kinds': KINDS}


def inspect_privacy(config, request=None):
    """准备字段类型选择，界面不展示或传回真实数据样本。"""
    if request and request.get('mode') == 'vba':
        from desktop.vba_export import inspect_vba
        return inspect_vba(config, request)
    with ExitStack() as stack:
        return plan_public_local(_prepare(config, stack))


def sample_value(kind, row, column):
    """完全独立的固定合成器；输入只有用户选的类型及演示行列序号。"""
    if kind == 'date':
        return datetime(2020, 1, 1) + timedelta(days=row)
    values = {'text': f'示例文本_{row+1:02d}', 'id': f'DEMO-{row+1:04d}', 'name': f'示例姓名{row+1:02d}',
              'category': ['示例类别A', '示例类别B', '示例类别C'][row % 3], 'integer': (row+1)*3+column,
              'money': 1000 + row*137 + column*11.25, 'number': round(10 + row*1.25 + column*.1, 2),
              'percent': ((row*7) % 80+10)/100, 'boolean': row % 2 == 0,
              'phone': f'000-0000-{row+1:04d}', 'email': f'demo{row+1:03d}@example.invalid'}
    return values[kind]


def safe_style(source, target, palette, kind=None):
    """重建基础样式，不复制自定义数字格式、任意字体名称或内容相关条件规则。"""
    if source is None or source.font is None:
        return
    font = source.font
    target.font = Font(name=font.name if font.name in ('Calibri', 'Arial', 'Microsoft YaHei', '宋体', '黑体', '等线') else 'Calibri',
                       size=bounded(font.sz, 11, 1, 100), bold=bool(font.b), italic=bool(font.i),
                       color=color(font.color, palette, '#000000')[1:], underline=font.u, strike=font.strike, vertAlign=font.vertAlign)
    target.alignment = copy(source.alignment)
    fill = source.fill
    if getattr(fill, 'type', None) in ('linear', 'path'):
        target.fill = GradientFill(type=fill.type, degree=bounded(fill.degree, 0, 0, 360),
                                   stop=[Stop(color=color(s.color, palette, '#FFFFFF')[1:], position=s.position) for s in fill.stop])
    elif getattr(fill, 'patternType', None):
        target.fill = PatternFill(fill.patternType, fgColor=color(fill.fgColor, palette, '#FFFFFF')[1:], bgColor=color(fill.bgColor, palette, '#FFFFFF')[1:])
    border = source.border
    sides = {}
    for edge in ('left', 'right', 'top', 'bottom', 'diagonal'):
        side = getattr(border, edge)
        if side is not None:
            sides[edge] = Side(style=side.style, color=color(side.color, palette, '#000000')[1:])
    target.border = Border(**sides, diagonalUp=border.diagonalUp, diagonalDown=border.diagonalDown)
    # 数字格式里的引号文本也能带出客户名，因此只使用固定的通用格式。
    target.number_format = {'money': '#,##0.00', 'number': '0.00', 'percent': '0.0%', 'date': 'yyyy-mm-dd', 'integer': '0'}.get(kind, 'General')


def _write_table(book, table, sanitizer, mode, row_count, notices):
    """保留表头布局、列宽和代表性公式，数据区统一为固定的演示行。"""
    ws = book.create_sheet(table['id'])
    rows, header, layout, palette = table['rows'], table['header'], table['layout'], table['palette']
    fields = {field['column']: field for field in table['fields']}
    for col in range(1, table['width']+1):
        attrs = {}
        for dim in layout['columns']:
            if int(dim['min']) <= col <= int(dim['max']):
                attrs.update(dim)
        ws.column_dimensions[get_column_letter(col)].width = bounded(attrs.get('width', layout['defaultWidth']), 8.43, 1, 100)
        ws.column_dimensions[get_column_letter(col)].hidden = col not in table['visible']
    for r in range(1, header+1):
        ws.row_dimensions[r].height = bounded(layout['rows'].get(r, {}).get('ht', layout['defaultHeight']), 15, 1, 100)
        for col in table['visible']:
            cell = ws.cell(r, col)
            source = rows[r-1][col-1] if r <= len(rows) else None
            safe_style(source, cell, palette)
            if r == header:
                cell.value = fields[col]['id']
            elif source is not None and source.value is not None:
                cell.value = '演示表头'  # 标题、多级表头、日期等均不用原内容
    for left, top, right, bottom in layout['merges']:
        if bottom < header and right <= table['width'] and all(c in table['visible'] for c in range(left, right+1)):
            ws.merge_cells(start_row=top, end_row=bottom, start_column=left, end_column=right)
        elif top <= header <= bottom:
            notices.add('列名行的合并区域已展开，以便每个匿名字段都有独立列名')
    for field in table['fields']:
        col, kind = field['column'], field['kind']
        formula_source = next((row[col-1] for r, row in enumerate(rows[header:], header+1)
                               if r not in layout['hiddenRowIndices'] and row[col-1].data_type == 'f'), None)
        for n in range(row_count):
            cell = ws.cell(header+n+1, col)
            cell.font = Font(name='Calibri', size=11, color='243C31')
            cell.fill = PatternFill('solid', fgColor='F2F7F4' if n % 2 else 'FFFFFF')
            cell.alignment = Alignment(vertical='center')
            cell.number_format = {'money': '#,##0.00', 'number': '0.00', 'percent': '0.0%', 'date': 'yyyy-mm-dd', 'integer': '0'}.get(kind, 'General')
            cell.value = sample_value(kind, n, field['index']) if mode == 'synthetic' else None
            if formula_source is not None:
                try:
                    cell.value = sanitizer.sanitize(formula_source.value, formula_source.coordinate, cell.coordinate, table, field['id'])
                    table.setdefault('formulaSources', {})[formula_source.coordinate] = formula_source.value
                except (ValueError, TypeError, TokenizerError, TranslatorError):
                    notices.add('部分公式无法安全转换，已改为虚构值或空白；未导出原公式')
        # 各演示行统一行高，不复制真实记录间的样式差异。
    for n in range(row_count):
        ws.row_dimensions[header+n+1].height = 22
    ws.freeze_panes = f'A{header+1}'
    ws.sheet_view.showGridLines = layout['gridlines']
    return ws


def _xlsx_bytes(book):
    """保存新工作簿并检查没有图表、缓存、批注、连接或隐藏附件部件。"""
    output = BytesIO()
    book.save(output)
    data = output.getvalue()
    allowed = re.compile(r'(?:docProps/(?:app|core)\.xml|xl/(?:theme/theme1\.xml|styles\.xml|workbook\.xml|_rels/workbook\.xml\.rels|worksheets/sheet\d+\.xml)|_rels/\.rels|\[Content_Types\]\.xml)')
    with ZipFile(BytesIO(data)) as archive:
        for name in archive.namelist():
            if not allowed.fullmatch(name) or b'TargetMode="External"' in archive.read(name):
                raise ValueError('脱敏包检查失败，未生成可分享文件')
    return data


def ai_guide(export_id, schema):
    """给 AI 的契约只含公开的匿名结构，不夹带原路径、名称或环境日志。"""
    return f'''EasyExcel 严格脱敏演示包

本包不是实际数据：字段和 Sheet 已匿名，演示值是独立生成的虚构样例。
不要据此推断真实记录数量、取值、分布、统计、业务结论或类别。
不同演示表里相同的虚构编号不代表真实关联关系，请用户说明需要关联的匿名字段。
需要业务含义时，请向用户询问抽象的字段含义，不索要真实记录。
原图表、图片、批注、超链接、条件规则、验证、命名范围、外部连接、元数据和公式缓存不包含在本包。
空白模板或虚构模板仅保留表头布局、固定演示行及可安全转换的代表性公式。
跨行公式范围已规范到演示行；常量替换为演示参数；特殊公式可能移除。
演示表不用于回写或覆盖原表。真实公式参数只在本机映射中，切勿根据演示参数猜测真实值。

请编写在 EasyExcel 离线 Python 工作台运行的脚本。脚本开头使用：
ai_tables, ai_params = load_ai_context("{export_id}")
df = ai_tables["{schema[0]['id']}"]
# df 的列名是 字段_001 等匿名名，运行时值来自本机真实数据。
# ai_params["参数_0001"] 等参数运行时取原公式常量。没有此参数时不要引用。

不调用 pd.read_excel 读取演示包来完成真实分析，不使用演示行数切片真实数据。
可用 pd、np、plt、Decimal、display(data, title)、save_excel(data, filename)、save_figure(fig, filename)。
还可使用 openpyxl、numpy_financial、pyxirr、scipy、statsmodels、seaborn、xlsxwriter、babel、QuantLib。
所有实际处理和输出均在离线电脑完成；将结果写入工作台 OUTPUT_DIR。
不联网、不上传输入、日志或统计结果，不将真实数据写入返回给 AI 的文件。
若需要测试，使用虚构样例；AI 给出的是分析方法与代码，不是对真实数据的结论。

本机使用：保持同一批工作表和表头设置，导入 AI 生成的 .py 并运行。
load_ai_context 会核对文件、Sheet、表头和列结构；换电脑需要在离线电脑重新导出，不能把本机映射文件发给 AI。
请只转移这个 ZIP；程序用户目录、运行日志、分析结果和映射文件均不属于脱敏包。
用户需求（请只填写抽象业务含义）：
- 各匿名字段的用途：
- 多张表的关联键及一对多/一对一关系：
- 需要计算、绘制或导出的内容：
- 核对与舍入规则：
'''


def export_privacy(request, config):
    """原子输出对外 ZIP；包含真实对应关系的映射单独保存在本机。"""
    if request.get('mode') == 'vba':
        from desktop.vba_export import export_vba
        return export_vba(request, config)
    mode, count = request.get('mode'), request.get('sampleRows')
    if mode not in MODES or type(count) is not int or count not in (5, 20, 50):
        raise ValueError('请选择导出内容及 5、20 或 50 行演示数据')
    export_id = str(UUID(request['exportId']))
    if export_id != request['exportId']:
        raise ValueError('导出标识无效')
    output = Path(request['outputPath']).resolve()
    if output.suffix.lower() != '.zip' or not output.parent.is_dir():
        raise ValueError('请选择 ZIP 文件保存位置')
    if output in [Path(p).resolve() for p in config['selected_input_files']]:
        raise ValueError('不能覆盖源文件')
    mapping_dir = Path(request['mappingDir']).resolve()
    mapping_dir.mkdir(parents=True, exist_ok=True)
    mapping_path = mapping_dir / f'{export_id}.json'
    notices = set()
    with ExitStack() as stack:
        tables = _prepare(config, stack)
        plan = plan_public_local(tables)
        if request.get('token') != plan['token']:
            raise ValueError('表格结构或选择已变化，请重新打开脱敏导出')
        overrides = request.get('fieldKinds')
        if not isinstance(overrides, dict) or set(overrides) != {t['id'] for t in tables}:
            raise ValueError('字段类型选择不完整')
        for table in tables:
            kinds = overrides[table['id']]
            if not isinstance(kinds, list) or len(kinds) != len(table['fields']) or any(k not in KINDS for k in kinds):
                raise ValueError('字段类型无效，请重新选择')
            for field, kind in zip(table['fields'], kinds):
                field['kind'] = kind
        sanitizer = FormulaSanitizer(tables, count, mode)
        public = [{'id': t['id'], 'fields': [{'id': f['id'], 'kind': f['kind']} for f in t['fields']]} for t in tables]
        if mode != 'schema':
            for spec, table in zip(public, tables):
                spec['headerRow'] = table['header']
                for field_spec, field in zip(spec['fields'], table['fields']):
                    field_spec['excelColumn'] = get_column_letter(field['column'])
        content = None
        if mode != 'schema':
            book = openpyxl.Workbook()
            book.remove(book.active)
            book.properties.creator = 'EasyExcel'
            book.properties.lastModifiedBy = 'EasyExcel'
            book.properties.created = datetime(2000, 1, 1)
            book.calculation = CalcProperties(fullCalcOnLoad=True)
            for table in tables:
                _write_table(book, table, sanitizer, mode, count, notices)
            if sanitizer.parameters:
                ws = book.create_sheet('演示参数')
                ws.append(['参数标识', '虚构示例值'])
                for name, item in sanitizer.parameters.items():
                    ws.append([name, demo_parameter(item['kind'], mode)])
            content = _xlsx_bytes(book)
            book.close()
        mapping = {'version': 1, 'id': export_id, 'tables': [
            {'id': t['id'], 'path': t['path'], 'sheet': t['sheet'], 'header': t['header'],
             'fields': [{'id': f['id'], 'name': f['name'], 'index': f['index']} for f in t['fields']],
             'formulaSources': t.get('formulaSources', {})} for t in tables],
            'parameters': sanitizer.parameters}
        schema = {'version': 1, 'id': export_id, 'synthetic': True, 'mode': mode, 'tables': public,
                  'parameters': [{'id': key, 'kind': item['kind']} for key, item in sanitizer.parameters.items()],
                  'notices': sorted(notices)}
        if mode != 'schema':
            schema['sampleRows'] = count  # 这是用户指定的演示行数，绝不写入源行数。
        with NamedTemporaryFile(dir=output.parent, prefix='.easyexcel-', suffix='.tmp', delete=False) as temporary:
            temp_path = Path(temporary.name)
        wrote_mapping = False
        try:
            with ZipFile(temp_path, 'w', ZIP_DEFLATED) as archive:
                archive.writestr('结构说明.json', json.dumps(schema, ensure_ascii=False, indent=2, allow_nan=False))
                archive.writestr('给AI的说明.txt', ai_guide(export_id, public))
                if content:
                    archive.writestr('虚构演示.xlsx' if mode == 'synthetic' else '空白模板.xlsx', content)
            # 使用独占创建，不能覆盖任何旧映射；本机映射不在 ZIP 所在目录内。
            mapping_text = json.dumps(mapping, ensure_ascii=False, indent=2, allow_nan=False)
            if len(mapping_text.encode('utf-8')) > 8 * 1024 * 1024:
                raise ValueError('本机公式映射超过容量上限，请减少所选 Sheet')
            with mapping_path.open('x', encoding='utf-8') as stream:
                wrote_mapping = True
                stream.write(mapping_text)
            os.replace(temp_path, output)
        except BaseException:
            temp_path.unlink(missing_ok=True)
            if wrote_mapping:
                mapping_path.unlink(missing_ok=True)
            raise
    return {'output': str(output), 'exportId': export_id, 'sheets': len(tables), 'mode': mode,
            'sampleRows': count if mode != 'schema' else 0, 'warnings': sorted(notices)}


def load_ai_context(export_id, mapping_dir, tables, inputs):
    """由运行器在离线机载入对应真实数据，核对结构后才提供匿名字段。"""
    try:
        if str(UUID(export_id)) != export_id or not mapping_dir:
            raise ValueError
        path = Path(mapping_dir) / f'{export_id}.json'
        if path.stat().st_size > 8 * 1024 * 1024:
            raise ValueError
        binding = json.loads(path.read_text(encoding='utf-8'))
    except (ValueError, TypeError, OSError):
        raise ValueError('本机没有此演示包的映射，请在这台离线电脑重新导出后再生成脚本') from None
    if binding.get('version') != 1 or binding.get('id') != export_id:
        raise ValueError('本机映射格式无效，请重新导出')
    if len(binding['tables']) != len(inputs):
        raise ValueError('当前 Sheet 选择与脱敏包不同，请选择导出时的同一批工作表')
    result = {}
    for item in binding['tables']:
        matches = [i for i in inputs if os.path.normcase(i['path']) == os.path.normcase(item['path']) and i['sheet'] == item['sheet']]
        if len(matches) != 1 or matches[0]['headerRow'] != item['header']:
            raise ValueError('当前文件、Sheet 或表头与本机映射不符，请重新导出')
        frame = tables[matches[0]['key']]
        if [str(c) for c in frame.columns] != [field['name'] for field in item['fields']]:
            raise ValueError('列结构已变化，请重新导出脱敏包后再生成脚本')
        if item.get('formulaSources'):
            book = openpyxl.load_workbook(item['path'], read_only=True, data_only=False, keep_links=False)
            try:
                ws = book[item['sheet']]
                ws.reset_dimensions()
                seen = set()
                for row in ws.iter_rows(min_row=item['header']+1, max_row=item['header']+50, max_col=MAX_COLUMNS):
                    for cell in row:
                        coordinate = getattr(cell, 'coordinate', None)
                        if coordinate in item['formulaSources']:
                            seen.add(coordinate)
                            if cell.value != item['formulaSources'][coordinate]:
                                raise ValueError('源公式或公式参数已变化，请重新导出脱敏包后再生成脚本')
                if seen != set(item['formulaSources']):
                    raise ValueError('源公式位置已变化，请重新导出脱敏包')
            finally:
                book.close()
        aliased = frame.copy()
        aliased.columns = [field['id'] for field in item['fields']]
        result[item['id']] = aliased
    parameters = {key: item['value'] for key, item in binding['parameters'].items()}
    return result, parameters
