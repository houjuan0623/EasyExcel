"""桌面端到端测试样本，只写测试指定的临时目录。

Copyright (c) 2026 AbeLin · MIT License
"""
import datetime
import json
from pathlib import Path
import sys
from io import BytesIO
from zipfile import ZipFile

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, IconSetRule
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side


def visual_sheet(wb):
    """覆盖复杂单元格、原位置图表、双轴与 100 行之外的对象。"""
    ws = wb.create_sheet('图表与样式')
    ws.merge_cells('A1:F2')
    ws['A1'] = '2026 经营分析 · 图表与复杂样式'
    ws['A1'].font = Font(name='Microsoft YaHei', size=19, bold=True, color='FFFFFF')
    ws['A1'].fill = GradientFill(stop=['155D45', '4BA783'])
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 24
    for col, width in zip('ABCDEF', [16, 19, 19, 15, 18, 24]):
        ws.column_dimensions[col].width = width
    for c, label in enumerate(['月份', '销售额', '成本', '完成率', '净利润', '趋势'], 1):
        ws.cell(4, c, label)
        ws.cell(4, c).font = Font(bold=True, color='155D45')
        ws.cell(4, c).fill = PatternFill('solid', fgColor='DFF0E6')
    for r in range(5, 17):
        revenue = 18000 + (r-5) * 1600 + (r % 3) * 900
        cost = 17000 + (r-5) * 800
        values = [f'{r-4}月', revenue, cost, .65 + (r-5) * .028, revenue-cost, (r-5)/11]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.font = Font(name='Microsoft YaHei', size=11, color='243C31')
            cell.fill = PatternFill('solid', fgColor='F1F7F3' if r % 2 else 'FFFFFF')
            cell.alignment = Alignment(vertical='center')
            cell.border = Border(bottom=Side(style='thin', color='E1EBE5'))
        ws.row_dimensions[r].height = 27
        for col in (2, 3, 5):
            ws.cell(r, col).number_format = '#,##0;[Red](#,##0)'
        for col in (4, 6):
            ws.cell(r, col).number_format = '0.0%'
    ws.conditional_formatting.add('B5:B16', DataBarRule(start_type='min', end_type='max', color='4E9B78'))
    ws.conditional_formatting.add('D5:D16', ColorScaleRule(start_type='min', start_color='F9E6C5', end_type='max', end_color='8BCDAD'))
    ws.conditional_formatting.add('F5:F16', IconSetRule('3Arrows', 'percent', [0, 33, 67]))
    ws.conditional_formatting.add('E5:E16', CellIsRule(operator='greaterThan', formula=['10000'], font=Font(bold=True, color='19714C')))
    ws['A18'] = '项目 / 月份'
    ws['A18'].border = Border(diagonal=Side(style='thin', color='21845E'), diagonalDown=True)
    ws.row_dimensions[18].height = 36
    ws.merge_cells('B18:F18')
    ws['B18'] = CellRichText('提示：', TextBlock(InlineFont(b=True, color='21845E'), '绿色为较好表现'), ' · 数据条、色阶与图标保留')
    ws['B18'].alignment = Alignment(vertical='center')
    ws['A20'] = '图案填充'
    ws['A20'].fill = PatternFill('lightDown', fgColor='8BBDA4', bgColor='F4FAF6')
    bar, line = BarChart(), LineChart()
    bar.title = '月度销售与完成率'
    bar.y_axis.title = '销售额（元）'
    bar.y_axis.numFmt = '#,##0'
    line.y_axis.title = '完成率'
    line.y_axis.axId = 200
    line.y_axis.numFmt = '0%'
    line.y_axis.scaling.min, line.y_axis.scaling.max = 0, 1
    for chart, col in [(bar, 2), (line, 4)]:
        chart.add_data(Reference(ws, min_col=col, min_row=4, max_row=16), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=16))
    line.series[0].graphicalProperties.line.solidFill = 'ED7D31'
    bar += line
    bar.width, bar.height = 21, 11.5
    ws.add_chart(bar, 'H4')
    pie = DoughnutChart()
    pie.title = '第一季度销售占比'
    pie.add_data(Reference(ws, min_col=2, min_row=4, max_row=7), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7))
    ws.add_chart(pie, 'A110')
    return ws


def create(directory):
    """构造多 Sheet 的合成数据，不使用用户业务数据。"""
    path = directory / "2026年销售明细.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    departments = ["华东区域", "华南区域", "华北区域", "西南区域"]
    for index, name in enumerate(["10月销售明细", "9月销售明细", "8月销售明细", "归档数据"]):
        ws = wb.create_sheet(name)
        ws.append(["订单编号", "区域", "负责人", "产品名称", "销售金额", "日期"])
        for n in range(24):
            ws.append([f"SO-{index + 8:02d}{n + 1:04d}", departments[n % 4], f"负责人{n % 5 + 1}",
                       ["办公设备", "企业服务", "数据产品"][n % 3], 1280 + n * 350,
                       datetime.datetime(2026, 10 - index, n + 1)])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="21845E")
            cell.font = Font(color="FFFFFF", bold=True)
        if name == "归档数据":
            ws.sheet_state = "hidden"
    ws = wb.create_sheet("原样预览")
    ws.merge_cells("A1:F2")
    ws["A1"] = "十月销售总览"
    ws["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="155D45")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 22
    for label, width in zip("ABCDEF", [22, 17, 20, 16, 14, 34]):
        ws.column_dimensions[label].width = width
    for index, label in enumerate(["订单编号", "日期", "销售金额", "区域", "完成率", "备注"], 1):
        cell = ws.cell(4, index, label)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="155D45")
        cell.fill = PatternFill("solid", fgColor="DFEEE7")
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[4].height = 27
    edge = Side(style="thin", color="D4DED8")
    for row in range(5, 151):
        values = [f"ORDER-{row:04d}", datetime.datetime(2026, 10, (row % 28) + 1), 1200 + row * 125.5,
                  departments[row % 4], 0.72 + row / 1000, "预览保留原表样式"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row, col, value)
            cell.font = Font(name="Microsoft YaHei", size=10, color="243C31")
            cell.fill = PatternFill("solid", fgColor="F2F7F4" if row % 2 else "FFFFFF")
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 24
        ws.cell(row, 2).number_format = "yyyy/mm/dd"
        ws.cell(row, 3).number_format = '"¥"#,##0.00'
        ws.cell(row, 5).number_format = "0.0%"
    ws["F5"] = "文字自动换行\n第二行保持原样"
    ws["F5"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 38
    ws.row_dimensions[8].hidden = True
    ws.column_dimensions["G"].hidden = True
    ws["G5"] = "隐藏字段"
    ws.merge_cells("A99:B105")
    ws["A99"] = "跨越第 100 行的合并区域"
    ws["F100"] = "第 100 行，完整保留"
    ws["F101"] = "第 101 行，不应显示"
    visual_sheet(wb)
    wb.save(path)
    wb.close()
    print(json.dumps({"path": str(path)}, ensure_ascii=False))


def verify(directory):
    """检查仅所选月份导出，内容、表头样式均由 Python 保存。"""
    files = list((directory / "拆分结果").rglob("*.xlsx"))
    bundles = [path for path in files if path.parent.name == '总工作簿']
    files = [path for path in files if path.parent.name != '总工作簿']
    assert len(files) == 4, len(files)
    assert len(bundles) == 1, len(bundles)
    total = 0
    for path in files:
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["9月销售明细"], wb.sheetnames
        assert wb.active["A1"].fill.fgColor.rgb == "0021845E"
        total += wb.active.max_row - 1
        wb.close()
    assert total == 24, total
    bundled = openpyxl.load_workbook(bundles[0])
    assert set(bundled.sheetnames) == {'华东区域', '华南区域', '华北区域', '西南区域'}
    assert sum(ws.max_row - 1 for ws in bundled) == total
    for ws in bundled:
        assert ws['A1'].fill.fgColor.rgb == '0021845E'
    bundled.close()
    print(json.dumps({"files": len(files), "rows": total}))


def verify_privacy(directory):
    """展开实际 UI 生成的包，检查演示行数、匿名字段与敏感信息缺席。"""
    with ZipFile(directory / 'AI演示包.zip') as archive:
        schema = json.loads(archive.read('结构说明.json'))
        assert schema['mode'] == 'synthetic' and schema['sampleRows'] == 20
        assert schema['tables'][0]['id'] == 'Sheet_001'
        with ZipFile(BytesIO(archive.read('虚构演示.xlsx'))) as inner:
            raw = ''.join(inner.read(name).decode('utf-8') for name in inner.namelist())
            for private in ('2026年销售明细', '9月销售明细', '负责人1', 'SO-090001', '华东区域'):
                assert private not in raw
        wb = openpyxl.load_workbook(BytesIO(archive.read('虚构演示.xlsx')))
        assert wb.active.max_row == 21 and wb.active.max_column == 6
        assert wb.active['A1'].value == '字段_001'
        assert wb.active['C2'].value == '示例姓名01'
        assert wb.active['E2'].value == 1045
        wb.close()
    print(json.dumps({'exportId': schema['id']}))


def verify_vba(directory):
    """验证打包界面生成的 VBA 包保留原名称与原公式，清除业务数据。"""
    with ZipFile(directory / 'VBA结构包.zip') as archive:
        schema = json.loads(archive.read('结构说明.json'))
        assert schema['mode'] == 'vba' and schema['formulasUnchanged'] is True
        assert 'ThisWorkbook' in archive.read('给AI的VBA编程说明.txt').decode('utf-8')
        wb = openpyxl.load_workbook(BytesIO(archive.read(schema['workbooks'][0]['template'])))
        original = openpyxl.load_workbook(directory/'2026年销售明细.xlsx')
        assert wb.sheetnames == original.sheetnames
        assert wb['9月销售明细']['A1'].value == original['9月销售明细']['A1'].value
        assert wb['9月销售明细']['A2'].value is None
        for before, after in zip(original, wb):
            for row in before:
                for cell in row:
                    if cell.data_type == 'f':
                        assert after[cell.coordinate].value == cell.value
        wb.close()
        original.close()
    print(json.dumps({'originalNamesPreserved': True}))


if __name__ == "__main__":
    directory = Path(sys.argv[2])
    directory.mkdir(parents=True, exist_ok=True)
    {'create': create, 'verify': verify, 'privacy': verify_privacy, 'vba': verify_vba}[sys.argv[1]](directory)
