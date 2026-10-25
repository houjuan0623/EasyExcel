"""统计分析界面测试只使用虚构数据，生成物写入调用者指定的仓库外目录。

Copyright (c) 2026 AbeLin · MIT License
"""

from datetime import datetime
import json
from pathlib import Path
import sys

import openpyxl


def create(directory):
    path = directory / "统计分析演示.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "预算明细"
    sheet.append(["虚构预算 · 用于功能测试"])
    sheet.merge_cells("A1:F1")
    sheet.append([])
    sheet.append(["编号", "部门", "日期", "预算金额", "实际金额", "次数"])
    for index in range(1, 121):
        sheet.append(
            [
                f"ID{index:03d}",
                ["研发", "销售", "行政"][index % 3],
                datetime(2026, (index - 1) % 12 + 1, 1),
                index * 100,
                index * 90,
                index % 7 + 1,
            ]
        )
    other = book.create_sheet("补充指标")
    other.append(["编号", "利润"])
    for index in reversed(range(1, 121)):
        other.append([f"ID{index:03d}", index * 20])
    horizontal = book.create_sheet("横向排布")
    horizontal.append(["项目", "一月", "二月", "三月"])
    horizontal.append(["预算", 10, 20, 30])
    horizontal.append(["实际", 12, 24, 36])
    book.save(path)
    book.close()
    print(json.dumps({"path": str(path)}, ensure_ascii=False))


if __name__ == "__main__":
    folder = Path(sys.argv[1])
    folder.mkdir(parents=True, exist_ok=True)
    create(folder)
