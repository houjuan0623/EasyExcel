"""由独立的本机 Excel 后台实例渲染，再把图片和真实单元格坐标交给 Electron。

Copyright (c) 2026 AbeLin · MIT License
"""

import base64
from contextlib import ExitStack
import gc
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import re
import sys
from tempfile import TemporaryDirectory
import time
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.utils import get_column_letter

from desktop.analysis import MAX_ROWS, Reader, integer
from desktop.preview_format import display

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
MAX_PIXELS = 24_000_000
MAX_EDGE = 7000


def preview_copy(source, target):
    """只复制到临时目录，并关闭副本的自动重算；外部活动内容不交给 Excel 打开。"""
    if source.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("Excel 原样选区支持 .xlsx / .xlsm；请先将旧版 .xls 另存为 .xlsx，或选择基础选区")
    with ZipFile(source) as archive:
        for item in archive.infolist():
            name = item.filename.lower()
            if any(
                part in name
                for part in (
                    "macrosheets/",
                    "dialogsheets/",
                    "connections.xml",
                    "querytables/",
                    "externallinks/",
                    "activex/",
                    "embeddings/",
                )
            ):
                raise ValueError("此工作簿含外部连接、活动对象或旧式宏表，不能交给后台 Excel 自动打开；可使用基础选区")
            if name.endswith(".rels"):
                root = ET.fromstring(archive.read(item))
                if any(
                    r.get("TargetMode") == "External" and not r.get("Type", "").endswith("/hyperlink") for r in root
                ):
                    raise ValueError("此工作簿含外部对象引用，原样预览不会自动加载；可使用基础选区")
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                with archive.open(item) as stream:
                    for _, element in ET.iterparse(stream, events=("end",)):
                        if element.tag == NS + "f" and re.search(
                            r"(?:\b|_xlfn\.)(?:WEBSERVICE|RTD|IMAGE)\s*\(", element.text or "", re.I
                        ):
                            raise ValueError("此工作簿含联网或实时数据公式，原样预览不会自动执行；可使用基础选区")
                        element.clear()
        with ZipFile(target, "w", ZIP_DEFLATED) as output:
            for item in archive.infolist():
                content = archive.read(item)
                if item.filename == "xl/workbook.xml":
                    # 用原 XML 的字节替换保留扩展命名空间，避免重写整个 workbook 丢失 mc 前缀。
                    attrs = b' calcMode="manual" fullCalcOnLoad="0" forceFullCalc="0" calcOnSave="0"'
                    match = re.search(rb"<(?:[\w.-]+:)?calcPr\b[^>]*?/?>", content)
                    if match:
                        tag = re.sub(
                            rb"\s+(?:calcMode|fullCalcOnLoad|forceFullCalc|calcOnSave)\s*=\s*[\"\'][^\"\']*[\"\']",
                            b"",
                            match[0],
                        )
                        ending = b"/>" if tag.endswith(b"/>") else b">"
                        tag = re.sub(rb"\s*/?>$", attrs + ending, tag)
                        content = content[: match.start()] + tag + content[match.end() :]
                    else:
                        closing = re.search(rb"</((?:[\w.-]+:)?)workbook\s*>", content)
                        if closing:
                            content = (
                                content[: closing.start()]
                                + b"<"
                                + closing[1]
                                + b"calcPr"
                                + attrs
                                + b"/>"
                                + content[closing.start() :]
                            )
                output.writestr(item, content)


def expanded_bounds(top, left, total_rows, total_cols, merges):
    """分页范围包含完整合并格；超大合并区不偷偷改变选区或错配坐标。"""
    # 保留空白网格，浮动图表可能位于有值单元格之外。
    right, bottom = min(16384, left + 19), min(MAX_ROWS, top + 39)
    bounds = left, top, right, bottom
    for _ in range(len(merges) + 1):
        a, b, c, d = bounds
        for ml, mt, mr, mb in merges:
            if mr >= a and ml <= c and mb >= b and mt <= d:
                a, b, c, d = min(a, ml), min(b, mt), max(c, mr), max(d, mb)
        next_bounds = a, b, c, d
        if c - a + 1 > 80 or d - b + 1 > 200 or d > MAX_ROWS:
            raise ValueError("该页的合并区域过大，无法准确铺出可选坐标；请调整起始位置，或使用基础选区")
        if next_bounds == bounds:
            return bounds
        bounds = next_bounds
    raise ValueError("无法确定合并区域边界")


def cell_details(path, sheet, bounds, include_hidden, meta):
    """公式原文和格式化缓存独立保留，不依赖图片识别，也不改写公式。"""
    left, top, right, bottom = bounds
    result = {}
    with ExitStack() as stack:
        formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        value_book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        stack.callback(formula_book.close)
        stack.callback(value_book.close)
        formula_ws, value_ws = formula_book[sheet], value_book[sheet]
        formula_ws.reset_dimensions()
        value_ws.reset_dimensions()
        args = dict(min_row=top, max_row=bottom, min_col=left, max_col=right)
        for r, (formulas, values) in enumerate(zip(formula_ws.iter_rows(**args), value_ws.iter_rows(**args)), top):
            for c, (cell, saved) in enumerate(zip(formulas, values), left):
                if not include_hidden and (r in meta["hiddenRows"] or c in meta["hiddenCols"]):
                    continue
                formula = (
                    (cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", ""))
                    if cell.data_type == "f"
                    else ""
                )
                text, _ = display(saved.value, cell.number_format)
                if formula or text:
                    result[f"{get_column_letter(c)}{r}"] = {
                        "formula": formula,
                        "text": text,
                        "missingCache": bool(formula and saved.value is None),
                    }
    return result


def capture_picture(area, excel_pid, warnings):
    """Excel 的屏幕图包含浮动图表；短暂借用剪贴板后恢复，用户新复制的内容优先。"""
    import pythoncom
    import win32clipboard
    import win32process
    from PIL import Image, ImageGrab

    old = pythoncom.OleGetClipboard()
    before = win32clipboard.GetClipboardSequenceNumber()
    owned = None
    image = None
    try:
        area.CopyPicture(Appearance=1, Format=2)
        owned = win32clipboard.GetClipboardSequenceNumber()
        for _ in range(12):
            if win32clipboard.GetClipboardSequenceNumber() != owned:
                raise ValueError("渲染时剪贴板发生变化，请重试原样预览")
            image = ImageGrab.grabclipboard()
            if isinstance(image, Image.Image):
                return image.copy()
            pythoncom.PumpWaitingMessages()
            time.sleep(0.05)
        raise ValueError("Excel 未能生成原表画面，请重试")
    finally:
        if image is not None and hasattr(image, "close"):
            image.close()
        if owned is None and win32clipboard.GetClipboardSequenceNumber() != before:
            owner = win32clipboard.GetClipboardOwner()
            if owner and win32process.GetWindowThreadProcessId(owner)[1] == excel_pid:
                owned = win32clipboard.GetClipboardSequenceNumber()
        if owned is not None:
            for attempt in range(12):
                if win32clipboard.GetClipboardSequenceNumber() != owned:
                    break
                try:
                    pythoncom.OleSetClipboard(old)
                    owned = win32clipboard.GetClipboardSequenceNumber()
                    pythoncom.OleFlushClipboard()
                    break
                except Exception:
                    owner = win32clipboard.GetClipboardOwner()
                    if owner and win32process.GetWindowThreadProcessId(owner)[1] == os.getpid():
                        owned = win32clipboard.GetClipboardSequenceNumber()
                    if attempt == 11:
                        warnings.append("系统剪贴板繁忙，未能恢复渲染前的剪贴板内容")
                    pythoncom.PumpWaitingMessages()
                    time.sleep(0.1)


def render_excel(path, sheet, bounds, include_hidden, show_formulas, warnings):
    """仅使用新建的 Excel 实例；Job 对象确保后端中止时不会遗留该实例。"""
    import pythoncom
    import win32api
    import win32con
    import win32event
    import win32job
    import win32process
    from win32com.client import DispatchEx

    # 多个预览窗口/应用实例也不能同时占用系统剪贴板。
    mutex = win32event.CreateMutex(None, False, "Local\\EasyExcel.NativePreview")
    acquired = win32event.WaitForSingleObject(mutex, 60000)
    if acquired not in (win32event.WAIT_OBJECT_0, win32event.WAIT_ABANDONED):
        mutex.Close()
        raise ValueError("另一个 Excel 原样预览尚未完成，请稍后重试")
    pythoncom.OleInitialize()
    excel = book = job = process = ws = area = shape = None
    managed = False
    try:
        before = set(win32process.EnumProcesses())
        excel = DispatchEx("Excel.Application")
        pid = win32process.GetWindowThreadProcessId(excel.Hwnd)[1]
        if pid in before:
            excel = None  # 绝不关闭或接管用户已经打开的 Excel 进程。
            raise ValueError("无法建立独立的后台 Excel 实例，请稍后重试")
        job = win32job.CreateJobObject(None, "")
        info = win32job.QueryInformationJobObject(job, win32job.JobObjectExtendedLimitInformation)
        info["BasicLimitInformation"]["LimitFlags"] = win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        win32job.SetInformationJobObject(job, win32job.JobObjectExtendedLimitInformation, info)
        process = win32api.OpenProcess(win32con.PROCESS_SET_QUOTA | win32con.PROCESS_TERMINATE, False, pid)
        win32job.AssignProcessToJobObject(job, process)
        managed = True
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = False
        excel.AskToUpdateLinks = False
        excel.AutomationSecurity = 3
        excel.Workbooks.Add()
        excel.Calculation = -4135  # xlCalculationManual，打开源副本前即生效。
        book = excel.Workbooks.Open(
            str(path),
            UpdateLinks=0,
            ReadOnly=True,
            Password="",
            WriteResPassword="",
            IgnoreReadOnlyRecommended=True,
            Notify=False,
            AddToMru=False,
        )
        ws = book.Worksheets(sheet)
        ws.Activate()
        extent_row, extent_col = (
            int(ws.UsedRange.Row + ws.UsedRange.Rows.Count - 1),
            int(ws.UsedRange.Column + ws.UsedRange.Columns.Count - 1),
        )
        for i in range(1, min(int(ws.Shapes.Count), 1000) + 1):
            shape = ws.Shapes(i)
            extent_row = max(extent_row, int(shape.BottomRightCell.Row))
            extent_col = max(extent_col, int(shape.BottomRightCell.Column))
        excel.ActiveWindow.Zoom = 100
        excel.ActiveWindow.DisplayFormulas = show_formulas
        left, top, right, bottom = bounds
        if include_hidden:
            ws.Range(ws.Cells(top, left), ws.Cells(bottom, right)).EntireRow.Hidden = False
            ws.Range(ws.Cells(top, left), ws.Cells(bottom, right)).EntireColumn.Hidden = False
            warnings.append("已在临时预览副本中展开当前页隐藏行列，原文件的隐藏状态未修改")
        area = ws.Range(ws.Cells(top, left), ws.Cells(bottom, right))
        width, height = float(area.Width), float(area.Height)
        if min(width, height) <= 0:
            raise ValueError("这一页只有隐藏行列，请调整起始位置或在数据块中开启包含隐藏行列")
        if max(width, height) * 4 / 3 > MAX_EDGE or width * height * 16 / 9 > MAX_PIXELS:
            raise ValueError("这一页的行高或列宽过大，请调整起始行列，或使用基础选区")
        columns = [
            {"index": c, "start": float(ws.Cells(top, c).Left - area.Left), "size": float(ws.Columns(c).Width)}
            for c in range(left, right + 1)
        ]
        rows = [
            {"index": r, "start": float(ws.Cells(r, left).Top - area.Top), "size": float(ws.Rows(r).Height)}
            for r in range(top, bottom + 1)
        ]
        image = capture_picture(area, pid, warnings)
        try:
            if image.width * image.height > MAX_PIXELS or max(image.size) > MAX_EDGE:
                raise ValueError("Excel 返回的画面过大，请缩小预览范围")
            sx, sy = (image.width - 1) / width, (image.height - 1) / height
            for axis, scale in ((columns, sx), (rows, sy)):
                for item in axis:
                    item["start"] = round(item["start"] * scale, 3)
                    item["size"] = round(item["size"] * scale, 3)
            stream = BytesIO()
            image.save(stream, "PNG")
            return {
                "image": "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode("ascii"),
                "width": image.width,
                "height": image.height,
                "rows": rows,
                "columns": columns,
                "excelVersion": str(excel.Version),
                "totalRows": extent_row,
                "totalColumns": extent_col,
            }
        finally:
            image.close()
    finally:
        area = ws = shape = None
        try:
            if book is not None:
                book.Close(False)
        except Exception:
            pass
        book = None
        gc.collect()
        try:
            # 正常路径先释放 COM 引用，再关闭仅拥有该进程的 Job；避免 Quit 后向已退出的 RPC 服务 Release。
            if excel is not None and not managed:
                excel.Quit()
        except Exception:
            pass
        excel = None
        gc.collect()
        if process is not None:
            process.Close()
        if job is not None:
            job.Close()
        pythoncom.CoUninitialize()
        win32event.ReleaseMutex(mutex)
        mutex.Close()


def preview(request):
    """原样图片与透明选区层使用同一份 Excel 几何信息，数据全程留在本机。"""
    if sys.platform != "win32":
        raise ValueError("Excel 原样选区需要 Windows 与本机 Microsoft Excel；可选择基础选区")
    with ExitStack() as stack:
        reader = Reader(request["inputPath"], stack)
        path, _, meta, _ = reader.sheet(request["path"], request["sheet"])
        top = integer(request.get("startRow", 1), 1, MAX_ROWS, "预览起始行")
        left = integer(request.get("startColumn", 1), 1, 16384, "预览起始列")
        bounds = expanded_bounds(top, left, meta["rows"], meta["columns"], meta["merges"])
        include_hidden = request.get("includeHidden") is True
        details = (
            cell_details(path, request["sheet"], bounds, include_hidden, meta) if path.suffix.lower() != ".xls" else {}
        )
        warnings = []
        if any(cell["missingCache"] for cell in details.values()):
            warnings.append("本页部分公式没有已保存结果；Excel 画面可能显示空白或 0，统计时仍按缺失处理")
        with TemporaryDirectory(prefix="EasyExcel-preview-") as temporary:
            target = Path(temporary) / ("preview" + path.suffix.lower())
            preview_copy(path, target)
            try:
                result = render_excel(
                    target, request["sheet"], bounds, include_hidden, request.get("showFormulas") is True, warnings
                )
            except ImportError as error:
                raise ValueError("当前软件缺少 Excel 联动组件，请更新完整软件包，或使用基础选区") from error
            except ValueError:
                raise
            except Exception as error:
                raise ValueError(
                    "后台 Excel 未能完成原样渲染，请确认 Microsoft Excel 已安装且可正常启动；也可使用基础选区"
                ) from error
        if sha256(path.read_bytes()).hexdigest() != reader.hashes[os.path.normcase(str(path))]:
            raise ValueError("渲染过程中源文件发生变化，请重新读取")
        return {
            **result,
            "engine": "excel",
            "sheet": request["sheet"],
            "startRow": bounds[1],
            "startColumn": bounds[0],
            "endRow": bounds[3],
            "endColumn": bounds[2],
            "totalRows": max(meta["rows"], result["totalRows"]),
            "totalColumns": max(meta["columns"], result["totalColumns"]),
            "cells": details,
            "merges": [
                list(m)
                for m in meta["merges"]
                if m[0] <= bounds[2] and m[2] >= bounds[0] and m[1] <= bounds[3] and m[3] >= bounds[1]
            ],
            "warnings": warnings,
        }
