"""读取工作表前 100 行的文字、样式与布局，供 Electron 只读展示。

Copyright (c) 2026 AbeLin · MIT License
"""
import colorsys
import datetime
from itertools import islice
import json
import math
import re
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel
from openpyxl.xml.functions import iterparse

from desktop.preview_format import display

ROW_LIMIT = 100
COLUMN_LIMIT = 256
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
BORDER_STYLES = {
    "hair": (1, "solid"), "thin": (1, "solid"), "medium": (2, "solid"), "thick": (3, "solid"),
    "dotted": (1, "dotted"), "dashed": (1, "dashed"), "dashDot": (1, "dashed"),
    "dashDotDot": (1, "dashed"), "mediumDashed": (2, "dashed"), "mediumDashDot": (2, "dashed"),
    "mediumDashDotDot": (2, "dashed"), "slantDashDot": (2, "dashed"), "double": (3, "double"),
}


def bounded(value, default, low, high):
    """将文件里的尺寸限制在 Excel 合法范围，避免非有限 CSS 数值。"""
    try:
        number = float(value)
        return max(low, min(high, number)) if math.isfinite(number) else default
    except (TypeError, ValueError):
        return default


def themes(book):
    """按 OOXML 主题索引读取色板，而不是 XML 子节点顺序。"""
    if not book.loaded_theme:
        return []
    root = ET.fromstring(book.loaded_theme)
    scheme = root.find(".//{*}clrScheme")
    result = []
    for name in ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"):
        item = scheme.find("{*}" + name) if scheme is not None else None
        node = item[0] if item is not None and len(item) else None
        result.append(node.get("lastClr", node.get("val", "000000")) if node is not None else "000000")
    return result


def color(value, palette, fallback):
    """解析 RGB、索引色、主题色和明暗调整；忽略单元格颜色的 alpha。"""
    if value is None:
        return fallback
    rgb = None
    if value.type == "rgb":
        rgb = value.rgb
    elif value.type == "indexed" and 0 <= value.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[value.indexed]
    elif value.type == "theme" and 0 <= value.theme < len(palette):
        rgb = palette[value.theme]
    if not isinstance(rgb, str) or not re.fullmatch(r"[\da-fA-F]{6}|[\da-fA-F]{8}", rgb):
        return fallback
    rgb = rgb[-6:]
    if value.tint:
        channels = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        hue, light, saturation = colorsys.rgb_to_hls(*channels)
        light = light * (1 + value.tint) if value.tint < 0 else light * (1 - value.tint) + value.tint
        rgb = "".join(f"{round(c * 255):02X}" for c in colorsys.hls_to_rgb(hue, light, saturation))
    return "#" + rgb.upper()


def cell_style(cell, palette):
    """只输出明确允许的样式属性，不把文件内容当作 HTML 或 CSS 执行。"""
    font, fill, border, alignment = cell.font, cell.fill, cell.border, cell.alignment
    horizontal = {"centerContinuous": "center", "distributed": "justify", "fill": "left"}.get(alignment.horizontal, alignment.horizontal)
    style = {
        "fontFamily": json.dumps(font.name or "Calibri", ensure_ascii=False) + ', "Microsoft YaHei", sans-serif',
        "fontSize": round(bounded(font.sz, 11, 1, 409) * 4 / 3, 2),
        "fontWeight": "bold" if font.b else "normal", "fontStyle": "italic" if font.i else "normal",
        "color": color(font.color, palette, "#000000"),
        "backgroundColor": color(fill.fgColor, palette, "#FFFFFF") if getattr(fill, "patternType", None) else "#FFFFFF",
        "textAlign": horizontal if horizontal in ("left", "center", "right", "justify") else "general",
        "verticalAlign": "middle" if alignment.vertical in ("center", "distributed", "justify") else alignment.vertical or "bottom",
        "wrapText": bool(alignment.wrap_text), "shrinkToFit": bool(alignment.shrink_to_fit),
        "indent": min(250, int(alignment.indent or 0)) * 10,
        "rotation": int(alignment.text_rotation or 0),
        "textDecoration": " ".join(part for active, part in ((font.u, "underline"), (font.strike, "line-through")) if active) or "none",
    }
    for edge in ("left", "right", "top", "bottom"):
        side = getattr(border, edge)
        if side is not None and side.style in BORDER_STYLES:
            width, kind = BORDER_STYLES[side.style]
            style["border" + edge.title()] = f"{width}px {kind} {color(side.color, palette, '#000000')}"
    return style


def metadata(ws):
    """流式读布局和公式缓存；只保留前 100 行，避免为预览加载整张大表。"""
    data = {"rows": {}, "columns": [], "merges": [], "cache": {}, "warnings": set(),
            "defaultHeight": 15, "defaultWidth": 8.43, "gridlines": True, "lastRow": 0, "lastColumn": 1}
    sheet_data = None
    with ws.parent._archive.open(ws._worksheet_path) as stream:
        for event, element in iterparse(stream, events=("start", "end")):
            tag = element.tag.rsplit("}", 1)[-1]
            if event == "start":
                if tag == "sheetData":
                    sheet_data = element
                continue
            attrs = element.attrib
            if tag == "row":
                row = int(attrs.get("r", data["lastRow"] + 1))
                data["lastRow"] = max(data["lastRow"], row)
                if row <= ROW_LIMIT:
                    data["rows"][row] = dict(attrs)
                    for cell in element:
                        reference = re.match(r"[A-Za-z]+", cell.get("r", ""))
                        if reference:
                            data["lastColumn"] = max(data["lastColumn"], column_index_from_string(reference.group()))
                        if cell.find(NS + "f") is not None:
                            value = cell.find(NS + "v")
                            if value is not None and (value.text is not None or cell.get("t") == "str"):
                                data["cache"][cell.get("r")] = (cell.get("t", "n"), value.text or "")
                element.clear()
                if sheet_data is not None:
                    sheet_data.clear()
            elif tag == "col":
                data["columns"].append(dict(attrs))
                element.clear()
            elif tag == "mergeCell":
                bounds = range_boundaries(attrs["ref"])
                if bounds[1] <= ROW_LIMIT and bounds[0] <= COLUMN_LIMIT:
                    data["merges"].append(bounds)
                element.clear()
            elif tag == "sheetFormatPr":
                data["defaultHeight"] = attrs.get("defaultRowHeight", 15)
                data["defaultWidth"] = attrs.get("defaultColWidth", 8.43)
            elif tag == "sheetView":
                data["gridlines"] = attrs.get("showGridLines", "1") not in ("0", "false")
            elif tag == "conditionalFormatting":
                data["warnings"].add("条件格式未渲染，显示单元格基础样式")
                element.clear()
            elif tag == "drawing":
                data["warnings"].add("图片和图表不在表格预览范围内")
    return data


def xlsx_data(path, sheet):
    """用 openpyxl 读取前 100 行样式，并从文件缓存中取公式的显示值。"""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet not in book.sheetnames or book[sheet].sheet_state != "visible":
            raise ValueError("该工作表不存在或已被隐藏，请重新读取文件")
        ws = book[sheet]
        data = metadata(ws)
        total = max(ws.max_row or 0, data["lastRow"])
        width = max(data["lastColumn"], ws.max_column or 1, max((merge[2] for merge in data["merges"]), default=1))
        height = min(ROW_LIMIT, total)
        palette = themes(book)
        # reset_dimensions 避免文件错误声明 A1:A1 时漏掉实际单元格。
        ws.reset_dimensions()
        rows = list(islice(ws.iter_rows(min_row=1, max_row=height, max_col=COLUMN_LIMIT), height))
        for row in rows:
            width = max(width, max((i + 1 for i, cell in enumerate(row) if cell.value is not None or getattr(cell, "has_style", False)), default=1))
        cells = {}
        for r, row in enumerate(rows, 1):
            for c, cell in enumerate(row[:min(width, COLUMN_LIMIT)], 1):
                if cell.font is None:
                    continue
                value = cell.value
                if cell.data_type == "f":
                    cached = data["cache"].get(cell.coordinate)
                    if cached is None:
                        data["warnings"].add("部分公式没有缓存结果，暂显示公式文本；可用 Excel 保存计算结果后刷新")
                        value = str(value) if isinstance(value, str) else "（公式结果未缓存）"
                    else:
                        kind, raw = cached
                        if kind == "b":
                            value = raw == "1"
                        elif kind in ("str", "e"):
                            value = raw
                        else:
                            value = float(raw)
                            if is_date_format(cell.number_format):
                                value = from_excel(value, book.epoch)
                text, approximate = display(value, cell.number_format)
                if approximate:
                    data["warnings"].add("部分特殊数字格式暂显示原始值")
                style = cell_style(cell, palette)
                cells[r, c] = {"text": text, "style": style, "numeric": isinstance(value, (int, float, datetime.date, datetime.time, datetime.timedelta)) and not isinstance(value, bool),
                               "center": isinstance(value, bool) or cell.data_type == "e"}
        return assemble(data, cells, width, total)
    finally:
        book.close()


def xls_data(path, sheet):
    """旧版 .xls 同样由 Python 读取格式、尺寸和合并范围。"""
    import xlrd
    book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
    try:
        if sheet not in book.sheet_names():
            raise ValueError("没有找到这个工作表，请重新读取文件")
        ws = book.sheet_by_name(sheet)
        if ws.visibility != 0:
            raise ValueError("该工作表不存在或已被隐藏，请重新读取文件")

        def rgb(index, default="000000"):
            channels = book.colour_map.get(index)
            return "".join(f"{n:02X}" for n in channels) if channels is not None else default

        styles = {}
        for index, xf in enumerate(book.xf_list):
            font = book.font_list[xf.font_index]
            alignment, border = xf.alignment, xf.border
            edges = {}
            kinds = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair", 8: "mediumDashed", 9: "dashDot", 10: "mediumDashDot", 11: "dashDotDot", 12: "mediumDashDotDot", 13: "slantDashDot"}
            for edge in ("left", "right", "top", "bottom"):
                edges[edge] = Side(style=kinds.get(getattr(border, edge + "_line_style")), color=rgb(getattr(border, edge + "_colour_index")))
            rotation = alignment.rotation
            styles[index] = cell_style(SimpleNamespace(
                font=Font(name=font.name, size=font.height / 20, bold=bool(font.bold), italic=bool(font.italic),
                          strike=bool(font.struck_out), underline="single" if font.underlined else None, color=rgb(font.colour_index)),
                fill=PatternFill("solid", fgColor=rgb(xf.background.pattern_colour_index, "FFFFFF")) if xf.background.fill_pattern else PatternFill(),
                border=Border(**edges),
                alignment=Alignment(horizontal={1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous", 7: "distributed"}.get(alignment.hor_align, "general"),
                                    vertical={0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}.get(alignment.vert_align, "bottom"),
                                    wrap_text=bool(alignment.text_wrapped), shrink_to_fit=bool(alignment.shrink_to_fit),
                                    indent=alignment.indent_level, text_rotation=rotation if 0 <= rotation <= 180 or rotation == 255 else 0)), [])
        data = {"rows": {r + 1: {"ht": info.height / 20, "hidden": str(info.hidden)} for r, info in ws.rowinfo_map.items() if r < ROW_LIMIT},
                "columns": [{"min": c + 1, "max": c + 1, "width": info.width / 256, "hidden": str(info.hidden)} for c, info in ws.colinfo_map.items()],
                "merges": [(cl + 1, rl + 1, ch, rh) for rl, rh, cl, ch in ws.merged_cells if rl < ROW_LIMIT and cl < COLUMN_LIMIT],
                "defaultHeight": (ws.default_row_height or 300) / 20,
                "defaultWidth": (ws.defcolwidth or 8), "gridlines": bool(ws.show_grid_lines), "warnings": set()}
        cells = {}
        for r in range(min(ROW_LIMIT, ws.nrows)):
            for c in range(min(COLUMN_LIMIT, ws.ncols)):
                cell = ws.cell(r, c)
                xf = book.xf_list[cell.xf_index]
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(value, "#ERROR!")
                fmt = book.format_map.get(xf.format_key)
                text, approximate = display(value, fmt.format_str if fmt else "General")
                if approximate:
                    data["warnings"].add("部分特殊数字格式暂显示原始值")
                cells[r + 1, c + 1] = {"text": text, "style": styles[cell.xf_index], "numeric": cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE),
                                     "center": cell.ctype in (xlrd.XL_CELL_BOOLEAN, xlrd.XL_CELL_ERROR)}
        return assemble(data, cells, max(ws.ncols, 1), ws.nrows)
    finally:
        book.release_resources()


def assemble(data, cells, width, total):
    """生成可见网格，合并区域按 100 行窗口裁剪；隐藏内容不传给界面。"""
    end = min(ROW_LIMIT, total)
    columns = []
    for index in range(1, min(COLUMN_LIMIT, width) + 1):
        attrs = {}
        for item in data["columns"]:
            if int(item["min"]) <= index <= int(item["max"]):
                attrs.update(item)
        if attrs.get("hidden") in ("1", "true") or float(attrs.get("width", data["defaultWidth"])) == 0:
            continue
        units = bounded(attrs.get("width", data["defaultWidth"]), 8.43, 0, 255)
        pixels = round(units * 12) if units < 1 else math.floor(units * 7 + 5)
        columns.append({"index": index, "label": get_column_letter(index), "width": max(1, pixels)})
    rows = []
    for index in range(1, end + 1):
        attrs = data["rows"].get(index, {})
        if attrs.get("hidden") in ("1", "true") or float(attrs.get("ht", data["defaultHeight"])) == 0:
            continue
        rows.append({"index": index, "height": round(bounded(attrs.get("ht", data["defaultHeight"]), 15, 0, 409.5) * 4 / 3, 2), "cells": []})
    merged, covered, merged_edges = {}, set(), {}
    for left, top, right, bottom in data["merges"]:
        visible_rows = [row["index"] for row in rows if top <= row["index"] <= bottom]
        visible_cols = [col["index"] for col in columns if left <= col["index"] <= right]
        if not visible_rows or not visible_cols:
            continue
        anchor = visible_rows[0], visible_cols[0]
        merged[anchor] = (len(visible_rows), len(visible_cols), (top, left))
        covered.update((r, c) for r in visible_rows for c in visible_cols if (r, c) != anchor)
        # Excel 可能仅在合并区域边缘的单元格保存边框，补齐锚点缺失的边。
        edges = {}
        for edge, positions in (
            ("Top", [(top, c) for c in visible_cols]), ("Bottom", [(bottom, c) for c in visible_cols]),
            ("Left", [(r, left) for r in visible_rows]), ("Right", [(r, right) for r in visible_rows]),
        ):
            for position in positions:
                border = cells.get(position, {}).get("style", {}).get("border" + edge)
                if border:
                    edges["border" + edge] = border
                    break
        merged_edges[anchor] = edges
    styles, style_keys = [], {}
    for row in rows:
        for col in columns:
            position = row["index"], col["index"]
            if position in covered:
                continue
            source = cells.get(position, {"text": "", "style": {}})
            span_rows, span_cols, original = merged.get(position, (1, 1, position))
            style = dict(source["style"])
            for edge, border in merged_edges.get(position, {}).items():
                style.setdefault(edge, border)
            text = source["text"] if original == position else ""
            if style.get("textAlign") == "general":
                style["textAlign"] = "center" if source.get("center") else "right" if source.get("numeric") else "left"
            key = json.dumps(style, sort_keys=True)
            if key not in style_keys:
                style_keys[key] = len(styles)
                styles.append(style)
            row["cells"].append({"column": col["index"], "text": text, "style": style_keys[key], "rowSpan": span_rows, "colSpan": span_cols})
    if width > COLUMN_LIMIT:
        data["warnings"].add(f"工作表较宽，本次预览前 {COLUMN_LIMIT} 列")
    return {"columns": columns, "rows": rows, "styles": styles, "rowLimit": ROW_LIMIT, "endRow": end,
            "totalRows": total, "hiddenRows": end - len(rows), "hiddenColumns": min(COLUMN_LIMIT, width) - len(columns),
            "truncatedColumns": width > COLUMN_LIMIT, "showGridLines": data["gridlines"], "warnings": sorted(data["warnings"])}


def read_preview(path, sheet):
    """根据真实文件格式选择 Python 读取器。"""
    return xls_data(path, sheet) if path.suffix.lower() == ".xls" else xlsx_data(path, sheet)
