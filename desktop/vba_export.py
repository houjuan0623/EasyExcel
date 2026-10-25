"""供原工作簿 VBA 编程的结构副本：保留身份和原公式，常量仅按表头白名单复制。

Copyright (c) 2026 AbeLin · MIT License
"""
from collections import Counter
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import posixpath
import re
from tempfile import NamedTemporaryFile
from uuid import UUID
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
C = 'http://schemas.openxmlformats.org/package/2006/content-types'
NS = '{' + S + '}'
RID = '{' + R + '}id'
MAX_CELLS = 1_000_000
MAX_HEADER_CELLS = 20_000
MAX_SHEETS = 64
MAX_BYTES = 200 * 1024 * 1024


def xml(data):
    """仅解析包内 XML，不解析外部实体。"""
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('工作簿包含不支持的 XML 声明，未导出')
    return ET.fromstring(data)


def encoded(element):
    return ET.tostring(element, encoding='utf-8', xml_declaration=True)


def local(tag):
    return tag.rsplit('}', 1)[-1]


def clean_node(element):
    """丢弃扩展节点及属性，避免把不可见扩展数据带出。"""
    element = deepcopy(element)
    for node in element.iter():
        for child in list(node):
            if not child.tag.startswith(NS) or local(child.tag) == 'extLst':
                node.remove(child)
        for key in list(node.attrib):
            if key.startswith('{') and key != RID:
                del node.attrib[key]
    return element


def part_target(source, target):
    """关系必须指向同一压缩包内的规范路径。"""
    if not target or '\\' in target or re.match(r'^[a-zA-Z]+:', target):
        raise ValueError('工作簿关系路径无效，未导出')
    result = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(source), target))
    if result.startswith('../') or result == '..':
        raise ValueError('工作簿关系越界，未导出')
    return result


def relationships(archive, part):
    relpart = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    return {r.get('Id'): r for r in xml(archive.read(relpart))} if relpart in archive.namelist() else {}


def rels_bytes(entries):
    root = ET.Element('{' + P + '}Relationships')
    for attrs in entries:
        ET.SubElement(root, '{' + P + '}Relationship', attrs)
    return encoded(root)


def cell_value(cell, strings):
    """只解码普通值；公式从不求值。"""
    if cell.find(NS + 'f') is not None:
        return None
    kind = cell.get('t', 'n')
    value = cell.findtext(NS + 'v')
    if kind == 's':
        return strings[int(value)] if value is not None else None
    if kind == 'inlineStr':
        item = cell.find(NS + 'is')
        return rich_text(item) if item is not None else ''
    return value


def rich_text(element):
    return ''.join((child.text or '') if child.tag == NS + 't' else ''.join(t.text or '' for t in child.findall(NS + 't'))
                   for child in element if child.tag in (NS + 't', NS + 'r'))


def parse_ranges(value):
    """表头必须是显式有限区域；空字符串表示不保留任何常量。"""
    if not isinstance(value, str) or len(value) > 2000:
        raise ValueError('表头区域格式无效，请使用 A1:D3；多个区域用分号分隔')
    ranges, size = [], 0
    for text in re.split('[;；,，]', value):
        text = text.strip().upper().replace('$', '')
        if not text:
            continue
        if not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?', text):
            raise ValueError('表头区域必须包含行号和列号，例如 A7:K9；不要填写整行或整列')
        left, top, right, bottom = range_boundaries(text)
        if left > right or top > bottom or right > 16384 or bottom > 1048576:
            raise ValueError('表头区域超出 Excel 范围')
        size += (right-left+1) * (bottom-top+1)
        if size > MAX_HEADER_CELLS:
            raise ValueError('每张表的表头区域最多包含 20,000 个单元格')
        span = f'{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}'
        if span not in [r[0] for r in ranges]:
            ranges.append((span, left, top, right, bottom))
    return ranges


def inside(address, ranges):
    left, top, _, _ = range_boundaries(address)
    return any(a <= left <= c and b <= top <= d for _, a, b, c, d in ranges)


def data_extent(sheet):
    """部分工作簿不写 dimension，此时从实际单元格确定结构边界。"""
    dimension = sheet.find(NS + 'dimension')
    if dimension is not None and dimension.get('ref'):
        return dimension.get('ref')
    cells = [range_boundaries(c.get('r')) for c in sheet.iter(NS + 'c')]
    if not cells:
        return 'A1'
    return f'{get_column_letter(min(c[0] for c in cells))}{min(c[1] for c in cells)}:{get_column_letter(max(c[0] for c in cells))}{max(c[1] for c in cells)}'


class Review:
    """按相同固定值和引用汇总检查项；不把百万个同类公式塞进界面。"""
    def __init__(self):
        self.groups = {}

    def add(self, kind, value, location, formula=''):
        key = kind, value
        item = self.groups.setdefault(key, {'kind': kind, 'value': value, 'count': 0, 'locations': [], 'formula': formula})
        item['count'] += 1
        if len(item['locations']) < 8:
            item['locations'].append(location)
        if len(self.groups) > 5000:
            raise ValueError('公式检查项超过 5,000 类，请先在本机按工作簿分批导出')

    def formula(self, expression, location):
        if not expression:
            return
        expression = expression if expression.startswith('=') else '=' + expression
        try:
            tokens = Tokenizer(expression).items
            literals = [(t.subtype, t.value) for t in tokens if t.type == 'OPERAND' and t.subtype in ('TEXT', 'NUMBER', 'LOGICAL')]
            if literals:
                self.add('公式固定值', '；'.join(v for _, v in literals), location, expression)
            external = [t.value for t in tokens if t.type == 'OPERAND' and t.subtype == 'RANGE' and ('[' in t.value and '!' in t.value or '://' in t.value or '\\' in t.value)]
            for ref in external:
                self.add('外部公式引用', ref, location, expression)
            for token in tokens:
                if token.type == 'FUNC' and token.value.upper().rstrip('(').split('.')[-1] in ('WEBSERVICE', 'HYPERLINK', 'RTD'):
                    self.add('外部访问函数', token.value, location, expression)
        except Exception:
            self.add('需人工检查的公式', expression, location, expression)


def _suggest_ranges(sheet, header, strings, tables):
    """从列名行及相邻合并表头建议区域；最终由用户检查和修改。"""
    cells = {c.get('r'): c for c in sheet.iter(NS + 'c')}
    by_row = Counter()
    max_col = 1
    for address, cell in cells.items():
        col, row, _, _ = range_boundaries(address)
        max_col = max(max_col, col)
        if cell.get('t') in ('s', 'inlineStr', 'str') and cell_value(cell, strings):
            by_row[row] += 1
    spans = []
    if header and 1 <= header <= 1000:
        top = header
        for merge in sheet.iter(NS + 'mergeCell'):
            _left, start, _right, end = range_boundaries(merge.get('ref'))
            if start <= header <= end:
                top = min(top, start)
        while top > max(1, header-3) and by_row[top-1] >= 2:
            top -= 1
        spans.append(f'A{top}:{get_column_letter(max_col)}{header}')
    for table in tables:
        if table.get('headerRowCount', '1') != '0':
            left, top, right, _ = range_boundaries(table.get('ref'))
            spans.append(f'{get_column_letter(left)}{top}:{get_column_letter(right)}{top}')
    return ';'.join(dict.fromkeys(spans))


def _base_package():
    """从空白工作簿取得标准内容类型及主题，不沿用原文档的元数据。"""
    wb = openpyxl.Workbook()
    wb.properties.creator = 'EasyExcel'
    wb.properties.lastModifiedBy = 'EasyExcel'
    wb.properties.created = datetime(2000, 1, 1)
    wb.properties.modified = datetime(2000, 1, 1)
    data = BytesIO()
    wb.save(data)
    wb.close()
    with ZipFile(data) as archive:
        return {name: archive.read(name) for name in ('docProps/app.xml', 'docProps/core.xml', 'xl/theme/theme1.xml')}


def safe_styles(data):
    """保留样式索引和基础格式，移除可能嵌有业务文字的格式及自定义名称。"""
    root = clean_node(xml(data))
    changed = 0
    for node in root.iter(NS + 'name'):
        node.set('val', 'Calibri')
    for node in root.iter(NS + 'numFmt'):
        code = node.get('formatCode', '')
        replacement = 'yyyy-mm-dd hh:mm:ss' if is_date_format(code) else '0.00%' if '%' in code else '#,##0.00'
        if code != replacement:
            changed += 1
        node.set('formatCode', replacement)
    for index, node in enumerate(root.iter(NS + 'cellStyle')):
        node.set('name', 'Normal' if index == 0 else f'Style_{index}')
    styles = root.find(NS + 'tableStyles')
    if styles is not None:
        styles.clear()
        styles.attrib.update(count='0', defaultTableStyle='TableStyleMedium2', defaultPivotStyle='PivotStyleLight16')
    return encoded(root), changed


def _build(path, number, config, overrides, review):
    """重建一个完整工作簿；不把跨文件的同名 Sheet 合并或改名。"""
    from desktop.bridge import inspect_book
    raw = path.read_bytes()
    if path.suffix.lower() not in ('.xlsx', '.xlsm'):
        raise ValueError('VBA 结构模板支持 .xlsx 和 .xlsm，请先在本机转换旧格式')
    parts = _base_package()
    workbook_id = f'workbook_{number:03d}'
    model = {'id': workbook_id, 'sourceName': path.name, 'template': f'{workbook_id}/VBA结构模板.xlsx', 'sheets': [], 'definedNames': [], 'externalLinks': [],
             'omitted': ['真实数据常量与公式缓存', '原 VBA 项目', '图表及透视缓存、图片、批注、超链接', '条件格式、数据验证、筛选条件', '文档属性、外部连接与自定义 XML'],
             'note': '仅保留用户确认的表头常量；公式、命名公式、表对象公式保留原文。基础样式经过清理。'}
    plans, names_seen = [], set()
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) != len(set(x.filename for x in entries)) or sum(x.file_size for x in entries) > MAX_BYTES:
            raise ValueError('工作簿包含重复包部件或解压后超过 200 MB，未导出')
        source_book = xml(archive.read('xl/workbook.xml'))
        source_rels = relationships(archive, 'xl/workbook.xml')
        strings = [rich_text(si) for si in xml(archive.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in archive.namelist() else []
        info = {s['name']: s for s in inspect_book(path, config)['sheets']}
        workbook = ET.Element(NS + 'workbook')
        source_pr = source_book.find(NS + 'workbookPr')
        pr = {'updateLinks': 'never'}
        if source_pr is not None and source_pr.get('date1904') is not None:
            pr['date1904'] = source_pr.get('date1904')
        ET.SubElement(workbook, NS + 'workbookPr', pr)
        views = source_book.find(NS + 'bookViews')
        if views is not None:
            workbook.append(clean_node(views))
        else:
            ET.SubElement(ET.SubElement(workbook, NS + 'bookViews'), NS + 'workbookView')
        sheets_node = ET.SubElement(workbook, NS + 'sheets')
        workbook_rels = []
        content_types = []
        total_cells = formula_count = table_index = 0
        for index, source_sheet in enumerate(source_book.find(NS + 'sheets'), 1):
            if index > MAX_SHEETS:
                raise ValueError('VBA 结构模板每本最多支持 64 张 Sheet')
            name, state = source_sheet.get('name'), source_sheet.get('state', 'visible')
            relation = source_rels.get(source_sheet.get(RID))
            if relation is None or relation.get('Type', '').rsplit('/', 1)[-1] != 'worksheet' or relation.get('TargetMode') == 'External':
                raise ValueError('VBA 结构模板暂不支持图表工作表或 Excel 4.0 宏表；未生成不完整副本')
            source_part = part_target('xl/workbook.xml', relation.get('Target'))
            source = xml(archive.read(source_part))
            source_sheet_rels = relationships(archive, source_part)
            source_tables = []
            for table_part in source.iter(NS + 'tablePart'):
                rel = source_sheet_rels.get(table_part.get(RID))
                if rel is None or rel.get('TargetMode') == 'External' or not rel.get('Type', '').endswith('/table'):
                    raise ValueError('Excel 表对象关系无效，未导出')
                source_tables.append(xml(archive.read(part_target(source_part, rel.get('Target')))))
            sheet_id = f'{workbook_id}_sheet_{index:03d}'
            names_seen.add(sheet_id)
            suggested = _suggest_ranges(source, info.get(name, {}).get('headerRow'), strings, source_tables)
            ranges = parse_ranges(overrides.get(sheet_id, suggested))
            for table in source_tables:
                if table.get('headerRowCount', '1') != '0':
                    left, top, right, _ = range_boundaries(table.get('ref'))
                    if any(not inside(f'{get_column_letter(c)}{top}', ranges) for c in range(left, right+1)):
                        raise ValueError(f'请保留 Excel 表对象的表头区域 {get_column_letter(left)}{top}:{get_column_letter(right)}{top}，以免字段名变化')
            span_text = ';'.join(r[0] for r in ranges)
            sheet_model = {'id': sheet_id, 'name': name, 'index': index, 'state': state, 'headerRanges': span_text,
                           'headers': [], 'formulas': [], 'tables': [], 'mergedRanges': [m.get('ref') for m in source.iter(NS + 'mergeCell')],
                           'protected': source.find(NS + 'sheetProtection') is not None,
                           'dataExtent': data_extent(source)}
            sheet_pr = source.find(NS + 'sheetPr')
            sheet_model['codeName'] = sheet_pr.get('codeName', '') if sheet_pr is not None else ''
            out = ET.Element(NS + 'worksheet')
            sheet_rels = []
            allowed = {'sheetPr', 'dimension', 'sheetViews', 'sheetFormatPr', 'cols', 'sheetData', 'autoFilter', 'mergeCells',
                       'printOptions', 'pageMargins', 'pageSetup', 'rowBreaks', 'colBreaks', 'tableParts'}
            for child in source:
                tag = local(child.tag)
                if tag not in allowed or not child.tag.startswith(NS):
                    continue
                node = clean_node(child)
                if tag == 'sheetData':
                    for row in node:
                        for cell in list(row):
                            if cell.tag != NS + 'c':
                                row.remove(cell)
                                continue
                            total_cells += 1
                            if total_cells > MAX_CELLS:
                                raise ValueError('VBA 结构模板每本最多处理 1,000,000 个实际单元格，请分批处理')
                            address = cell.get('r')
                            if not address or not re.fullmatch('[A-Z]{1,3}[1-9][0-9]*', address):
                                raise ValueError('工作表含无效单元格地址，未导出')
                            formula = cell.find(NS + 'f')
                            value = cell_value(cell, strings) if formula is None else None
                            keep = formula is None and inside(address, ranges) and value is not None
                            old_type = cell.get('t', 'n')
                            cell.attrib = {k: v for k, v in cell.attrib.items() if k in ('r', 's')}
                            for sub in list(cell):
                                cell.remove(sub)
                            if formula is not None:
                                cell.append(formula)
                                formula_count += 1
                                sheet_model['formulas'].append({'cell': address, 'text': formula.text or '', 'attributes': dict(formula.attrib)})
                                review.formula(formula.text, f'{path.name} / {name}!{address}')
                            elif keep:
                                if old_type in ('s', 'inlineStr', 'str'):
                                    cell.set('t', 'inlineStr')
                                    ET.SubElement(ET.SubElement(cell, NS + 'is'), NS + 't', {'{http://www.w3.org/XML/1998/namespace}space': 'preserve'}).text = value
                                    kind = 'text'
                                else:
                                    cell.set('t', old_type)
                                    ET.SubElement(cell, NS + 'v').text = value
                                    kind = old_type
                                sheet_model['headers'].append({'cell': address, 'value': value, 'type': kind})
                elif tag == 'autoFilter':
                    node = ET.Element(NS + 'autoFilter', {'ref': child.get('ref')})
                elif tag == 'pageSetup':
                    node.attrib.pop(RID, None)
                elif tag == 'tableParts':
                    node = ET.Element(NS + 'tableParts', {'count': str(len(source_tables))})
                    for table in source_tables:
                        table_index += 1
                        rid = f'rIdTable{table_index}'
                        ET.SubElement(node, NS + 'tablePart', {RID: rid})
                        target = f'xl/tables/table{table_index}.xml'
                        sheet_rels.append({'Id': rid, 'Type': R + '/table', 'Target': f'../tables/table{table_index}.xml'})
                        table_out = ET.Element(NS + 'table', {k: v for k, v in table.attrib.items() if k in ('id', 'name', 'displayName', 'ref', 'headerRowCount', 'totalsRowCount', 'totalsRowShown')})
                        columns = ET.SubElement(table_out, NS + 'tableColumns')
                        column_models = []
                        for col in table.findall(NS + 'tableColumns/' + NS + 'tableColumn'):
                            attrs = {k: v for k, v in col.attrib.items() if k in ('id', 'name', 'totalsRowFunction')}
                            out_col = ET.SubElement(columns, NS + 'tableColumn', attrs)
                            item = {'name': col.get('name'), 'id': col.get('id')}
                            for formula_tag in ('calculatedColumnFormula', 'totalsRowFormula'):
                                expression = col.find(NS + formula_tag)
                                if expression is not None:
                                    out_col.append(clean_node(expression))
                                    item[formula_tag] = expression.text or ''
                                    review.formula(expression.text, f'{path.name} / {name} / {table.get("displayName")}[{col.get("name")}]')
                            column_models.append(item)
                        columns.set('count', str(len(column_models)))
                        style = table.find(NS + 'tableStyleInfo')
                        if style is not None:
                            safe = clean_node(style)
                            if not re.fullmatch(r'TableStyle(?:Light|Medium|Dark)\d+', safe.get('name', '')):
                                safe.set('name', 'TableStyleMedium2')
                            table_out.append(safe)
                        parts[target] = encoded(table_out)
                        content_types.append((target, 'application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml'))
                        sheet_model['tables'].append({'name': table.get('name'), 'displayName': table.get('displayName'), 'ref': table.get('ref'), 'columns': column_models})
                out.append(node)
            target = f'xl/worksheets/sheet{index}.xml'
            parts[target] = encoded(out)
            if sheet_rels:
                parts[f'xl/worksheets/_rels/sheet{index}.xml.rels'] = rels_bytes(sheet_rels)
            rid = f'rIdSheet{index}'
            ET.SubElement(sheets_node, NS + 'sheet', {'name': name, 'sheetId': str(index), 'state': state, RID: rid})
            workbook_rels.append({'Id': rid, 'Type': R + '/worksheet', 'Target': f'worksheets/sheet{index}.xml'})
            content_types.append((target, 'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'))
            model['sheets'].append(sheet_model)
            plans.append({'id': sheet_id, 'source': path.name, 'name': name, 'state': state, 'headerRanges': span_text,
                          'headers': sheet_model['headers'], 'formulaCount': len(sheet_model['formulas']),
                          'dataExtent': sheet_model['dataExtent'], 'tableCount': len(source_tables),
                          'codeName': sheet_model['codeName'], 'tables': sheet_model['tables']})
        defined = ET.SubElement(workbook, NS + 'definedNames')
        for item in source_book.findall(NS + 'definedNames/' + NS + 'definedName'):
            attrs = {k: v for k, v in item.attrib.items() if k in ('name', 'localSheetId', 'hidden')}
            ET.SubElement(defined, NS + 'definedName', attrs).text = item.text
            model['definedNames'].append({**attrs, 'formula': item.text or ''})
            review.formula(item.text, f'{path.name} / 名称 {item.get("name")}')
        external = source_book.find(NS + 'externalReferences')
        if external is not None:
            external_out = ET.SubElement(workbook, NS + 'externalReferences')
            for index, item in enumerate(external, 1):
                rel = source_rels.get(item.get(RID))
                if rel is None or not rel.get('Type', '').endswith('/externalLink') or rel.get('TargetMode') == 'External':
                    raise ValueError('外部引用关系无效，未导出')
                part = part_target('xl/workbook.xml', rel.get('Target'))
                linked = xml(archive.read(part))
                external_book = linked.find(NS + 'externalBook')
                if external_book is None:
                    raise ValueError('VBA 模板暂不支持 DDE/OLE 外部链接，未导出')
                source_links = relationships(archive, part)
                link = source_links.get(external_book.get(RID))
                if link is None or not link.get('Type', '').endswith('/externalLinkPath') or link.get('TargetMode') != 'External':
                    raise ValueError('外部工作簿引用无效，未导出')
                link_target = link.get('Target', '')
                review.add('外部路径', link_target, path.name)
                model['externalLinks'].append({'index': index, 'target': link_target})
                linked_out = ET.Element(NS + 'externalLink')
                external_book_out = ET.SubElement(linked_out, NS + 'externalBook', {RID: 'rId1'})
                for tag in ('sheetNames', 'definedNames'):
                    child = external_book.find(NS + tag)
                    if child is not None:
                        external_book_out.append(clean_node(child))
                        if tag == 'definedNames':
                            for definition in child:
                                review.formula(definition.get('refersTo'), path.name + ' / 外部名称 ' + definition.get('name', ''))
                target = f'xl/externalLinks/externalLink{index}.xml'
                parts[target] = encoded(linked_out)
                parts[f'xl/externalLinks/_rels/externalLink{index}.xml.rels'] = rels_bytes([{'Id': 'rId1', 'Type': R + '/externalLinkPath', 'Target': link_target, 'TargetMode': 'External'}])
                rid = f'rIdExternal{index}'
                ET.SubElement(external_out, NS + 'externalReference', {RID: rid})
                workbook_rels.append({'Id': rid, 'Type': R + '/externalLink', 'Target': f'externalLinks/externalLink{index}.xml'})
                content_types.append((target, 'application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml'))
        ET.SubElement(workbook, NS + 'calcPr', {'calcMode': 'manual', 'fullCalcOnLoad': '0', 'forceFullCalc': '0', 'calcOnSave': '0'})
        parts['xl/workbook.xml'] = encoded(workbook)
        styles, changed = safe_styles(archive.read('xl/styles.xml'))
        parts['xl/styles.xml'] = styles
        model['sanitizedNumberFormats'] = changed
        model['formulaCount'] = formula_count
        model['protected'] = source_book.find(NS + 'workbookProtection') is not None
        workbook_rels += [{'Id': 'rIdStyles', 'Type': R + '/styles', 'Target': 'styles.xml'}, {'Id': 'rIdTheme', 'Type': R + '/theme', 'Target': 'theme/theme1.xml'}]
        parts['xl/_rels/workbook.xml.rels'] = rels_bytes(workbook_rels)
        parts['_rels/.rels'] = rels_bytes([
            {'Id': 'rId1', 'Type': R + '/officeDocument', 'Target': 'xl/workbook.xml'},
            {'Id': 'rId2', 'Type': P + '/metadata/core-properties', 'Target': 'docProps/core.xml'},
            {'Id': 'rId3', 'Type': R + '/extended-properties', 'Target': 'docProps/app.xml'}])
        content_types += [('xl/workbook.xml', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml'),
                          ('xl/styles.xml', 'application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml'),
                          ('xl/theme/theme1.xml', 'application/vnd.openxmlformats-officedocument.theme+xml'),
                          ('docProps/core.xml', 'application/vnd.openxmlformats-package.core-properties+xml'),
                          ('docProps/app.xml', 'application/vnd.openxmlformats-officedocument.extended-properties+xml')]
        types = ET.Element('{' + C + '}Types')
        ET.SubElement(types, '{' + C + '}Default', {'Extension': 'rels', 'ContentType': 'application/vnd.openxmlformats-package.relationships+xml'})
        ET.SubElement(types, '{' + C + '}Default', {'Extension': 'xml', 'ContentType': 'application/xml'})
        for name, content_type in content_types:
            ET.SubElement(types, '{' + C + '}Override', {'PartName': '/' + name, 'ContentType': content_type})
        parts['[Content_Types].xml'] = encoded(types)
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return {'model': model, 'plans': plans, 'bytes': stream.getvalue(), 'hash': sha256(raw).hexdigest(), 'path': str(path), 'ids': names_seen}


def prepare_vba(config, request):
    """检查清单绑定文件内容及表头区域；导出时重新核对。"""
    overrides = request.get('headerRanges', {})
    if not isinstance(overrides, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in overrides.items()):
        raise ValueError('表头区域设置无效')
    paths = [Path(p).resolve() for p in config['selected_input_files']]
    if len(paths) > 16:
        raise ValueError('一次最多导出 16 个 VBA 结构工作簿')
    review = Review()
    books = [_build(p, index, config, overrides, review) for index, p in enumerate(paths, 1)]
    if set(overrides) - set().union(*(b['ids'] for b in books)):
        raise ValueError('工作簿选择已变化，请重新检查表头区域')
    identity = [{'path': b['path'], 'hash': b['hash'], 'headers': {s['id']: s['headerRanges'] for s in b['plans']}} for b in books]
    token = sha256(json.dumps(identity, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    plan = {'token': token, 'tables': [], 'kinds': {}, 'vba': {
        'workbooks': [{'id': b['model']['id'], 'name': b['model']['sourceName'], 'sheets': b['plans'],
                       'definedNames': b['model']['definedNames'], 'externalLinks': b['model']['externalLinks']} for b in books],
        'review': list(review.groups.values()), 'formulaCount': sum(b['model']['formulaCount'] for b in books),
        'scope': '所选文件的全部 Sheet（包括隐藏 Sheet），用于保留跨 Sheet 引用；非表头常量全部清空。'}}
    return books, plan


def inspect_vba(config, request):
    return prepare_vba(config, request)[1]


def ai_guide():
    return '''EasyExcel · 为原工作簿编写 VBA

目标：用户将你生成的完整 VBA 粘贴到原工作簿的标准模块，在原始真实数据上运行。
请先向用户确认具体业务目标、输入区域、输出位置、空值/重复值规则和 Excel 版本，再生成完整代码。

本包保留原 Sheet 名、顺序、隐藏状态、用户确认的表头文字及其地址、合并区域、行列位置、Excel 表对象名称/区域/列名、命名区域及原公式。
工作表单元格公式（含共享/数组公式）、表对象计算公式和命名公式保持原文，没有匿名、平移、裁剪或替换常量。
除确认的表头区域以外，真实常量与统计结果已清空。原公式的缓存也已清空。公式中的固定文本/数字以及外部路径经用户检查后保留，不能把它们当作虚构值。
模板包含全部 Sheet 结构（包括隐藏 Sheet）；真实数据并不在模板中。模板手动计算、禁止自动更新外部链接；不要为检查模板而计算、刷新连接或访问外部文件。
原 VBA 项目、图表/图片/透视对象、条件格式、数据验证、筛选条件、文档属性和外部连接未复制；基础格式已清理。结构说明记录原表的保护状态，模板本身未加锁。

VBA 编写约定：
1. 使用 Option Explicit。宏粘贴到原工作簿的标准模块；Set wb = ThisWorkbook。
   所有 Worksheets、Range、Cells、Rows 都限定 wb/ws，避免依赖 ActiveWorkbook、ActiveSheet、Select 或当前打开的模板。
2. 严格使用结构说明中的原 Sheet 名、表头文字、表对象名和命名区域；不要改成 Sheet_001 或“字段_001”，不要引用本包文件名。
3. 本包提供“运行前结构检查.bas”作为结构检查起点。它没有业务处理逻辑。生成最终代码时把检查函数和你的业务过程放在同一标准模块中。
   .bas 可在 VBE 中通过“文件 → 导入文件”导入；若提供直接粘贴的代码，去掉 Attribute VB_Name 行。
   不要把依赖 Excel 版本、保护口令、其他工作簿或缺失业务规则的代码描述为已经保证可以运行。
4. 在原工作簿运行时确定末行：优先使用正确 ListObject.DataBodyRange；普通区域选择有业务意义的必填列，处理全空/只有表头的情况。
   不要从空模板的 End(xlUp)、CurrentRegion、UsedRange 或预览 100 行推断真实数据范围；不要把本次行数写死为永久上限。
   并排表块分别处理；合并表头跨多行时，从结构说明的表头底部之后开始读取。
5. 公式列默认保留；筛选/排序/写入时避免把整块区域 Value2 覆盖回去导致公式丢失。需要新增公式时先明确 Formula/Formula2 和兼容版本。
   现有跨 Sheet、结构化引用、共享/数组公式保持语义；不要用静态值替换。外部工作簿引用应在用户本机核验，不自动联网。
6. 原代码若需图表、透视、验证规则或既有宏，请先询问具体对象要求；模板中缺少这些对象不表示原表中不存在。
7. 输出优先写新 Sheet 或用户指定位置；只读分析不修改源数据。用错误处理保证 ScreenUpdating、EnableEvents、Calculation 等状态恢复。
8. 如果必须验证业务逻辑，请另造完全虚构的测试数据，保持真实字段名与结构；不要从本包推断实际金额、分布或业务结论。

交付完整可复制模块、入口过程名、粘贴位置、运行步骤和必要的使用条件，不要求用户自行补齐省略的核心代码。
用户操作：先备份原工作簿，将 .xlsx 另存为支持宏的 .xlsm（已有 .xlsm 可直接使用）；Alt+F11 → 原工作簿工程 → 插入模块 → 粘贴完整代码 → 返回 Excel 运行入口宏。
本包只帮助 AI 对齐代码与原表结构；最终业务逻辑仍需在本机副本验证。

参考：
https://learn.microsoft.com/en-us/office/vba/api/excel.application.thisworkbook
https://learn.microsoft.com/en-us/office/vba/api/excel.range.end
'''


def vba_guard(model):
    """ASCII 模块用 ChrW 恢复原中文名称，避免 VBE 导入时编码乱码。"""
    def string(value):
        value = str(value)
        if all(32 <= ord(c) <= 126 for c in value):
            return '"' + value.replace('"', '""') + '"'
        data = value.encode('utf-16-be').hex().upper()
        return 'EE_U("' + data + '")'

    lines = ['Attribute VB_Name = "EasyExcelLayout"', 'Option Explicit', '',
             "' Import into the ORIGINAL workbook, then add the requested business routine.",
             'Public Function EE_CheckLayout(ByVal wb As Workbook) As Boolean',
             '    On Error GoTo InvalidLayout']
    for index in range(len(model['sheets'])):
        lines.append(f'    If Not EE_Sheet{index+1}(wb) Then GoTo InvalidLayout')
    lines += ['    EE_CheckLayout = True', '    Exit Function', 'InvalidLayout:',
              '    MsgBox "Workbook layout differs from the export. Check sheet names and headers.", vbExclamation',
              'End Function', '']
    for index, sheet in enumerate(model['sheets'], 1):
        lines += [f'Private Function EE_Sheet{index}(ByVal wb As Workbook) As Boolean', '    Dim ws As Worksheet',
                  '    On Error GoTo InvalidLayout', '    Set ws = wb.Worksheets(' + string(sheet['name']) + ')']
        # 大表完整表头清单保存在 JSON；模块只抽查有限定位点，避免 VBA 过程长度超限。
        selected = [h for h in sheet['headers'] if h['type'] == 'text' and len(h['value']) <= 80][:12]
        for header in selected:
            lines.append(f'    If CStr(ws.Range("{header["cell"]}").Value2) <> {string(header["value"])} Then GoTo InvalidLayout')
        lines += [f'    EE_Sheet{index} = True', 'InvalidLayout:', 'End Function', '']
    lines += ['Public Sub EE_CheckOnly()', '    If EE_CheckLayout(ThisWorkbook) Then MsgBox "Layout checks passed. No data was changed."',
              'End Sub', '', 'Private Function EE_U(ByVal hexText As String) As String',
              '    Dim i As Long', '    For i = 1 To Len(hexText) Step 4',
              '        EE_U = EE_U & ChrW(CLng("&H" & Mid$(hexText, i, 4)))', '    Next i', 'End Function', '']
    return '\r\n'.join(lines)


def export_vba(request, config):
    """写入完整结构包；用户必须检查名称、表头及公式常量，源文件内容不能变化。"""
    if request.get('reviewConfirmed') is not True:
        raise ValueError('请先检查表头、原公式固定值及外部引用，再勾选检查确认')
    export_id = str(UUID(request['exportId']))
    if export_id != request['exportId']:
        raise ValueError('导出标识无效')
    output = Path(request['outputPath']).resolve()
    if output.suffix.lower() != '.zip' or not output.parent.is_dir():
        raise ValueError('请选择 ZIP 文件保存位置')
    if output in [Path(p).resolve() for p in config['selected_input_files']]:
        raise ValueError('不能覆盖源文件')
    books, plan = prepare_vba(config, request)
    if request.get('token') != plan['token']:
        raise ValueError('文件内容或表头区域已变化，请重新生成检查清单后导出')
    schema = {'version': 1, 'id': export_id, 'mode': 'vba', 'dataRemoved': True, 'formulasUnchanged': True,
              'workbooks': [b['model'] for b in books], 'checks': plan['vba']['review']}
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, prefix='.easyexcel-vba-', suffix='.tmp', delete=False) as stream:
            temporary = Path(stream.name)
        with ZipFile(temporary, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('给AI的VBA编程说明.txt', ai_guide())
            archive.writestr('结构说明.json', json.dumps(schema, ensure_ascii=False, indent=2, allow_nan=False))
            for item in books:
                archive.writestr(item['model']['template'], item['bytes'])
                archive.writestr(item['model']['id'] + '/运行前结构检查.bas', vba_guard(item['model']).encode('ascii'))
        os.replace(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {'output': str(output), 'exportId': export_id, 'sheets': sum(len(b['plans']) for b in books), 'mode': 'vba',
            'sampleRows': 0, 'workbooks': len(books), 'formulaCount': plan['vba']['formulaCount'],
            'warnings': ['原 Sheet 名、表头、公式固定值及外部引用按检查结果保留；这个模式不是完全匿名导出。',
                         '空模板供编程参考，不用于覆盖原表；AI 代码需粘贴到原工作簿标准模块，并在本机副本验证。']}
